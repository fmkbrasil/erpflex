from __future__ import annotations

import io
import json
import os
import smtplib
import ssl
import threading
import time
from email.message import EmailMessage
from email.utils import formataddr
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import quote_plus, urlencode

from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import and_, func, or_, select
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from starlette.middleware.sessions import SessionMiddleware

from .connectors import ERPFlexClient, NFStockClient
from .db import (
    AccessProfile, AppUser, AuditLog, BoxType, Branch, Company, CollectionAutomationConfig, CollectionCustomerRule, CollectionSendLog, Delivery, DeliveryAssignment, DeliveryMovement,
    DeliveryOccurrence, DeliveryRule, FinancialTitle, LogisticsImport, Partner, Product, ProfilePermission, PurchaseInvoice, PurchaseItem, RawRecord, SalesInvoiceItem,
    Region, Responsible, Route, RouteCost, RouteStop, SalesInvoice, SalesOrder, SessionLocal, SyncAutomationConfig, SyncAutomationModule, SyncProgress, SyncRun, SyncWorkerHeartbeat,
    UserBranchAccess, UserCompanyAccess, UserViewPreference, create_sqlite_backup, database_status, init_db
)
from .access import access_context, add_audit, auth_configured, can, current_user
from .security import hash_password, verify_password
from .sync_service import start_sync, start_sync_batch, recover_interrupted_sync_runs, page_index_summary, _purchase_items_from_record, _purchase_item_value
from .integration_settings import (
    encryption_fingerprint, encryption_is_default, encryption_source, get_erpflex_settings, get_nfstock_settings, get_smtp_settings,
    integration_public_view, mark_test_result, save_erpflex, save_nfstock, save_smtp,
)
from .logistics import alerta_endereco, classificar_regiao, cliente_exibicao, endereco_busca, maps_url, raw_channel, regiao_especial, waze_url
from .logistics_import import import_logistics_file
from .utils import deep_find, fnum, text, access_key_from, normalize_date, payload_hash, safe_json

BASE_DIR = Path(__file__).resolve().parent
APP_NAME = os.getenv("APP_NAME") or "Plataforma Gestão Integrada"
APP_USER = os.getenv("APP_USER") or "admin"
APP_PASSWORD = os.getenv("APP_PASSWORD") or ""
SECRET_KEY = os.getenv("SECRET_KEY") or "desenvolvimento-troque-em-producao"

app = FastAPI(title=APP_NAME, version="1.3.30")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))
templates.env.globals["cliente_exibicao"] = cliente_exibicao
templates.env.globals["endereco_busca"] = endereco_busca
templates.env.globals["maps_url"] = maps_url
templates.env.globals["waze_url"] = waze_url
templates.env.globals["alerta_endereco"] = alerta_endereco


@app.on_event("startup")
def startup():
    global _collection_worker_started
    init_db()
    recover_interrupted_sync_runs()
    if not _collection_worker_started:
        threading.Thread(target=_collection_worker, name="collection-auto", daemon=True).start()
        _collection_worker_started = True


@app.middleware("http")
async def auth_middleware(request: Request, call_next):
    path = request.url.path
    public = path in {"/login", "/health"} or path.startswith("/static/")
    with SessionLocal() as db:
        configured = auth_configured(db)
        if not public and configured and not request.session.get("user_id"):
            if request.method == "GET" and not path.startswith("/api/"):
                return RedirectResponse("/login", status_code=303)
            return JSONResponse({"error": "não autenticado"}, status_code=401)
        if not public and configured:
            try:
                ctx = access_context(db, request)
            except Exception:
                request.session.clear()
                return RedirectResponse("/login", status_code=303)
            required = required_permission(request.method, path)
            if required and not can(ctx, required):
                if path.startswith("/api/"):
                    return JSONResponse({"error": "sem permissão", "permission": required}, status_code=403)
                return HTMLResponse("Acesso não autorizado para este módulo.", status_code=403)
    return await call_next(request)


# SessionMiddleware precisa envolver o middleware de autenticação para request.session existir.
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, same_site="lax", https_only=False)


def required_permission(method: str, path: str) -> str | None:
    method = method.upper()
    rules = [
        ("/administracao", "admin.manage"),
        ("/sincronizacao", "sync.manage"),
        ("/integracoes", "sync.manage"),
        ("/sistema/backup", "admin.manage"),
        ("/cadastros", "cadastros.view"),
        ("/faturamento", "faturamento.view"),
        ("/financeiro", "financeiro.view"),
        ("/compras", "compras.view"),
        ("/importar-logistica", "importar.manage"),
        ("/roteirizar", "roteirizar.manage" if method != "GET" else "roteirizar.view"),
        ("/painel-entregas", "painel_entregas.view"),
        ("/painel-gerencial", "painel_gerencial.view"),
        ("/relatorios-gerenciais", "relatorios.view"),
        ("/logistica", "logistica.edit" if method != "GET" else "logistica.view"),
        ("/roteiros", "roteiros.manage" if method != "GET" else "roteiros.view"),
        ("/baixa-entregas", "baixa.manage"),
        ("/entregas/", "baixa.manage" if method != "GET" else "logistica.view"),
        ("/ocorrencias", "ocorrencias.manage" if method != "GET" else "ocorrencias.view"),
        ("/responsaveis", "responsaveis.manage" if method != "GET" else "responsaveis.view"),
        ("/regras-entrega", "regras.manage" if method != "GET" else "logistica.view"),
        ("/parametros-logistica", "parametros.manage" if method != "GET" else "logistica.view"),
        ("/custos-roteiro", "roteiros.manage" if method != "GET" else "roteiros.view"),
        ("/diagnostico", "diagnostico.view"),
        ("/api/v1/", "api.read"),
        ("/api/sync/", "sync.manage"),
        ("/docs", "diagnostico.view"),
        ("/openapi.json", "diagnostico.view"),
    ]
    if path == "/":
        return "dashboard.view"
    for prefix, perm in rules:
        if path.startswith(prefix):
            return perm
    return None


def render(request: Request, template: str, **ctx):
    with SessionLocal() as db:
        configured = auth_configured(db)
        ac = None
        if not configured or request.session.get("user_id"):
            try:
                ac = access_context(db, request)
            except Exception:
                ac = None
    if ac:
        ctx.update({
            "company": ctx.get("company") or ac.company, "branch": ac.branch, "current_user": ac.user,
            "current_profile": ac.profile, "permissions": ac.permissions, "available_companies": ac.companies,
            "available_branches": ac.branches, "can_perm": lambda p: can(ac, p),
        })
    else:
        ctx.update({"branch": None, "current_user": None, "current_profile": None, "permissions": set(),
                    "available_companies": [], "available_branches": [], "can_perm": lambda p: False})
    ctx.update({"request": request, "app_name": APP_NAME, "auth_enabled": configured})
    return templates.TemplateResponse(template, ctx)


def company_context(db, request: Request):
    ac = access_context(db, request)
    return ac.company.id, ac.company


@app.get("/health")
def health():
    return {"ok": True, "version": "1.3.30", "schema": database_status().get("schema_version"), "erpflex_engine": (os.getenv("ERPFLEX_ENGINE") or "go_v78")}


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return render(request, "login.html", error=None)


@app.post("/login")
def login(request: Request, user: str = Form(""), password: str = Form("")):
    with SessionLocal() as db:
        row = db.scalar(select(AppUser).where(func.lower(AppUser.username) == user.strip().lower(), AppUser.active.is_(True)).limit(1))
        if row and verify_password(password, row.password_hash):
            request.session.clear()
            request.session["user_id"] = row.id
            row.last_login_at = datetime.utcnow()
            db.commit()
            ac = access_context(db, request)
            request.session["company_id"] = ac.company.id
            if ac.branch:
                request.session["branch_id"] = ac.branch.id
            return RedirectResponse("/", status_code=303)
    return render(request, "login.html", error="Usuário ou senha inválidos.")


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        counts = {
            "orders": db.scalar(select(func.count(SalesOrder.id)).where(SalesOrder.company_id == cid)) or 0,
            "sales_invoices": db.scalar(select(func.count(SalesInvoice.id)).where(SalesInvoice.company_id == cid)) or 0,
            "purchase_invoices": db.scalar(select(func.count(PurchaseInvoice.id)).where(PurchaseInvoice.company_id == cid)) or 0,
            "products": db.scalar(select(func.count(Product.id)).where(Product.company_id == cid)) or 0,
            "partners": db.scalar(select(func.count(Partner.id)).where(Partner.company_id == cid)) or 0,
            "deliveries": db.scalar(select(func.count(Delivery.id)).where(Delivery.company_id == cid)) or 0,
            "receivable": db.scalar(select(func.count(FinancialTitle.id)).where(FinancialTitle.company_id == cid, FinancialTitle.kind == "RECEBER")) or 0,
            "payable": db.scalar(select(func.count(FinancialTitle.id)).where(FinancialTitle.company_id == cid, FinancialTitle.kind.in_(["PAGAR", "DESPESA"]))) or 0,
        }
        recon = dict(db.execute(
            select(PurchaseInvoice.reconciliation_status, func.count(PurchaseInvoice.id))
            .where(PurchaseInvoice.company_id == cid)
            .group_by(PurchaseInvoice.reconciliation_status)
        ).all())
        last_runs = db.scalars(select(SyncRun).where(SyncRun.company_id == cid).order_by(SyncRun.id.desc()).limit(8)).all()
        latest = db.scalar(select(SyncRun).where(SyncRun.company_id == cid).order_by(SyncRun.id.desc()).limit(1))
    return render(request, "dashboard.html", company=company, counts=counts, recon=recon, runs=last_runs, latest=latest)


@app.get("/sincronizacao", response_class=HTMLResponse)
def sync_page(request: Request):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        runs = db.scalars(select(SyncRun).where(SyncRun.company_id == cid).order_by(SyncRun.id.desc()).limit(30)).all()
        integrations = integration_public_view(db, cid)
        page_index = page_index_summary(cid)
        worker_cfg = db.scalar(select(SyncAutomationConfig).where(SyncAutomationConfig.company_id == cid).limit(1))
        worker_modules = db.scalars(select(SyncAutomationModule).where(
            SyncAutomationModule.company_id == cid
        ).order_by(SyncAutomationModule.id.asc())).all()
        worker_hb = db.scalar(select(SyncWorkerHeartbeat).where(
            SyncWorkerHeartbeat.company_id == cid
        ).order_by(SyncWorkerHeartbeat.heartbeat_at.desc()).limit(1))
        worker_view = {"online": False, "age_seconds": None, "status": "SEM WORKER", "message": "Worker ainda não iniciou."}
        if worker_hb:
            age = max(0, int((datetime.utcnow() - worker_hb.heartbeat_at).total_seconds()))
            worker_view = {
                "online": age <= 90, "age_seconds": age, "status": worker_hb.status,
                "message": worker_hb.message, "current_module": worker_hb.current_module,
                "current_run_id": worker_hb.current_run_id, "instance_id": worker_hb.instance_id,
                "heartbeat_at": worker_hb.heartbeat_at,
            }
    today = date.today()
    return render(
        request, "sync.html", company=company, runs=runs,
        erp_configured=integrations["erpflex"]["configured"],
        nf_configured=integrations["nfstock"]["configured"],
        integrations=integrations, page_index=page_index, worker_cfg=worker_cfg,
        worker_modules=worker_modules, worker_view=worker_view,
        default_start=(today - timedelta(days=7)).isoformat(), default_end=today.isoformat(),
    )


@app.post("/sincronizacao/worker-config")
async def save_sync_worker_config(request: Request):
    form = await request.form()
    allowed = ["products", "banks", "clientes", "orders", "faturamento", "receber", "pagar", "compras", "despesas"]
    enabled = set(form.getlist("modules"))
    now = datetime.utcnow()
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        cfg = db.scalar(select(SyncAutomationConfig).where(SyncAutomationConfig.company_id == cid).limit(1))
        if not cfg:
            cfg = SyncAutomationConfig(company_id=cid)
            db.add(cfg)
            db.flush()
        cfg.active = bool(form.get("active"))
        def iv(name, default, lo, hi):
            try: return max(lo, min(hi, int(form.get(name) or default)))
            except Exception: return default
        cfg.poll_seconds = iv("poll_seconds", 30, 5, 300)
        cfg.allowed_start_hour = iv("allowed_start_hour", 0, 0, 23)
        cfg.allowed_end_hour = iv("allowed_end_hour", 23, 0, 23)
        cfg.max_blocks = iv("max_blocks", 100, 1, 1000)
        cfg.master_max_pages = iv("master_max_pages", 500, 1, 2000)
        cfg.max_clients = iv("max_clients", 1000, 1, 5000)
        cfg.sales_details = bool(form.get("sales_details"))
        cfg.purchase_details = bool(form.get("purchase_details"))
        for module in allowed:
            row = db.scalar(select(SyncAutomationModule).where(
                SyncAutomationModule.company_id == cid, SyncAutomationModule.module == module
            ).limit(1))
            if not row:
                row = SyncAutomationModule(company_id=cid, module=module)
                db.add(row)
                db.flush()
            was_enabled = bool(row.enabled)
            row.enabled = module in enabled
            row.interval_minutes = iv(f"interval_{module}", row.interval_minutes or 30, 1, 10080)
            if row.enabled and (not was_enabled or row.next_run_at is None):
                row.next_run_at = now
        db.commit()
    return RedirectResponse("/sincronizacao?worker=saved", status_code=303)


@app.post("/sincronizacao/worker-executar")
async def queue_sync_worker_now(request: Request):
    form = await request.form()
    requested = set(form.getlist("modules"))
    now = datetime.utcnow()
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        rows = db.scalars(select(SyncAutomationModule).where(
            SyncAutomationModule.company_id == cid, SyncAutomationModule.enabled.is_(True)
        )).all()
        for row in rows:
            if not requested or row.module in requested:
                row.next_run_at = now
        db.commit()
    return RedirectResponse("/sincronizacao?worker=queued", status_code=303)


@app.get("/integracoes", response_class=HTMLResponse)
def integracoes(request: Request, saved: str = "", tested: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        integrations = integration_public_view(db, cid)
    return render(
        request, "integracoes.html", company=company, integrations=integrations, db_status=database_status(),
        encryption_fingerprint=encryption_fingerprint(), encryption_default=encryption_is_default(), encryption_source=encryption_source(),
        saved=saved, tested=tested,
    )


@app.post("/integracoes/erpflex")
def integracoes_erpflex(
    request: Request, base_url: str = Form("https://api.erpflex.com.br"), username: str = Form(""),
    password: str = Form(""), active: str | None = Form(None),
):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        row = save_erpflex(db, cid, base_url=base_url, username=username, password=password, active=bool(active))
        ac = access_context(db, request)
        add_audit(db, ac, "INTEGRACAO_SALVA", "integration", row.id, "ERPFlex atualizado; segredo não registrado no log")
        db.commit()
    return RedirectResponse("/integracoes?saved=erpflex", status_code=303)


@app.post("/integracoes/nfstock")
def integracoes_nfstock(
    request: Request, base_url: str = Form("https://ms-exportacao-nfstock.pack.alterdata.com.br"),
    token: str = Form(""), crm: str = Form(""), cnpj: str = Form(""), page_size: int = Form(25),
    active: str | None = Form(None),
):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        row = save_nfstock(db, cid, base_url=base_url, token=token, crm=crm, cnpj=cnpj, page_size=page_size, active=bool(active))
        ac = access_context(db, request)
        add_audit(db, ac, "INTEGRACAO_SALVA", "integration", row.id, "NF Stock atualizado; token não registrado no log")
        db.commit()
    return RedirectResponse("/integracoes?saved=nfstock", status_code=303)


@app.post("/integracoes/smtp")
def integracoes_smtp(
    request: Request, host: str = Form(""), port: int = Form(587), username: str = Form(""),
    password: str = Form(""), from_email: str = Form(""), from_name: str = Form(""),
    security: str = Form("starttls"), active: str | None = Form(None),
):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        row = save_smtp(db, cid, host=host, port=port, username=username, password=password,
                        from_email=from_email, from_name=from_name, security=security, active=bool(active))
        ac = access_context(db, request)
        add_audit(db, ac, "INTEGRACAO_SALVA", "email_config", row.id, "SMTP atualizado; senha não registrada no log")
        db.commit()
    return RedirectResponse("/integracoes?saved=smtp", status_code=303)


def _smtp_connect(settings: dict):
    if not settings.get("active") or not settings.get("host"):
        raise RuntimeError("SMTP não configurado/ativo.")
    host = settings["host"]; port = int(settings.get("port") or 587)
    security = (settings.get("security") or "starttls").lower()
    if security == "ssl":
        server = smtplib.SMTP_SSL(host, port, timeout=25, context=ssl.create_default_context())
    else:
        server = smtplib.SMTP(host, port, timeout=25)
        server.ehlo()
        if security == "starttls":
            server.starttls(context=ssl.create_default_context()); server.ehlo()
    if settings.get("username"):
        server.login(settings.get("username"), settings.get("password") or "")
    return server


def _smtp_send_boleto(settings: dict, recipient: str, subject: str, body: str, boleto_html: str, filename: str, xml_text: str = "", xml_filename: str = "") -> None:
    msg = EmailMessage()
    from_email = settings.get("from_email") or settings.get("username") or ""
    from_name = settings.get("from_name") or ""
    msg["From"] = formataddr((from_name, from_email)) if from_name else from_email
    msg["To"] = recipient
    msg["Subject"] = subject
    msg.set_content(body)
    msg.add_attachment(boleto_html.encode("utf-8"), maintype="text", subtype="html", filename=filename)
    if xml_text:
        msg.add_attachment(xml_text.encode("utf-8"), maintype="application", subtype="xml", filename=(xml_filename or "nfe.xml"))
    server = _smtp_connect(settings)
    try:
        server.send_message(msg)
    finally:
        try: server.quit()
        except Exception: server.close()


@app.post("/sistema/backup")
def sistema_backup(request: Request):
    with SessionLocal() as db:
        _cid, _company = company_context(db, request)
        ac = access_context(db, request)
        path = create_sqlite_backup("manual")
        add_audit(db, ac, "BACKUP_MANUAL", "database", None, str(path) if path else "Banco não é SQLite ou arquivo inexistente")
        db.commit()
    return RedirectResponse("/integracoes?tested=backup-ok" if path else "/integracoes?tested=backup-nao-aplicavel", status_code=303)


@app.post("/sincronizacao/erpflex")
def start_erpflex_sync(
    request: Request,
    module: str = Form(...),
    max_blocks: int = Form(20),
    start_offset: int = Form(0),
    start_page: int = Form(1),
    start_pos: int = Form(1),
    mode: str = Form("recent"),
    purchase_details: str | None = Form(None),
    sales_details: str | None = Form(None),
):
    allowed = {"products", "orders", "faturamento", "receber", "pagar", "compras", "despesas"}
    if module not in allowed:
        return JSONResponse({"error": "módulo inválido"}, status_code=400)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
    run_id = start_sync(cid, "ERPFLEX", module, {
        "max_blocks": max_blocks, "start_offset": start_offset, "start_page": start_page, "start_pos": start_pos,
        "mode": mode, "purchase_details": purchase_details, "sales_details": sales_details
    })
    return RedirectResponse(f"/sincronizacao?run={run_id}", status_code=303)


@app.post("/sincronizacao/erpflex-lote")
def start_erpflex_batch_route(
    request: Request,
    modules: list[str] = Form([]),
    action: str = Form("selected"),
    start_date: str = Form(""),
    end_date: str = Form(""),
    max_blocks: int = Form(100),
    master_max_pages: int = Form(500),
    max_clients: int = Form(1000),
    purchase_details: str | None = Form(None),
    sales_details: str | None = Form(None),
):
    ordered_allowed = ["products", "banks", "orders", "faturamento", "receber", "pagar", "compras", "despesas", "clientes"]
    if action == "all":
        selected = list(ordered_allowed)
    else:
        wanted = set(modules or [])
        selected = [m for m in ordered_allowed if m in wanted]
    if not selected:
        return RedirectResponse("/sincronizacao?error=selecione-modulos", status_code=303)
    try:
        d1 = date.fromisoformat(start_date)
        d2 = date.fromisoformat(end_date)
        if d2 < d1:
            start_date, end_date = end_date, start_date
    except Exception:
        return RedirectResponse("/sincronizacao?error=periodo-invalido", status_code=303)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
    run_ids = start_sync_batch(cid, selected, {
        "mode": "period",
        "start_date": start_date,
        "end_date": end_date,
        "max_blocks": max(1, min(1000, max_blocks)),
        "master_max_pages": max(1, min(2000, master_max_pages)),
        "max_clients": max(1, min(5000, max_clients)),
        "purchase_details": purchase_details,
        "sales_details": sales_details,
    })
    if not run_ids:
        return RedirectResponse("/sincronizacao?error=nenhuma-execucao", status_code=303)
    return RedirectResponse("/sincronizacao?runs=" + ",".join(str(x) for x in run_ids), status_code=303)


@app.post("/sincronizacao/erpflex-forcar")
def force_erpflex_sync(
    request: Request,
    modules: list[str] = Form([]),
    start_date: str = Form(...),
    end_date: str = Form(...),
    max_blocks: int = Form(300),
    purchase_details: str | None = Form(None),
    sales_details: str | None = Form(None),
):
    """Reprocessa deliberadamente um período sem avançar/retroceder o cursor principal.

    O índice local de páginas é usado como atalho quando já conhece o período.
    Se ainda não conhece, o mesmo motor V7.8 localiza o intervalo e passa a
    alimentar o índice para as próximas execuções.
    """
    ordered_allowed = ["orders", "faturamento", "receber", "pagar", "compras", "despesas"]
    wanted = set(modules or [])
    selected = [m for m in ordered_allowed if m in wanted]
    if not selected:
        return RedirectResponse("/sincronizacao?error=selecione-modulos-forcar", status_code=303)
    try:
        d1 = date.fromisoformat(start_date)
        d2 = date.fromisoformat(end_date)
        if d2 < d1:
            start_date, end_date = end_date, start_date
    except Exception:
        return RedirectResponse("/sincronizacao?error=periodo-invalido", status_code=303)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
    run_ids = start_sync_batch(cid, selected, {
        "mode": "force_period",
        "start_date": start_date,
        "end_date": end_date,
        "max_blocks": max(1, min(2000, max_blocks)),
        "purchase_details": purchase_details,
        "sales_details": sales_details,
    })
    if not run_ids:
        return RedirectResponse("/sincronizacao?error=nenhuma-execucao", status_code=303)
    return RedirectResponse("/sincronizacao?forced=1&runs=" + ",".join(str(x) for x in run_ids), status_code=303)


@app.post("/sincronizacao/nfstock")
def start_nfstock_sync(
    request: Request,
    start_date: str = Form(...),
    end_date: str = Form(...),
    max_pages: int = Form(100),
):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
    run_id = start_sync(cid, "NFSTOCK", "nfe", {"start_date": start_date, "end_date": end_date, "max_pages": max_pages})
    return RedirectResponse(f"/sincronizacao?run={run_id}", status_code=303)


@app.post("/sincronizacao/testar/{source}")
def test_connection(request: Request, source: str):
    source_key = source.lower()
    provider = "ERPFLEX" if source_key == "erpflex" else ("NFSTOCK" if source_key == "nfstock" else ("SMTP" if source_key == "smtp" else ""))
    if not provider:
        return JSONResponse({"error": "fonte inválida"}, status_code=400)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        try:
            if provider == "ERPFLEX":
                result = ERPFlexClient(get_erpflex_settings(db, cid)).test()
            elif provider == "NFSTOCK":
                result = NFStockClient(get_nfstock_settings(db, cid)).test()
            else:
                cfg = get_smtp_settings(db, cid)
                server = _smtp_connect(cfg)
                try:
                    server.noop()
                finally:
                    try: server.quit()
                    except Exception: server.close()
                result = {"ok": True, "message": "Conexão SMTP autenticada com sucesso."}
            if provider in {"ERPFLEX", "NFSTOCK"}:
                msg = f"Conexão OK. {result}"
                mark_test_result(db, cid, provider, True, msg)
                db.commit()
            return JSONResponse(result)
        except Exception as e:
            if provider in {"ERPFLEX", "NFSTOCK"}:
                mark_test_result(db, cid, provider, False, str(e)); db.commit()
            return JSONResponse({"ok": False, "error": str(e)}, status_code=400)


@app.get("/api/sync/{run_id}")
def sync_status(request: Request, run_id: int):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        r = db.get(SyncRun, run_id)
        if not r or r.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        progress = db.get(SyncProgress, r.id)
        now = datetime.utcnow()
        return {
            "id": r.id, "source": r.source, "module": r.module, "status": r.status,
            "found": r.total_found, "processed": r.total_processed, "inserted": r.total_inserted,
            "updated": r.total_updated, "errors": r.total_errors,
            "step": r.current_step, "message": r.message,
            "started_at": r.started_at.isoformat() if r.started_at else None,
            "finished_at": r.finished_at.isoformat() if r.finished_at else None,
            "server_time": now.isoformat(),
            "progress": {
                "percent": round(float(progress.percent or 0.0), 1) if progress else (100.0 if r.status not in {"PENDENTE", "EM_EXECUCAO"} else 0.0),
                "phase": progress.phase if progress else None,
                "current": progress.current if progress else None,
                "total": progress.total if progress else None,
                "detail": progress.detail if progress else None,
                "heartbeat_at": progress.heartbeat_at.isoformat() if progress and progress.heartbeat_at else None,
                "stage_updated_at": progress.stage_updated_at.isoformat() if progress and progress.stage_updated_at else None,
            },
        }


DOCUMENT_PER_PAGE = {25, 50, 100, 250}

DOCUMENT_COLUMN_DEFS = {
    "faturamento": [
        ("number", "NF-e"), ("order_number", "Pedido"), ("customer", "Cliente"), ("products", "Produtos"), ("items", "Itens"),
        ("issue_date", "Emissão"), ("status_nf", "Situação NF-e"), ("payment_status", "Status cobrança"), ("payment", "Financeiro"),
        ("forecast_date", "Previsão/entrega"), ("total", "Valor"), ("carrier", "Transportadora"),
        ("city", "Município"), ("volumes", "Volumes"), ("access_key", "Chave NF-e"),
    ],
    "compras": [
        ("number", "NF"), ("supplier", "Fornecedor"), ("products", "Produtos"), ("items", "Itens"), ("issue_date", "Emissão"),
        ("series", "Série"), ("access_key", "Chave NF-e"), ("total_erpflex", "Valor ERPFlex"),
        ("total_nfstock", "Valor NF Stock"), ("payment_status", "Status pagamento"), ("status", "Conciliação"),
    ],
    "financeiro": [
        ("title_number", "Título"), ("nf_number", "NF-e"), ("order_number", "Pedido"),
        ("partner", "Cliente/Fornecedor"), ("issue_date", "Emissão"), ("due_date", "Vencimento"),
        ("bank", "Banco"), ("wallet", "Carteira"), ("value", "Valor"), ("paid_value", "Pago"),
        ("balance", "Saldo"), ("status", "Status"), ("remittance", "Remessa gerada"),
    ],
}

DOCUMENT_DEFAULT_COLS = {
    "faturamento": ["number", "order_number", "customer", "products", "items", "issue_date", "status_nf", "payment_status", "payment", "total", "volumes"],
    "compras": ["number", "supplier", "products", "items", "issue_date", "total_erpflex", "payment_status", "status"],
    "financeiro": ["title_number", "nf_number", "order_number", "partner", "issue_date", "due_date", "bank", "wallet", "value", "balance", "remittance"],
}


def _module_pref(db, request: Request, company_id: int, view_key: str):
    ac = access_context(db, request)
    if not ac.user:
        return None
    return db.scalar(select(UserViewPreference).where(
        UserViewPreference.user_id == ac.user.id,
        UserViewPreference.company_id == company_id,
        UserViewPreference.view_key == view_key,
    ).limit(1))


def _module_view_settings(db, request: Request, company_id: int, module: str, cols: str, per_page: int, save_view: int):
    defs = DOCUMENT_COLUMN_DEFS[module]
    valid = dict(defs)
    pref = _module_pref(db, request, company_id, f"documentos:{module}:v2")
    pref_cols = []
    if pref and pref.columns_json:
        try:
            pref_cols = [x for x in json.loads(pref.columns_json) if x in valid]
        except Exception:
            pref_cols = []
    requested = [x for x in (cols or "").split(",") if x in valid]
    selected = requested or pref_cols or DOCUMENT_DEFAULT_COLS[module]
    # Exibe novos indicadores financeiros também para usuários com preferência salva de versões anteriores.
    if not requested and module in {"faturamento", "compras"} and "payment_status" in valid and "payment_status" not in selected:
        anchor = "payment" if module == "faturamento" else "status"
        pos = selected.index(anchor) if anchor in selected else len(selected)
        selected = selected[:pos] + ["payment_status"] + selected[pos:]
    pp = per_page if per_page in DOCUMENT_PER_PAGE else (pref.per_page if pref and pref.per_page in DOCUMENT_PER_PAGE else 50)
    if save_view and access_context(db, request).user:
        ac = access_context(db, request)
        if not pref:
            pref = UserViewPreference(user_id=ac.user.id, company_id=company_id, view_key=f"documentos:{module}:v2")
            db.add(pref)
        pref.columns_json = json.dumps(selected, ensure_ascii=False)
        pref.per_page = pp
        db.commit()
    return defs, selected, pp


def _date_from_text(value: str | None):
    value = (value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value[:10], fmt).date()
        except Exception:
            pass
    return None


def _period_bounds(period: str, date_from: str = "", date_to: str = ""):
    period = (period or "all").lower()
    if period not in {"all", "today", "week", "month", "custom"}:
        period = "all"
    today = date.today()
    start = end = None
    if period == "today":
        start = end = today
    elif period == "week":
        start = today - timedelta(days=today.weekday())
        end = start + timedelta(days=6)
    elif period == "month":
        start = today.replace(day=1)
        if start.month == 12:
            nxt = start.replace(year=start.year + 1, month=1)
        else:
            nxt = start.replace(month=start.month + 1)
        end = nxt - timedelta(days=1)
    elif period == "custom":
        start = _date_from_text(date_from)
        end = _date_from_text(date_to)
        if start and not end:
            end = start
        if end and not start:
            start = end
        if start and end and start > end:
            start, end = end, start
    return period, start, end


def _apply_date_range(stmt, column, start, end):
    if start:
        stmt = stmt.where(column >= start.isoformat())
    if end:
        stmt = stmt.where(column <= end.isoformat())
    return stmt


def _period_label(period: str, start, end):
    labels = {"all": "Todos", "today": "Hoje", "week": "Semana", "month": "Mês", "custom": "Personalizado"}
    if start and end:
        return f"{labels.get(period, period)} · {start.strftime('%d/%m/%Y')} a {end.strftime('%d/%m/%Y')}"
    return labels.get(period, period)


def _page_urls(path: str, page: int, pages: int, params: dict):
    base = {k: v for k, v in params.items() if v not in (None, "", 0)}
    def make(p):
        q = dict(base); q["page"] = p
        return f"{path}?{urlencode(q)}"
    return make(max(1, page-1)) if page > 1 else None, make(min(pages, page+1)) if page < pages else None


def _load_raw_payload(db, raw_id: int | None):
    if not raw_id:
        return {}
    raw = db.get(RawRecord, raw_id)
    if not raw or not raw.payload_json:
        return {}
    try:
        return json.loads(raw.payload_json)
    except Exception:
        return {}


def _candidate_item_lists(node):
    found = []
    if isinstance(node, dict):
        for key, value in node.items():
            if isinstance(value, list) and value and all(isinstance(x, dict) for x in value):
                found.append((str(key).lower(), value))
            found.extend(_candidate_item_lists(value))
    elif isinstance(node, list):
        for value in node[:20]:
            found.extend(_candidate_item_lists(value))
    return found


def _sales_items_from_payload(payload: dict):
    best, best_score = [], -1
    hints = ("item", "itens", "produto", "produtos", "detalhe", "detalhes")
    for key, values in _candidate_item_lists(payload):
        sample = values[:5]
        score = (8 if any(h in key for h in hints) else 0) + min(len(values), 10)
        for item in sample:
            keys = " ".join(str(k).lower() for k in item.keys())
            score += sum(3 for h in ("desc_produto", "descricao", "produto_id", "quant", "preco", "valor", "ncm", "cfop") if h in keys)
        if score > best_score:
            best, best_score = values, score
    rows = []
    for it in best:
        rows.append({
            "item_id": text(deep_find(it, ("item_id", "id_item", "id"))),
            "product_id": text(deep_find(it, ("produto_id", "id_produto", "product_id"))),
            "code": text(deep_find(it, ("codigo_produto", "cod_produto", "produto_codigo", "codigo", "code", "sku", "referencia"))),
            "description": text(deep_find(it, ("desc_produto", "descricao_produto", "produto_desc", "descricao_item", "descricao", "description", "produto", "nome"))),
            "ean": text(deep_find(it, ("EAN", "ean", "gtin", "codigo_barras", "cEAN"))),
            "ncm": text(deep_find(it, ("ncm", "NCM", "codigo_ncm"))),
            "cfop": text(deep_find(it, ("cfop", "CFOP", "cod_cfop"))),
            "unit": text(deep_find(it, ("unidade", "un", "unit", "ucom", "uCom"))),
            "qty": fnum(deep_find(it, ("quantidade", "qtd", "qty", "qtde", "qcom", "qCom"))),
            "unit_price": fnum(deep_find(it, ("preco_unitario", "valor_unitario", "unit_price", "vl_unitario", "preco", "vuncom", "vUnCom"))),
            "total": fnum(deep_find(it, ("preco_total_item", "valor_total", "total", "valor_item", "vprod", "vProd"))),
            "icms": fnum(deep_find(it, ("valor_icms", "icms", "vICMS"))),
            "ipi": fnum(deep_find(it, ("valor_ipi", "ipi", "vIPI"))),
        })
    return rows


def _catalog_product(db, cid: int, product_id: str = "", code: str = "", ean: str = ""):
    conds = []
    if product_id:
        conds.append(Product.erpflex_id == str(product_id))
    if code:
        conds.append(Product.code == str(code))
    if ean:
        conds.append(Product.ean == str(ean))
    if not conds:
        return None
    return db.scalar(select(Product).where(Product.company_id == cid, or_(*conds)).order_by(Product.id.desc()).limit(1))


def _enrich_item_from_catalog(db, cid: int, item: dict) -> dict:
    row = dict(item)
    product = _catalog_product(db, cid, text(row.get("product_id")), text(row.get("code")), text(row.get("ean")))
    if product:
        row["code"] = text(row.get("code")) or text(product.code)
        row["description"] = text(row.get("description")) or text(product.description)
        row["ean"] = text(row.get("ean")) or text(product.ean)
        row["ncm"] = text(row.get("ncm")) or text(product.ncm)
        row["unit"] = text(row.get("unit")) or text(product.unit)
    return row


def _persist_sales_invoice_items(db, cid: int, inv: SalesInvoice, payload: dict) -> int:
    items = [_enrich_item_from_catalog(db, cid, x) for x in _sales_items_from_payload(payload or {})]
    if not items:
        return 0
    db.query(SalesInvoiceItem).filter(SalesInvoiceItem.sales_invoice_id == inv.id).delete(synchronize_session=False)
    for item in items:
        db.add(SalesInvoiceItem(
            sales_invoice_id=inv.id,
            erpflex_item_id=text(item.get("item_id")) or None,
            product_id=text(item.get("product_id")) or None,
            product_code=text(item.get("code")) or None,
            ean=text(item.get("ean")) or None,
            description=text(item.get("description")) or "",
            ncm=text(item.get("ncm")) or None,
            cfop=text(item.get("cfop")) or None,
            unit=text(item.get("unit")) or None,
            qty=fnum(item.get("qty")), unit_price=fnum(item.get("unit_price")), total=fnum(item.get("total")),
            icms=fnum(item.get("icms")), ipi=fnum(item.get("ipi")),
        ))
    return len(items)


def _sales_invoice_detail_payload(db, cid: int, inv: SalesInvoice) -> dict:
    if not inv.erpflex_id:
        return {}
    raw = db.scalar(select(RawRecord).where(
        RawRecord.company_id == cid, RawRecord.source == "ERPFLEX",
        RawRecord.module == "faturamento_detalhe", RawRecord.external_id == str(inv.erpflex_id)
    ).order_by(RawRecord.id.desc()).limit(1))
    if not raw:
        return {}
    try:
        return json.loads(raw.payload_json or "{}")
    except Exception:
        return {}


def _save_raw_detail(db, cid: int, module: str, external_id: str, payload: dict) -> RawRecord:
    h = payload_hash(payload)
    raw = db.scalar(select(RawRecord).where(
        RawRecord.company_id == cid, RawRecord.source == "ERPFLEX", RawRecord.module == module,
        RawRecord.external_id == str(external_id)
    ).limit(1))
    if not raw:
        raw = RawRecord(company_id=cid, source="ERPFLEX", module=module, external_id=str(external_id), record_hash=h, payload_json=safe_json(payload))
        db.add(raw); db.flush()
    else:
        raw.record_hash = h; raw.payload_json = safe_json(payload); raw.updated_at = datetime.utcnow()
    return raw



def _specific_nf_number(payload: dict | None) -> str:
    """Número fiscal real da NF-e. Nunca usa `documento`, que pode ser pedido/título."""
    if not payload:
        return ""
    value = deep_find(payload, (
        "nr_nfe", "nfe", "numero_nfe", "numero_nf", "numero_nota", "nota_fiscal", "SF2_NrNfe", "nf"
    ))
    raw = text(value).strip()
    if not raw:
        return ""
    # O campo fiscal da conta é numérico. Preserva zeros à esquerda, mas rejeita 0.
    digits = "".join(ch for ch in raw if ch.isdigit())
    if not digits or int(digits or "0") <= 0:
        return ""
    # Se o valor original continha separadores/concatenações, o campo não é confiável como NF.
    if any(ch not in "0123456789" for ch in raw):
        return ""
    return raw


def _invoice_order_number(payload: dict | None) -> str:
    if not payload:
        return ""
    return text(deep_find(payload, ("documento", "numero_pedido", "pedido", "orcamento", "pedido_numero"))).strip()


def _invoice_status(payload: dict | None) -> dict:
    payload = payload or {}
    raw = text(deep_find(payload, ("status_nf", "status_nfe", "status_danfe", "situacao_nf", "situacao_nfe", "situacao", "status"))).strip()
    cancel_flag = deep_find(payload, ("cancelada", "cancelado", "nf_cancelada", "nfe_cancelada", "cancelamento"))
    cancel_text = text(cancel_flag).strip().lower()
    cancelled = ("cancel" in raw.lower()) or cancel_text in {"1", "s", "sim", "true", "cancelada", "cancelado"}
    state = "CANCELADA" if cancelled else "ATIVA"
    return {"state": state, "raw": raw or state, "cancelled": cancelled}


def _repair_sales_invoice_core(db, cid: int) -> int:
    """Repara versões antigas que usaram `documento` como NF e valor genérico como total."""
    changed = 0
    rows = db.execute(select(SalesInvoice, RawRecord).join(RawRecord, SalesInvoice.raw_record_id == RawRecord.id).where(
        SalesInvoice.company_id == cid, RawRecord.source == "ERPFLEX", RawRecord.module == "faturamento"
    )).all()
    for inv, raw in rows:
        try:
            payload = json.loads(raw.payload_json or "{}")
        except Exception:
            continue
        nf = _specific_nf_number(payload)
        if nf and inv.number != nf:
            inv.number = nf; changed += 1
        elif not nf and inv.number:
            # Para registros ERPFlex sem nr_nfe real, não mantenha o antigo documento/pedido como NF.
            inv.number = None; changed += 1
        total = fnum(deep_find(payload, ("valor_nf", "valor_total_da_nota", "valor_total", "total_nf")))
        if total and abs(float(inv.total or 0) - total) > 0.005:
            inv.total = total; changed += 1
    if changed:
        db.flush()
    return changed


def _title_payload(db, title: FinancialTitle) -> dict:
    base = _load_raw_payload(db, title.raw_record_id) if title.raw_record_id else {}
    detail = _receivable_detail_payload(db, title) if title.kind == "RECEBER" else {}
    merged = dict(base or {})
    if detail:
        merged.update(detail)
    return merged


def _title_nf_number(payload: dict | None) -> str:
    # Usa a mesma extração fiscal restritiva do faturamento. Nunca usa `documento`,
    # pois nesta conta ele representa o pedido/referência comercial.
    return _specific_nf_number(payload)


def _title_faturamento_id(payload: dict | None) -> str:
    return text(deep_find(payload or {}, ("faturamento_id", "id_nota_saida", "id_faturamento"))).strip()


def _title_number(payload: dict | None, fallback: str = "") -> str:
    """Identificador do título. Não usa documento, pois nesta conta documento representa o pedido."""
    value = deep_find(payload or {}, (
        "numero_titulo", "titulo_numero", "titulo", "nr_titulo", "num_titulo",
        "id_titulo", "titulo_id", "id_registro_titulos_receber"
    ))
    return text(value).strip() or text(fallback).strip()


def _title_order_number(payload: dict | None, stored_document: str = "") -> str:
    """Pedido/documento comercial separado do número fiscal e do número do título."""
    value = deep_find(payload or {}, ("documento", "numero_pedido", "pedido", "pedido_numero"))
    return text(value).strip() or text(stored_document).strip()


def _wallets_from_boleto_options(payload) -> dict[tuple[str, str], str]:
    """Extrai nomes de carteiras do endpoint comprovado /api_v2/boleto/options."""
    out: dict[tuple[str, str], str] = {}
    def local(node: dict, names: tuple[str, ...]):
        lowered = {str(k).lower(): v for k, v in node.items()}
        for name in names:
            v = lowered.get(name.lower())
            if v not in (None, "", [], {}):
                return v
        return None
    def walk(node, bank_ctx: str = "", wallet_ctx: bool = False):
        if isinstance(node, dict):
            bank_here = text(local(node, ("id_banco","banco_id","idBanco","bank_id"))).strip() or bank_ctx
            wid = text(local(node, ("id_carteira","carteira_id","idCarteira","wallet_id","value"))).strip()
            label = text(local(node, ("carteira_desc","nome_carteira","descricao_carteira","carteira","label","text","nome","descricao"))).strip()
            if wallet_ctx and wid and label and label != wid:
                out[(bank_here, wid)] = label
                out.setdefault(("", wid), label)
            for key, child in node.items():
                lk = str(key).lower()
                walk(child, bank_here, wallet_ctx or "carteir" in lk or "wallet" in lk)
        elif isinstance(node, list):
            for child in node:
                walk(child, bank_ctx, wallet_ctx)
    walk(payload)
    return out


def _refresh_boleto_options(db, client: ERPFlexClient, cid: int) -> int:
    try:
        payload = client.boleto_options()
    except Exception:
        return 0
    if payload in (None, "", [], {}):
        return 0
    _save_raw_detail(db, cid, "boleto_options", "default", payload)
    return len(_wallets_from_boleto_options(payload))


def _receivable_boleto_ids(db, title: FinancialTitle) -> dict:
    payload = _title_payload(db, title)
    detail = _receivable_detail_payload(db, title)
    id_registro = text(deep_find(detail, ("id_registro_titulos_receber", "idRegistro"))).strip()
    if not id_registro:
        # Fallback validado no Analytics V7.8 desta conta: id do Título V1.
        id_registro = text(deep_find(detail, ("id",))).strip()
    bank_id = text(deep_find(detail, ("id_banco","banco_id","idBanco"))).strip() or text(deep_find(payload, ("id_banco","banco_id","idBanco"))).strip()
    wallet_id = text(deep_find(detail, ("id_carteira","carteira_id","idCarteira"))).strip() or text(deep_find(payload, ("id_carteira","carteira_id","idCarteira"))).strip()
    return {
        "id_registro": id_registro, "bank_id": bank_id, "wallet_id": wallet_id,
        "numero_boleto": text(deep_find(detail, ("numero_boleto",))).strip(),
        "linha_digitavel": text(deep_find(detail, ("linha_digitavel","linhaDigital"))).strip(),
    }


def _boleto_status(db, title: FinancialTitle) -> tuple[str, str]:
    """Status operacional do boleto. Saldo tem prioridade para evitar falso `Pago`."""
    if title.partner_id:
        rule = db.scalar(select(CollectionCustomerRule).where(
            CollectionCustomerRule.company_id == title.company_id,
            CollectionCustomerRule.partner_id == title.partner_id,
            CollectionCustomerRule.has_collection_contract == True
        ).limit(1))
        if rule:
            return "CONTRATO", "Contrato de cobrança"
    value = max(0.0, float(title.value or 0))
    paid = max(0.0, float(title.paid_value or 0))
    balance = max(0.0, value - paid)
    # Data de baixa/status textual isolados não bastam: alguns retornos antigos trazem
    # esses campos mesmo com saldo em aberto. Pago somente quando o valor está liquidado.
    if value > 0 and balance <= 0.005:
        return "PAGO", "Pago"
    if _cached_boleto_html(db, title):
        return "DISPONIVEL", "Disponível"
    ids = _receivable_boleto_ids(db, title)
    if ids.get("numero_boleto") or ids.get("linha_digitavel"):
        return "DISPONIVEL", "Disponível"
    if ids.get("id_registro") and ids.get("bank_id") and ids.get("wallet_id") and ids.get("wallet_id") != "0":
        return "PRONTO", "Pronto para gerar"
    return "INDISPONIVEL", "Indisponível"


def _payment_status(title: FinancialTitle) -> tuple[str, str]:
    """Status financeiro normalizado para contas a pagar/receber sem inventar estado da API."""
    value = max(0.0, float(title.value or 0))
    paid = max(0.0, float(title.paid_value or 0))
    balance = max(0.0, value - paid)
    raw = (title.status or "").strip().lower()
    if any(x in raw for x in ("cancel", "estorn")):
        return "CANCELADO", "Cancelado"
    if value > 0 and balance <= 0.005:
        return "PAGO", "Pago" if title.kind == "PAGAR" else "Recebido"
    if paid > 0.005 and balance > 0.005:
        return "PARCIAL", "Parcial"
    today = date.today().isoformat()
    due = str(title.due_date or "")[:10]
    if due and due < today and balance > 0.005:
        return "VENCIDO", "Vencido"
    return "ABERTO", "Em aberto"


def _unique_invoice_by_order(db, cid: int, order_number: str, kind: str):
    """Fallback seguro: pedido exato + único, pesquisando primeiro apenas raws que contêm a referência."""
    target = text(order_number).strip()
    if not target:
        return None
    target_norm = target.lstrip("0") or "0"
    if kind == "RECEBER":
        q = (select(SalesInvoice, RawRecord)
             .join(RawRecord, SalesInvoice.raw_record_id == RawRecord.id)
             .where(SalesInvoice.company_id == cid, RawRecord.payload_json.ilike(f"%{target}%")))
    else:
        q = (select(PurchaseInvoice, RawRecord)
             .join(RawRecord, PurchaseInvoice.erpflex_raw_id == RawRecord.id)
             .where(PurchaseInvoice.company_id == cid, RawRecord.payload_json.ilike(f"%{target}%")))
    matches = []
    for inv, raw in db.execute(q).all():
        try:
            payload = json.loads(raw.payload_json or "{}")
        except Exception:
            payload = {}
        ref = _invoice_order_number(payload)
        if ref and (ref == target or (ref.lstrip("0") or "0") == target_norm):
            matches.append(inv)
            if len(matches) > 1:
                return None
    return matches[0] if len(matches) == 1 else None


def _sales_invoice_xml_for_title(db, cid: int, title: FinancialTitle) -> tuple[str, str]:
    """Retorna XML de saída somente quando ele já existe nos dados locais/retornos armazenados."""
    module, inv = _linked_invoice_for_title(db, cid, title)
    if module != "faturamento" or not inv:
        return "", ""
    payload = _sales_invoice_detail_payload(db, cid, inv) or {}
    value = deep_find(payload, ("xml", "xml_nfe", "nfe_xml", "conteudoXml", "conteudo_xml", "arquivoXml", "arquivo_xml"))
    xml_text = text(value).strip()
    if not xml_text.startswith("<"):
        xml_text = ""
    filename = f"NFe_{inv.number or inv.access_key or inv.id}.xml" if xml_text else ""
    return xml_text, filename


def _cached_boleto_html(db, title: FinancialTitle) -> str:
    raw = db.scalar(select(RawRecord).where(
        RawRecord.company_id == title.company_id, RawRecord.source == "ERPFLEX",
        RawRecord.module == "boleto_html", RawRecord.external_id == str(title.erpflex_id)
    ).order_by(RawRecord.id.desc()).limit(1))
    if not raw:
        return ""
    try:
        payload = json.loads(raw.payload_json or "{}")
        return str(payload.get("html") or "")
    except Exception:
        return ""


def _fetch_boleto_html(db, cid: int, title: FinancialTitle, client: ERPFlexClient, force: bool = False) -> tuple[str, dict]:
    html = "" if force else _cached_boleto_html(db, title)
    detail = _receivable_detail_payload(db, title)
    if not detail:
        detail = client.receivable_title_detail(title.erpflex_id) or {}
        if detail:
            bank_names, wallet_names = _finance_maps(db, cid)
            _save_receivable_detail(db, title, detail, bank_names, wallet_names)
    _refresh_bank_catalog(db, client, cid)
    _refresh_boleto_options(db, client, cid)
    bank_names, wallet_names = _finance_maps(db, cid)
    _save_receivable_detail(db, title, detail, bank_names, wallet_names)
    ids = _receivable_boleto_ids(db, title)
    if not html:
        if not ids["id_registro"] or not ids["bank_id"] or not ids["wallet_id"] or ids["wallet_id"] == "0":
            raise RuntimeError("O ERPFlex não forneceu idRegistro, banco e carteira necessários para gerar o boleto.")
        html = client.boleto_html(ids["id_registro"], ids["bank_id"], ids["wallet_id"])
        _save_raw_detail(db, cid, "boleto_html", str(title.erpflex_id), {"html": html, **ids})
    return html, ids


def _customer_email_for_title(db, cid: int, title: FinancialTitle, client: ERPFlexClient | None = None) -> str:
    partner = db.get(Partner, title.partner_id) if title.partner_id else None
    payload = _title_payload(db, title)
    direct = text(deep_find(payload, ("email_cobranca", "email", "email_financeiro"))).strip()
    if direct:
        return direct
    customer_id = (partner.erpflex_id if partner else "") or text(deep_find(payload, ("cliente_id","id_cliente"))).strip()
    if client and customer_id:
        try:
            detail = client.customer_detail(customer_id) or {}
            _save_raw_detail(db, cid, "customer_detail", customer_id, detail)
            return text(deep_find(detail, ("email_cobranca","email","email_financeiro"))).strip()
        except Exception:
            pass
    return ""


def _remittance_status(db, title: FinancialTitle) -> tuple[str, str]:
    """Indica remessa somente por campos efetivamente presentes nos payloads sincronizados."""
    payload = _title_payload(db, title) or {}
    value = deep_find(payload, (
        "remessa_gerada", "gerou_remessa", "arquivo_remessa_gerado", "arquivo_remessa",
        "nome_arquivo_remessa", "remessa", "enviado_remessa"
    ))
    if value in (None, "", [], {}):
        return "NAO", "Não"
    if isinstance(value, bool):
        return ("SIM", "Sim") if value else ("NAO", "Não")
    txt = text(value).strip().lower()
    if txt in {"0", "false", "nao", "não", "n", "pendente"}:
        return "NAO", "Não"
    return "SIM", "Sim"


def _collection_contract_map(db, cid: int, partner_ids: list[int]) -> dict[int, bool]:
    if not partner_ids:
        return {}
    rows = db.execute(select(CollectionCustomerRule.partner_id, CollectionCustomerRule.has_collection_contract).where(
        CollectionCustomerRule.company_id == cid, CollectionCustomerRule.partner_id.in_(partner_ids)
    )).all()
    return {int(pid): bool(flag) for pid, flag in rows}


def _collection_config(db, cid: int) -> CollectionAutomationConfig:
    cfg = db.scalar(select(CollectionAutomationConfig).where(CollectionAutomationConfig.company_id == cid).limit(1))
    if not cfg:
        cfg = CollectionAutomationConfig(company_id=cid, active=False, days_before_due=3, include_overdue=True, resend_interval_hours=24)
        db.add(cfg); db.flush()
    return cfg


def _run_automatic_collection_company(cid: int) -> dict:
    sent = failed = skipped = 0
    with SessionLocal() as db:
        cfg = _collection_config(db, cid)
        if not cfg.active:
            return {"sent":0,"failed":0,"skipped":0}
        smtp = get_smtp_settings(db, cid)
        if not smtp.get("active") or not smtp.get("host") or not smtp.get("from_email"):
            return {"sent":0,"failed":0,"skipped":0}
        today = date.today()
        cutoff = today + timedelta(days=max(0, int(cfg.days_before_due or 0)))
        titles = db.scalars(select(FinancialTitle).where(
            FinancialTitle.company_id == cid, FinancialTitle.kind == "RECEBER"
        ).order_by(FinancialTitle.due_date, FinancialTitle.id)).all()
        partner_ids = [t.partner_id for t in titles if t.partner_id]
        contracts = _collection_contract_map(db, cid, partner_ids)
        settings = get_erpflex_settings(db, cid)
        with ERPFlexClient(settings) as client:
            for title in titles:
                balance = max(0.0, float(title.value or 0)-float(title.paid_value or 0))
                if balance <= 0.005 or contracts.get(title.partner_id, False):
                    skipped += 1; continue
                due = _date_from_text(title.due_date)
                if not due or due > cutoff or (due < today and not cfg.include_overdue):
                    skipped += 1; continue
                last = db.scalar(select(CollectionSendLog).where(
                    CollectionSendLog.company_id == cid, CollectionSendLog.financial_title_id == title.id,
                    CollectionSendLog.status == "ENVIADO"
                ).order_by(CollectionSendLog.sent_at.desc()).limit(1))
                if last and last.sent_at and last.sent_at > datetime.utcnow()-timedelta(hours=max(1,int(cfg.resend_interval_hours or 24))):
                    skipped += 1; continue
                try:
                    recipient = _customer_email_for_title(db, cid, title, client)
                    if not recipient or "@" not in recipient:
                        raise RuntimeError("Cliente sem e-mail de cobrança disponível.")
                    html, _ids = _fetch_boleto_html(db, cid, title, client, force=False)
                    display = _financial_display(db, title, db.get(Partner, title.partner_id) if title.partner_id else None)
                    nf = display.get("nf_number") or ""
                    xml_text, xml_filename = _sales_invoice_xml_for_title(db, cid, title)
                    subject = (cfg.subject_template or "Cobrança - título {titulo} - vencimento {vencimento}").format(
                        titulo=display.get("title_number") or title.erpflex_id, vencimento=title.due_date or "", nf=nf
                    )
                    body = (cfg.body_template or "Prezados,\n\nSegue boleto referente ao título {titulo}, vencimento {vencimento}, no valor de R$ {valor:.2f}.\n\nAtenciosamente.").format(
                        titulo=display.get("title_number") or title.erpflex_id, vencimento=title.due_date or "", valor=balance, nf=nf
                    )
                    filename = f"boleto_{title.erpflex_id}.html"
                    _smtp_send_boleto(smtp, recipient, subject, body, html, filename, xml_text, xml_filename)
                    db.add(CollectionSendLog(company_id=cid, financial_title_id=title.id, recipient=recipient, trigger_type="AUTO", status="ENVIADO", message=f"NF-e {nf}" if nf else None))
                    db.commit(); sent += 1
                except Exception as exc:
                    db.rollback()
                    db.add(CollectionSendLog(company_id=cid, financial_title_id=title.id, trigger_type="AUTO", status="ERRO", message=str(exc)[:1000]))
                    db.commit(); failed += 1
    return {"sent":sent,"failed":failed,"skipped":skipped}


_collection_worker_started = False
def _collection_worker():
    while True:
        try:
            with SessionLocal() as db:
                companies = db.scalars(select(Company.id).where(Company.active == True)).all()
            for cid in companies:
                _run_automatic_collection_company(int(cid))
        except Exception:
            pass
        time.sleep(3600)


def _invoice_payment_map(db, cid: int, invoices: list[SalesInvoice]) -> dict[int, dict]:
    result = {inv.id: {"titles": [], "summary": {"count":0,"total":0.0,"paid":0.0,"balance":0.0,"next_due":"","banks":[],"wallets":[]}} for inv in invoices}
    if not invoices:
        return result
    by_fid = {str(inv.erpflex_id): inv.id for inv in invoices if inv.erpflex_id}
    by_nf = {(str(inv.number).lstrip("0") or "0"): inv.id for inv in invoices if inv.number}
    order_refs: dict[str, list[int]] = {}
    for inv in invoices:
        ref = _invoice_order_number(_load_raw_payload(db, inv.raw_record_id))
        if ref:
            order_refs.setdefault(ref.lstrip("0") or "0", []).append(inv.id)
    titles = db.scalars(select(FinancialTitle).where(FinancialTitle.company_id == cid, FinancialTitle.kind == "RECEBER")).all()
    bank_names, wallet_names = _finance_maps(db, cid)
    for title in titles:
        payload = _title_payload(db, title)
        fid = _title_faturamento_id(payload)
        nf = _title_nf_number(payload)
        inv_id = by_fid.get(fid) if fid else None
        if not inv_id and nf:
            inv_id = by_nf.get(nf.lstrip("0") or "0")
        if not inv_id:
            order = _title_order_number(payload, title.document or "")
            ids = order_refs.get(order.lstrip("0") or "0", []) if order else []
            if len(ids) == 1:
                inv_id = ids[0]
        if not inv_id:
            continue
        partner = db.get(Partner, title.partner_id) if title.partner_id else None
        disp = _financial_display(db, title, partner, bank_names, wallet_names)
        balance = max(0.0, float(title.value or 0) - float(title.paid_value or 0))
        entry = {"title": title, "display": disp, "payload": payload, "nf_number": nf, "balance": balance}
        result[inv_id]["titles"].append(entry)
    for inv_id, block in result.items():
        titles = block["titles"]
        due = sorted([x["title"].due_date for x in titles if x["title"].due_date])
        banks=[]; wallets=[]
        for x in titles:
            b=x["display"].get("bank"); w=x["display"].get("wallet")
            if b and b not in banks: banks.append(b)
            if w and w not in wallets: wallets.append(w)
        total_value = sum(float(x["title"].value or 0) for x in titles)
        paid_value = sum(float(x["title"].paid_value or 0) for x in titles)
        balance_value = sum(float(x["balance"] or 0) for x in titles)
        if not titles:
            pay_status = "SEM TÍTULO"
        elif balance_value <= 0.005:
            pay_status = "PAGO"
        elif paid_value > 0.005:
            pay_status = "PARCIAL"
        else:
            pay_status = "EM ABERTO"
        block["summary"] = {
            "count": len(titles), "total": total_value,
            "paid": paid_value, "balance": balance_value, "status": pay_status,
            "next_due": due[0] if due else "", "banks": banks, "wallets": wallets,
            "code": "PAGO" if pay_status=="PAGO" else "PARCIAL" if pay_status=="PARCIAL" else "ABERTO" if pay_status=="EM ABERTO" else "SEM_TITULO",
        }
    return result


def _purchase_payment_map(db, cid: int, invoices: list[PurchaseInvoice]) -> dict[int, dict]:
    result = {inv.id: {"titles": [], "summary": {"count":0,"total":0.0,"paid":0.0,"balance":0.0,"status":"SEM TÍTULO"}} for inv in invoices}
    if not invoices:
        return result
    wanted = {inv.id for inv in invoices}
    titles = db.scalars(select(FinancialTitle).where(FinancialTitle.company_id == cid, FinancialTitle.kind == "PAGAR")).all()
    for title in titles:
        module, inv = _linked_invoice_for_title(db, cid, title)
        if module != "compras" or not inv or inv.id not in wanted:
            continue
        result[inv.id]["titles"].append(title)
    for inv_id, block in result.items():
        titles = block["titles"]
        total = sum(float(t.value or 0) for t in titles)
        paid = sum(float(t.paid_value or 0) for t in titles)
        balance = sum(max(0.0, float(t.value or 0)-float(t.paid_value or 0)) for t in titles)
        if not titles: code,label = "SEM_TITULO", "Sem título"
        elif balance <= 0.005: code,label = "PAGO", "Pago"
        elif paid > 0.005: code,label = "PARCIAL", "Parcial"
        elif any(_payment_status(t)[0] == "VENCIDO" for t in titles): code,label = "VENCIDO", "Vencido"
        else: code,label = "ABERTO", "Em aberto"
        block["summary"] = {"count":len(titles),"total":total,"paid":paid,"balance":balance,"status":label,"code":code}
    return result

def _product_summary(items, limit: int = 2):
    names = []
    for it in items or []:
        name = text((it.get("description") if isinstance(it, dict) else getattr(it, "description", "")) or
                    (it.get("produto") if isinstance(it, dict) else ""))
        if name and name not in names:
            names.append(name)
    shown = names[:limit]
    summary = " · ".join(shown)
    if len(names) > limit:
        summary += f" · +{len(names)-limit}"
    return {"count": len(items or []), "summary": summary or "Itens não disponíveis"}


def _invoice_order_refs(payload: dict) -> list[str]:
    refs: list[str] = []
    for name in ("orcamento_id", "id_orcamento", "pedido_id", "id_pedido", "orcamento", "solicitacao_id", "id_solicitacao"):
        value = text(deep_find(payload or {}, (name,)))
        if value and value not in refs:
            refs.append(value)
    return refs


def _order_payload_candidates(db, cid: int, refs: list[str]):
    """Retorna detalhes/resumos locais de pedidos sem inventar vínculo por cliente/data."""
    seen = set()
    for ref in refs:
        # Primeiro, detalhes individuais já consultados pela API documentada.
        raws = db.scalars(select(RawRecord).where(
            RawRecord.company_id == cid, RawRecord.source == "ERPFLEX",
            RawRecord.module == "order_detail", RawRecord.external_id == str(ref)
        ).order_by(RawRecord.id.desc()).limit(3)).all()
        for raw in raws:
            if raw.id in seen: continue
            seen.add(raw.id)
            try: yield json.loads(raw.payload_json or "{}")
            except Exception: pass

        # Depois o pedido sincronizado na base central.
        order = db.scalar(select(SalesOrder).where(
            SalesOrder.company_id == cid, or_(SalesOrder.erpflex_id == str(ref), SalesOrder.number == str(ref))
        ).order_by(SalesOrder.id.desc()).limit(1))
        if order and order.raw_record_id and order.raw_record_id not in seen:
            seen.add(order.raw_record_id)
            yield _load_raw_payload(db, order.raw_record_id)

        # Compatibilidade com registros antigos cujo external_id foi salvo com outro identificador.
        raws = db.scalars(select(RawRecord).where(
            RawRecord.company_id == cid, RawRecord.source == "ERPFLEX", RawRecord.module == "orders",
            RawRecord.payload_json.ilike(f"%{ref}%")
        ).order_by(RawRecord.id.desc()).limit(20)).all()
        for raw in raws:
            if raw.id in seen: continue
            seen.add(raw.id)
            try:
                rp = json.loads(raw.payload_json or "{}")
            except Exception:
                continue
            ids = {text(deep_find(rp, (name,))) for name in ("id", "solicitacao_id", "pedido_id", "id_solicitacao", "numero_pedido", "pedido", "codigo")}
            ids.discard("")
            if str(ref) in ids:
                yield rp


def _refresh_bank_catalog(db, client: ERPFlexClient, cid: int) -> int:
    """Atualiza /api/bancos/ (endpoint comprovado no V7.8) e preserva os campos SA6_* reais."""
    try:
        rows = client.banks()
    except Exception:
        return 0
    saved = 0
    for rec in rows or []:
        if not isinstance(rec, dict):
            continue
        bid = text(deep_find(rec, ("SA6_ID", "id", "banco_id", "id_banco", "SA6_Codigo", "codigo")))
        if not bid:
            continue
        _save_raw_detail(db, cid, "banks", bid, rec)
        saved += 1
    return saved


def _hydrate_sales_invoice_items(db, cid: int, inv: SalesInvoice, client: ERPFlexClient | None = None) -> tuple[int, str]:
    """Garante itens da NF usando somente fontes comprovadas.

    Ordem: itens já persistidos -> itens no faturamento -> pedido local ->
    GET /api/venda/solicitacao/{id} (endpoint oficial documentado).
    Nunca consulta /api_v2/faturamento/itens/{id}, que não está validado.
    """
    stored = db.scalars(select(SalesInvoiceItem).where(SalesInvoiceItem.sales_invoice_id == inv.id)).all()
    if stored:
        return len(stored), "faturamento_api"
    payload = _load_raw_payload(db, inv.raw_record_id)
    count = _persist_sales_invoice_items(db, cid, inv, payload)
    if count:
        return count, "faturamento"
    refs = _invoice_order_refs(payload)
    for rp in _order_payload_candidates(db, cid, refs):
        items = _sales_items_from_payload(rp)
        if items:
            count = _persist_sales_invoice_items(db, cid, inv, {"itens": items})
            if count:
                return count, "pedido"
    if client:
        for ref in refs:
            try:
                detail = client.order_detail(ref)
            except Exception:
                detail = None
            if not detail:
                continue
            _save_raw_detail(db, cid, "order_detail", str(ref), detail)
            items = _sales_items_from_payload(detail)
            if items:
                count = _persist_sales_invoice_items(db, cid, inv, {"itens": items})
                if count:
                    return count, "pedido_api"
    return 0, "sem_itens"


def _sales_items_with_fallback(db, cid: int, inv: SalesInvoice, payload: dict):
    stored = db.scalars(select(SalesInvoiceItem).where(SalesInvoiceItem.sales_invoice_id == inv.id).order_by(SalesInvoiceItem.id)).all()
    if stored:
        return stored, "faturamento_api"
    items = _sales_items_from_payload(payload or {})
    if items:
        return [_enrich_item_from_catalog(db, cid, x) for x in items], "faturamento"
    for rp in _order_payload_candidates(db, cid, _invoice_order_refs(payload or {})):
        order_items = _sales_items_from_payload(rp)
        if order_items:
            return [_enrich_item_from_catalog(db, cid, x) for x in order_items], "pedido"
    return [], "sem_itens"


def _purchase_items_display(db, inv: PurchaseInvoice):
    rows = db.scalars(select(PurchaseItem).where(PurchaseItem.purchase_invoice_id == inv.id).order_by(PurchaseItem.id)).all()
    if rows:
        # Repara descrições antigas vazias cruzando com o cadastro central.
        changed = False
        for r in rows:
            if not r.description:
                prod = _catalog_product(db, inv.company_id, text(r.product_code), text(r.product_code), text(r.ean))
                if prod and prod.description:
                    r.description = prod.description; changed = True
                    r.ean = r.ean or prod.ean; r.ncm = r.ncm or prod.ncm; r.unit = r.unit or prod.unit
        if changed: db.flush()
        return rows, _product_summary(rows)
    payload = _load_raw_payload(db, inv.erpflex_raw_id)
    parsed = [_enrich_item_from_catalog(db, inv.company_id, x) for x in _sales_items_from_payload(payload)] if payload else []
    return parsed, _product_summary(parsed)


def _persist_purchase_items(db, inv: PurchaseInvoice, payload: dict) -> int:
    raw_items = _purchase_items_from_record(payload or {})
    if not raw_items: return 0
    db.query(PurchaseItem).filter(PurchaseItem.purchase_invoice_id == inv.id).delete(synchronize_session=False)
    for raw_item in raw_items:
        item = {
            "product_id": text(_purchase_item_value(raw_item, ("produto_id", "id_produto", "product_id"))),
            "code": text(_purchase_item_value(raw_item, ("codigo_produto", "produto_codigo", "codigo", "sku", "referencia"))),
            "ean": text(_purchase_item_value(raw_item, ("EAN", "ean", "gtin", "codigo_barras", "cEAN"))),
            "description": text(_purchase_item_value(raw_item, ("desc_produto", "descricao_produto", "descricao_item", "produto_desc", "produto", "descricao", "nome", "natureza_descricao"))),
            "ncm": text(_purchase_item_value(raw_item, ("ncm", "NCM"))), "cfop": text(_purchase_item_value(raw_item, ("cfop", "CFOP"))),
            "unit": text(_purchase_item_value(raw_item, ("unidade", "unit", "uCom"))),
        }
        item = _enrich_item_from_catalog(db, inv.company_id, item)
        product_code = item["code"] or item["product_id"] or None
        db.add(PurchaseItem(
            purchase_invoice_id=inv.id, product_code=product_code, ean=item["ean"] or None, description=item["description"] or "",
            ncm=item["ncm"] or None, cfop=item["cfop"] or None, unit=item["unit"] or None,
            qty=fnum(_purchase_item_value(raw_item, ("quantidade", "qtd", "qCom"))),
            unit_price=fnum(_purchase_item_value(raw_item, ("preco_unitario", "valor_unitario", "vUnCom", "unitario"))),
            total=fnum(_purchase_item_value(raw_item, ("preco_total_item", "valor_item", "valor_total", "total", "vProd"))),
            icms=fnum(_purchase_item_value(raw_item, ("valor_icms", "icms", "vICMS"))),
            icms_st=fnum(_purchase_item_value(raw_item, ("valor_icmsst", "icms_st", "vICMSST"))),
            ipi=fnum(_purchase_item_value(raw_item, ("valor_ipi", "ipi", "vIPI"))),
            pis=fnum(_purchase_item_value(raw_item, ("valor_pis", "pis", "vPIS"))),
            cofins=fnum(_purchase_item_value(raw_item, ("valor_cofins", "cofins", "vCOFINS"))),
        ))
    return len(raw_items)


def _wallets_from_bank_payload(payload: dict, bank_id: str) -> dict[tuple[str,str], str]:
    result = {}
    def walk(node, carteira_context=False):
        if isinstance(node, dict):
            for key, child in node.items():
                ctx = carteira_context or ("carteir" in str(key).lower())
                if ctx and isinstance(child, (dict,list)):
                    walk(child, True)
            if carteira_context:
                wid = text(deep_find(node, ("id_carteira", "carteira_id", "idCarteira", "SA6_Carteira", "id")))
                label = text(deep_find(node, ("nome_carteira", "carteira_desc", "nome", "descricao", "codigo", "carteira")))
                if wid and label and label != wid:
                    result[(bank_id, wid)] = label
        elif isinstance(node, list):
            for child in node: walk(child, carteira_context)
    walk(payload, False)
    return result


def _finance_maps(db, cid: int):
    bank_names: dict[str,str] = {}; wallet_names: dict[tuple[str,str],str] = {}
    raws = db.scalars(select(RawRecord).where(
        RawRecord.company_id == cid, RawRecord.source == "ERPFLEX", RawRecord.module.in_(["banks", "bank_detail", "boleto_options"])
    )).all()
    for raw in raws:
        try: payload = json.loads(raw.payload_json or "{}")
        except Exception: continue
        if raw.module == "boleto_options":
            wallet_names.update(_wallets_from_boleto_options(payload))
            continue
        # /api/bancos/ da conta real retorna nomenclatura legada SA6_* (ex.: SA6_ID=25325, SA6_Desc=ITAU VIN).
        bid = text(deep_find(payload, ("SA6_ID", "id", "banco_id", "id_banco"))) or (str(raw.external_id) if raw.module=="bank_detail" else "")
        code = text(deep_find(payload, ("SA6_Codigo", "codigo", "codigo_banco")))
        name = text(deep_find(payload, ("SA6_Desc", "nome", "descricao", "banco", "razao_social", "banco_desc", "nome_banco")))
        if bid and name:
            bank_names[bid] = name
        if code and name:
            bank_names.setdefault(code, name)
        if bid:
            wallet_names.update(_wallets_from_bank_payload(payload, bid))
    return bank_names, wallet_names


def _bank_map(db, cid: int):
    return _finance_maps(db, cid)[0]

def _receivable_detail_payload(db, title: FinancialTitle) -> dict:
    if title.kind != "RECEBER" or not title.erpflex_id:
        return {}
    raw = db.scalar(select(RawRecord).where(
        RawRecord.company_id == title.company_id, RawRecord.source == "ERPFLEX",
        RawRecord.module == "receber_detalhe", RawRecord.external_id == str(title.erpflex_id)
    ).order_by(RawRecord.id.desc()).limit(1))
    if not raw:
        return {}
    try:
        return json.loads(raw.payload_json or "{}")
    except Exception:
        return {}


def _generic_financial_label(value: str | None, prefixes: tuple[str, ...]) -> bool:
    v = text(value).strip().lower()
    if not v:
        return True
    return any(v.startswith(x.lower()) for x in prefixes)


def _financial_display(db, title: FinancialTitle, partner: Partner | None, bank_names: dict[str, str] | None = None,
                       wallet_names: dict[tuple[str,str],str] | None = None):
    if bank_names is None or wallet_names is None:
        bank_names, wallet_names = _finance_maps(db, title.company_id)
    payload = _load_raw_payload(db, title.raw_record_id)
    detail = _receivable_detail_payload(db, title)
    partner_name = (partner.name if partner and partner.name and partner.name != "Não identificado" else "")
    if not partner_name:
        fields = ("cliente_desc", "cliente_nome", "cliente", "razao_social", "nome_cliente") if title.kind == "RECEBER" else ("fornecedor_desc", "fornecedor_nome", "fornecedor", "razao_social", "nome_fornecedor")
        partner_name = text(deep_find(detail, fields)) or text(deep_find(payload, fields)) or "—"
    issue = title.issue_date or normalize_date(deep_find(detail, ("emissao", "data_emissao", "data_inclusao", "dt_inclusao", "data"))) or normalize_date(deep_find(payload, ("emissao", "data_emissao", "data_inclusao", "dt_inclusao", "data"))) or "—"

    bank_id = text(deep_find(detail, ("id_banco", "banco_id", "idBanco"))) or text(deep_find(payload, ("banco_id", "id_banco", "idBanco")))
    bank_desc = text(deep_find(detail, ("banco_desc", "nome_banco", "banco"))) or text(deep_find(payload, ("banco_desc", "nome_banco", "banco")))
    stored_bank = "" if _generic_financial_label(title.bank, ("Sem banco", "Banco ", "Banco ID ")) else text(title.bank)
    bank = bank_desc or stored_bank or bank_names.get(bank_id, "")
    if not bank:
        bank = "Sem banco informado pela API" if bank_id in {"", "0"} else f"Banco {bank_id}"

    wallet_id = text(deep_find(detail, ("id_carteira", "carteira_id", "idCarteira"))) or text(deep_find(payload, ("id_carteira", "carteira_id", "idCarteira")))
    wallet_desc = text(deep_find(detail, ("carteira_desc", "nome_carteira", "carteira"))) or text(deep_find(payload, ("carteira_desc", "nome_carteira", "carteira")))
    stored_wallet = "" if _generic_financial_label(title.wallet, ("Sem carteira", "Carteira ", "Carteira não ")) else text(title.wallet)
    wallet = wallet_desc or stored_wallet or wallet_names.get((bank_id, wallet_id), "") or wallet_names.get(("", wallet_id), "")
    if not wallet:
        wallet = "Carteira não informada pela API" if wallet_id in {"", "0"} else f"Carteira {wallet_id}"
    merged_fin = {**(payload or {}), **(detail or {})}
    nf_number = _title_nf_number(merged_fin)
    faturamento_id = _title_faturamento_id(merged_fin)
    title_number = _title_number(merged_fin, title.erpflex_id)
    order_number = _title_order_number(merged_fin, title.document or "")
    return {"partner": partner_name, "issue_date": issue, "bank": bank, "bank_id": bank_id, "wallet": wallet, "wallet_id": wallet_id,
            "nf_number": nf_number, "faturamento_id": faturamento_id, "title_number": title_number, "order_number": order_number,
            "payload": payload, "detail_payload": detail}


def _save_receivable_detail(db, title: FinancialTitle, detail: dict, bank_names: dict[str, str] | None = None,
                            wallet_names: dict[tuple[str,str],str] | None = None) -> None:
    if not detail:
        return
    ext = str(title.erpflex_id)
    _save_raw_detail(db, title.company_id, "receber_detalhe", ext, detail)
    if bank_names is None or wallet_names is None:
        bank_names, wallet_names = _finance_maps(db, title.company_id)
    bank_id = text(deep_find(detail, ("id_banco", "banco_id", "idBanco")))
    bank_desc = text(deep_find(detail, ("banco_desc", "nome_banco", "banco")))
    wallet_id = text(deep_find(detail, ("id_carteira", "carteira_id", "idCarteira")))
    wallet_desc = text(deep_find(detail, ("carteira_desc", "nome_carteira", "carteira")))
    if bank_desc:
        title.bank = bank_desc
    elif bank_id:
        title.bank = bank_names.get(bank_id) or f"Banco {bank_id}"
    if wallet_desc:
        title.wallet = wallet_desc
    elif wallet_id:
        title.wallet = wallet_names.get((bank_id, wallet_id)) or wallet_names.get(("", wallet_id)) or f"Carteira {wallet_id}"
    issue = normalize_date(deep_find(detail, ("emissao", "data_emissao", "data_inclusao", "dt_inclusao", "data")))
    if issue and not title.issue_date:
        title.issue_date = issue


def _ensure_bank_detail(db, client: ERPFlexClient, cid: int, bank_id: str) -> bool:
    bid = text(bank_id)
    if not bid or bid == "0":
        return False
    existing = db.scalar(select(RawRecord).where(
        RawRecord.company_id == cid, RawRecord.source == "ERPFLEX", RawRecord.module == "bank_detail",
        RawRecord.external_id == bid
    ).limit(1))
    if existing:
        return True
    try:
        detail = client.bank_detail(bid)
    except Exception:
        return False
    if not detail:
        return False
    _save_raw_detail(db, cid, "bank_detail", bid, detail)
    return True

def _sales_note_context(db, cid: int, invoice_id: int):
    row = db.execute(select(SalesInvoice, Partner).outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
                     .where(SalesInvoice.company_id == cid, SalesInvoice.id == invoice_id)).first()
    if not row:
        return None
    inv, customer = row
    payload = _load_raw_payload(db, inv.raw_record_id)
    detail_payload = _sales_invoice_detail_payload(db, cid, inv)
    if detail_payload:
        merged_payload = dict(payload or {})
        merged_payload.update(detail_payload)
        payload = merged_payload
    def val(*names, default=""):
        v = deep_find(payload, tuple(names)) if payload else None
        return v if v not in (None, "") else default
    header = {
        "number": inv.number or _specific_nf_number(payload),
        "series": inv.series or text(val("serie_da_nf", "serie", "serie_nf")),
        "model": text(val("modelo_da_nf", "modelo", "modelo_nf")),
        "key": inv.access_key or access_key_from(payload) or "",
        "issue_date": inv.issue_date or text(val("data_emissao", "emissao", "dt_emissao")),
        "nature": text(val("natureza_da_operacao", "natureza_operacao", "natureza")),
        "cfop": text(val("cfop", "cod_cfop")),
        "status": _invoice_status(payload).get("raw") or text(val("status_danfe", "status", "situacao")),
        "customer_name": (customer.name if customer else "") or text(val("cliente", "razao_social", "destinatario")),
        "customer_trade": (customer.trade_name if customer else "") or text(val("nome_fantasia", "fantasia")),
        "customer_cnpj": (customer.cnpj_cpf if customer else "") or text(val("cnpj", "cpf_cnpj", "cnpj_cpf")),
        "products_total": fnum(val("valor_total_produtos", "valor_produtos", "total_produtos", "vprod")),
        "freight": fnum(val("valor_frete", "frete", "vfrete")),
        "discount": fnum(val("valor_desconto", "desconto", "vdesc")),
        "base_icms": fnum(val("base_icms", "valor_base_icms", "vbc")),
        "icms": fnum(val("valor_icms", "icms", "vicms")),
        "icms_st": fnum(val("valor_icmsst", "valor_icms_st", "icms_st", "vst")),
        "ipi": fnum(val("valor_ipi", "ipi", "vipi")),
        "pis": fnum(val("valor_pis", "pis", "vpis")),
        "cofins": fnum(val("valor_cofins", "cofins", "vcofins")),
        "total": inv.total or fnum(val("valor_nf", "valor_total_da_nota", "valor_total", "total", "valor", "vnf")),
        "carrier": inv.carrier or text(val("transportadora", "transportador", "nome_transportadora")),
        "volumes": inv.volumes or fnum(val("quantidade_volume", "qtd_volume", "volumes", "qvol")),
        "delivery_address": inv.delivery_address or text(val("endereco_entrega", "logradouro_entrega", "endereco")),
        "delivery_city": inv.delivery_city or text(val("municipio_entrega", "cidade_entrega", "municipio")),
        "delivery_state": inv.delivery_state or text(val("uf_entrega", "uf")),
        "customer_address": (customer.address if customer else "") or text(val("destinatario_endereco", "endereco_destinatario", "endereco_entrega")),
        "customer_district": (customer.district if customer else "") or text(val("destinatario_bairro", "bairro_destinatario", "bairro_entrega")),
        "customer_city": (customer.city if customer else "") or text(val("destinatario_municipio", "municipio_destinatario", "municipio_entrega")),
        "customer_state": (customer.state if customer else "") or text(val("destinatario_uf", "uf_destinatario", "uf_entrega")),
        "customer_zip": (customer.zip_code if customer else "") or text(val("destinatario_cep", "cep_destinatario", "cep_entrega")),
        "customer_ie": text(val("ie_destinatario", "inscricao_estadual_destinatario")),
        "customer_phone": text(val("telefone_destinatario", "fone_destinatario")),
        "issuer_ie": text(val("ie_emitente", "inscricao_estadual_emitente")),
        "issuer_address": text(val("endereco_emitente", "logradouro_emitente")),
        "issuer_district": text(val("bairro_emitente")),
        "issuer_city": text(val("municipio_emitente", "cidade_emitente")),
        "issuer_state": text(val("uf_emitente")),
        "issuer_zip": text(val("cep_emitente")),
        "issuer_phone": text(val("telefone_emitente", "fone_emitente")),
        "protocol": text(val("protocolo_autorizacao", "protocolo", "numero_protocolo")),
        "authorization_date": text(val("data_autorizacao", "dh_recibo", "data_protocolo")),
        "entry_exit_date": text(val("data_saida", "data_entrada_saida", "saida")),
        "entry_exit_time": text(val("hora_saida", "hora_entrada_saida")),
        "freight_mode": text(val("modalidade_frete", "mod_frete")),
        "gross_weight": fnum(val("peso_bruto", "pesoBruto")),
        "net_weight": fnum(val("peso_liquido", "pesoLiquido")),
        "volume_brand": text(val("marca_volume", "marca")),
        "volume_number": text(val("numero_volume", "numeracao_volume")),
        "additional_info": text(val("informacoes_complementares", "informacoes_adicionais", "observacao", "observacoes")),
    }
    items, items_source = _sales_items_with_fallback(db, cid, inv, payload)
    header["items_source"] = items_source
    return inv, customer, items, payload, header


@app.get("/faturamento", response_class=HTMLResponse)
def faturamento(request: Request, q: str = "", period: str = "all", date_from: str = "", date_to: str = "",
                nf_status: str = "all", page: int = 1, per_page: int = 0, cols: str = "", save_view: int = 0,
                items_updated: int = 0, items_failed: int = 0):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        # Repara automaticamente registros de versões antigas onde `documento`/pedido foi usado como NF.
        if _repair_sales_invoice_core(db, cid):
            db.commit()
        column_defs, selected_cols, pp = _module_view_settings(db, request, cid, "faturamento", cols, per_page, save_view)
        period, start, end = _period_bounds(period, date_from, date_to)
        valid_nf = and_(
            SalesInvoice.number.is_not(None), SalesInvoice.number != "",
            func.ltrim(func.trim(SalesInvoice.number), "0") != ""
        )
        base_where = (SalesInvoice.company_id == cid, valid_nf)
        stmt = select(SalesInvoice, Partner).outerjoin(Partner, SalesInvoice.customer_id == Partner.id).outerjoin(RawRecord, SalesInvoice.raw_record_id == RawRecord.id).where(*base_where)
        count_stmt = select(func.count(SalesInvoice.id)).outerjoin(Partner, SalesInvoice.customer_id == Partner.id).outerjoin(RawRecord, SalesInvoice.raw_record_id == RawRecord.id).where(*base_where)
        if q.strip():
            like = f"%{q.strip()}%"
            cond = or_(SalesInvoice.number.ilike(like), SalesInvoice.access_key.ilike(like), Partner.name.ilike(like), RawRecord.payload_json.ilike(like))
            stmt = stmt.where(cond); count_stmt = count_stmt.where(cond)
        stmt = _apply_date_range(stmt, SalesInvoice.issue_date, start, end)
        count_stmt = _apply_date_range(count_stmt, SalesInvoice.issue_date, start, end)
        nf_status = nf_status if nf_status in {"all", "active", "cancelled"} else "all"
        ordered_stmt = stmt.order_by(SalesInvoice.issue_date.desc(), SalesInvoice.id.desc())
        if nf_status == "all":
            total_rows = db.scalar(count_stmt) or 0
            pages = max(1, (total_rows + pp - 1)//pp); page = min(max(1, page), pages)
            rows = db.execute(ordered_stmt.offset((page-1)*pp).limit(pp)).all()
        else:
            # Status fiscal é interpretado somente pelos campos reais presentes no payload.
            # Filtramos após leitura para não inventar uma regra SQL baseada em texto parcial.
            status_rows = db.execute(ordered_stmt).all()
            want_cancelled = nf_status == "cancelled"
            status_rows = [r for r in status_rows if _invoice_status(_load_raw_payload(db, r[0].raw_record_id)).get("cancelled") == want_cancelled]
            total_rows = len(status_rows)
            pages = max(1, (total_rows + pp - 1)//pp); page = min(max(1, page), pages)
            rows = status_rows[(page-1)*pp:page*pp]
        invoices = [inv for inv, _p in rows]
        payments = _invoice_payment_map(db, cid, invoices)
        item_info = {}; invoice_display = {}
        for inv, _partner_row in rows:
            payload = _load_raw_payload(db, inv.raw_record_id)
            items, source = _sales_items_with_fallback(db, cid, inv, payload)
            info = _product_summary(items); info["source"] = source; item_info[inv.id] = info
            invoice_display[inv.id] = {
                "order_number": _invoice_order_number(payload),
                "status": _invoice_status(payload),
                "payment": payments.get(inv.id, {}).get("summary", {}),
            }
        params = {"q": q, "period": period, "date_from": date_from, "date_to": date_to, "nf_status": nf_status, "per_page": pp, "cols": ",".join(selected_cols)}
        prev_url, next_url = _page_urls("/faturamento", page, pages, params)
    return render(request, "faturamento.html", company=company, rows=rows, q=q, period=period,
                  date_from=date_from, date_to=date_to, period_label=_period_label(period,start,end),
                  column_defs=column_defs, selected_cols=selected_cols, per_page=pp, page=page, pages=pages,
                  total_rows=total_rows, prev_url=prev_url, next_url=next_url, item_info=item_info, nf_status=nf_status,
                  invoice_display=invoice_display, payments=payments, items_updated=items_updated, items_failed=items_failed)


@app.post("/faturamento/atualizar-itens")
def faturamento_atualizar_itens(request: Request, invoice_id: list[int] = Form(default=[]), q: str = Form(""),
                                period: str = Form("all"), date_from: str = Form(""), date_to: str = Form(""), nf_status: str = Form("all"),
                                per_page: int = Form(25), page: int = Form(1), cols: str = Form("")):
    # V1.3.23: rota à prova de erro. Itens são obtidos pelo pedido relacionado
    # usando o endpoint oficial /api/venda/solicitacao/{id}.
    updated = failed = 0
    try:
        with SessionLocal() as db:
            cid, _company = company_context(db, request)
            ids = [int(x) for x in (invoice_id or []) if str(x).isdigit()]
            invoices = db.scalars(select(SalesInvoice).where(
                SalesInvoice.company_id == cid, SalesInvoice.id.in_(ids or [-1])
            )).all()
            settings = get_erpflex_settings(db, cid)
            client = ERPFlexClient(settings)
            try:
                for inv in invoices:
                    try:
                        count, _source = _hydrate_sales_invoice_items(db, cid, inv, client)
                        if count:
                            updated += 1
                        else:
                            failed += 1
                    except Exception:
                        failed += 1
                db.commit()
            finally:
                client.close()
    except Exception:
        # Nunca deixar uma ação de enriquecimento derrubar a página inteira.
        failed = max(failed, len(invoice_id or []) or 1)
    params={"q":q,"period":period,"date_from":date_from,"date_to":date_to,"per_page":per_page,"page":page,"cols":cols,"items_updated":updated,"items_failed":failed}
    return RedirectResponse("/faturamento?"+urlencode(params), status_code=303)


@app.post("/faturamento/{invoice_id}/atualizar-itens")
def faturamento_atualizar_item_unico(request: Request, invoice_id: int):
    ok = 0
    try:
        with SessionLocal() as db:
            cid, _company = company_context(db, request)
            inv = db.scalar(select(SalesInvoice).where(SalesInvoice.company_id == cid, SalesInvoice.id == invoice_id).limit(1))
            if inv:
                settings = get_erpflex_settings(db, cid)
                client = ERPFlexClient(settings)
                try:
                    count, _source = _hydrate_sales_invoice_items(db, cid, inv, client)
                    ok = 1 if count else 0
                    db.commit()
                finally:
                    client.close()
    except Exception:
        ok = 0
    return RedirectResponse(f"/faturamento/{invoice_id}?items_updated={ok}", status_code=303)


@app.get("/faturamento/{invoice_id}", response_class=HTMLResponse)
def faturamento_detalhe(request: Request, invoice_id: int, items_updated: int = 0):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        ctx = _sales_note_context(db, cid, invoice_id)
        if not ctx:
            return HTMLResponse("NF-e não encontrada.", status_code=404)
        inv, customer, items, payload, header = ctx
        payment_block = _invoice_payment_map(db, cid, [inv]).get(inv.id, {"titles":[],"summary":{}})
        status_info = _invoice_status(payload)
        order_number = _invoice_order_number(payload)
    return render(request, "faturamento_nota.html", company=company, inv=inv, customer=customer, items=items,
                  payload=payload, header=header, print_mode=False, items_updated=items_updated,
                  payment_block=payment_block, status_info=status_info, order_number=order_number)


@app.get("/faturamento/{invoice_id}/imprimir", response_class=HTMLResponse)
def faturamento_imprimir(request: Request, invoice_id: int):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        ctx = _sales_note_context(db, cid, invoice_id)
        if not ctx:
            return HTMLResponse("NF-e não encontrada.", status_code=404)
        inv, customer, items, payload, header = ctx
    return render(request, "danfe_impressao.html", company=company, inv=inv, customer=customer, items=items,
                  payload=payload, header=header, document_kind="saida")


@app.get("/compras", response_class=HTMLResponse)
def compras(request: Request, status: str = "", q: str = "", period: str = "all", date_from: str = "", date_to: str = "", nf_filter: str = "valid",
            page: int = 1, per_page: int = 0, cols: str = "", save_view: int = 0, items_updated: int = 0, items_failed: int = 0):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        column_defs, selected_cols, pp = _module_view_settings(db, request, cid, "compras", cols, per_page, save_view)
        period, start, end = _period_bounds(period, date_from, date_to)
        stmt = select(PurchaseInvoice, Partner).outerjoin(Partner, PurchaseInvoice.supplier_id == Partner.id).outerjoin(RawRecord, PurchaseInvoice.erpflex_raw_id == RawRecord.id).where(PurchaseInvoice.company_id == cid)
        count_stmt = select(func.count(PurchaseInvoice.id)).outerjoin(Partner, PurchaseInvoice.supplier_id == Partner.id).outerjoin(RawRecord, PurchaseInvoice.erpflex_raw_id == RawRecord.id).where(PurchaseInvoice.company_id == cid)
        nf_filter = nf_filter if nf_filter in {"valid", "all", "zero"} else "valid"
        zero_nf = or_(PurchaseInvoice.number.is_(None), func.ltrim(func.trim(PurchaseInvoice.number), "0") == "")
        if nf_filter == "valid":
            stmt = stmt.where(~zero_nf); count_stmt = count_stmt.where(~zero_nf)
        elif nf_filter == "zero":
            stmt = stmt.where(zero_nf); count_stmt = count_stmt.where(zero_nf)
        if status:
            stmt = stmt.where(PurchaseInvoice.reconciliation_status == status); count_stmt = count_stmt.where(PurchaseInvoice.reconciliation_status == status)
        if q.strip():
            like = f"%{q.strip()}%"
            cond = or_(PurchaseInvoice.number.ilike(like), PurchaseInvoice.access_key.ilike(like), Partner.name.ilike(like), RawRecord.payload_json.ilike(like))
            stmt = stmt.where(cond); count_stmt = count_stmt.where(cond)
        stmt = _apply_date_range(stmt, PurchaseInvoice.issue_date, start, end)
        count_stmt = _apply_date_range(count_stmt, PurchaseInvoice.issue_date, start, end)
        total_rows = db.scalar(count_stmt) or 0
        pages = max(1, (total_rows + pp - 1)//pp); page = min(max(1,page), pages)
        rows = db.execute(stmt.order_by(PurchaseInvoice.issue_date.desc(), PurchaseInvoice.id.desc()).offset((page-1)*pp).limit(pp)).all()
        ids = [inv.id for inv,_ in rows]
        item_info = {}
        for inv, _supplier_row in rows:
            _items, info = _purchase_items_display(db, inv)
            item_info[inv.id] = info
        payments = _purchase_payment_map(db, cid, [inv for inv,_p in rows])
        statuses = db.execute(select(PurchaseInvoice.reconciliation_status, func.count(PurchaseInvoice.id)).where(PurchaseInvoice.company_id == cid).group_by(PurchaseInvoice.reconciliation_status)).all()
        params={"status":status,"q":q,"period":period,"date_from":date_from,"date_to":date_to,"nf_filter":nf_filter,"per_page":pp,"cols":",".join(selected_cols)}
        prev_url,next_url=_page_urls("/compras",page,pages,params)
    return render(request, "compras.html", company=company, rows=rows, statuses=statuses, status=status, q=q,
                  period=period,date_from=date_from,date_to=date_to,nf_filter=nf_filter,period_label=_period_label(period,start,end),
                  column_defs=column_defs,selected_cols=selected_cols,per_page=pp,page=page,pages=pages,total_rows=total_rows,
                  prev_url=prev_url,next_url=next_url,item_info=item_info,payments=payments,items_updated=items_updated,items_failed=items_failed)


@app.post("/compras/atualizar-itens")
def compras_atualizar_itens(request: Request, purchase_id: list[int] = Form(default=[])):
    updated = failed = 0
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        settings = get_erpflex_settings(db, cid)
        invoices = db.scalars(select(PurchaseInvoice).where(
            PurchaseInvoice.company_id == cid, PurchaseInvoice.id.in_(purchase_id or [-1])
        )).all()
        try:
            with ERPFlexClient(settings) as client:
                for inv in invoices:
                    payload = _load_raw_payload(db, inv.erpflex_raw_id)
                    api_id = text(deep_find(payload or {}, ("id", "compra_id", "id_compra", "id_informacoes_fiscais"))) or text(inv.erpflex_id)
                    if not api_id:
                        failed += 1
                        continue
                    try:
                        detail = client.purchase_detail(api_id)
                    except Exception:
                        detail = None
                    if not detail:
                        failed += 1
                        continue
                    merged = dict(payload or {})
                    merged.update(detail)
                    if inv.erpflex_raw_id:
                        raw = db.get(RawRecord, inv.erpflex_raw_id)
                        if raw:
                            raw.payload_json = safe_json(merged)
                            raw.record_hash = payload_hash(merged)
                            raw.updated_at = datetime.utcnow()
                    count = _persist_purchase_items(db, inv, merged)
                    if count:
                        updated += 1
                    else:
                        failed += 1
                db.commit()
        except Exception:
            failed += len(invoices)
    return RedirectResponse(f"/compras?items_updated={updated}&items_failed={failed}", status_code=303)


def _purchase_note_context(db, cid: int, purchase_id: int):
    row = db.execute(
        select(PurchaseInvoice, Partner)
        .outerjoin(Partner, PurchaseInvoice.supplier_id == Partner.id)
        .where(PurchaseInvoice.company_id == cid, PurchaseInvoice.id == purchase_id)
    ).first()
    if not row:
        return None
    inv, supplier = row
    items = db.scalars(select(PurchaseItem).where(PurchaseItem.purchase_invoice_id == inv.id).order_by(PurchaseItem.id)).all()
    payload = _load_raw_payload(db, inv.erpflex_raw_id)
    if not items and payload:
        items = _sales_items_from_payload(payload)
    def val(*names, default=""):
        v = deep_find(payload, tuple(names)) if payload else None
        return v if v not in (None, "") else default
    header = {
        "number": inv.number or _specific_nf_number(payload),
        "series": inv.series or text(val("serie_da_nf", "serie", "serie_nf")),
        "model": text(val("modelo_da_nf", "modelo", "modelo_nf")),
        "key": inv.access_key or access_key_from(payload) or "",
        "issue_date": inv.issue_date or text(val("emissao", "emissao_original", "data_emissao")),
        "entry_date": text(val("data_saida", "data_entrada", "entrada")),
        "nature": text(val("natureza_da_operacao", "natureza_operacao", "natureza")),
        "cfop": text(val("cod_cfop", "cfop")),
        "status": _invoice_status(payload).get("raw") or text(val("status_danfe", "status", "situacao")),
        "supplier_cnpj": (supplier.cnpj_cpf if supplier else "") or text(val("cnpj", "cpf_cnpj", "cnpj_cpf")),
        "supplier_name": (supplier.name if supplier else "") or text(val("cliente", "razao_social", "fornecedor")),
        "supplier_trade": (supplier.trade_name if supplier else "") or text(val("nome_fantasia", "fantasia")),
        "products_total": fnum(val("valor_total_produtos", "valor_produtos", "total_produtos")),
        "freight": fnum(val("valor_frete", "frete")),
        "discount": fnum(val("valor_desconto", "desconto")),
        "base_icms": fnum(val("base_icms", "valor_base_icms")),
        "icms": fnum(val("valor_icms", "icms")),
        "icms_st": fnum(val("valor_icmsst", "valor_icms_st", "icms_st")),
        "ipi": fnum(val("valor_ipi", "ipi")),
        "pis": fnum(val("valor_pis", "pis")),
        "cofins": fnum(val("valor_cofins", "cofins")),
        "total": inv.total_erpflex or fnum(val("valor_total_da_nota", "valor_total", "total", "valor")),
        "supplier_address": (supplier.address if supplier else "") or text(val("endereco_fornecedor", "logradouro_fornecedor", "endereco_emitente")),
        "supplier_district": (supplier.district if supplier else "") or text(val("bairro_fornecedor", "bairro_emitente")),
        "supplier_city": (supplier.city if supplier else "") or text(val("municipio_fornecedor", "cidade_fornecedor", "municipio_emitente")),
        "supplier_state": (supplier.state if supplier else "") or text(val("uf_fornecedor", "uf_emitente")),
        "supplier_zip": (supplier.zip_code if supplier else "") or text(val("cep_fornecedor", "cep_emitente")),
        "supplier_ie": text(val("ie_fornecedor", "inscricao_estadual_fornecedor", "ie_emitente")),
        "protocol": text(val("protocolo_autorizacao", "protocolo", "numero_protocolo")),
        "authorization_date": text(val("data_autorizacao", "dh_recibo", "data_protocolo")),
        "carrier": text(val("transportadora", "transportador", "nome_transportadora")),
        "volumes": fnum(val("quantidade_volume", "qtd_volume", "volumes", "qvol")),
        "freight_mode": text(val("modalidade_frete", "mod_frete")),
        "gross_weight": fnum(val("peso_bruto", "pesoBruto")),
        "net_weight": fnum(val("peso_liquido", "pesoLiquido")),
        "additional_info": text(val("informacoes_complementares", "informacoes_adicionais", "observacao", "observacoes")),
    }
    return inv, supplier, items, payload, header


@app.get("/compras/{purchase_id}", response_class=HTMLResponse)
def compra_detalhe(request: Request, purchase_id: int):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        ctx = _purchase_note_context(db, cid, purchase_id)
        if not ctx:
            return HTMLResponse("Compra não encontrada.", status_code=404)
        inv, supplier, items, payload, header = ctx
    return render(request, "compra_nota.html", company=company, inv=inv, supplier=supplier, items=items, payload=payload, header=header, print_mode=False)


@app.get("/compras/{purchase_id}/imprimir", response_class=HTMLResponse)
def compra_imprimir(request: Request, purchase_id: int):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        ctx = _purchase_note_context(db, cid, purchase_id)
        if not ctx:
            return HTMLResponse("Compra não encontrada.", status_code=404)
        inv, supplier, items, payload, header = ctx
    return render(request, "danfe_impressao.html", company=company, inv=inv, supplier=supplier, items=items,
                  payload=payload, header=header, document_kind="entrada")


def _linked_invoice_for_title(db, cid: int, title: FinancialTitle):
    """Relaciona título por ID fiscal/NF e, como último fallback, por pedido exato e único."""
    payload = _title_payload(db, title)
    nf = _title_nf_number(payload)
    faturamento_id = _title_faturamento_id(payload)
    order_number = _title_order_number(payload, title.document or "")
    if title.kind == "RECEBER":
        inv = None
        if faturamento_id:
            inv = db.scalar(select(SalesInvoice).where(
                SalesInvoice.company_id == cid, SalesInvoice.erpflex_id == faturamento_id
            ).order_by(SalesInvoice.id.desc()).limit(1))
        if not inv and nf:
            stripped = nf.lstrip("0") or "0"
            inv = db.scalar(select(SalesInvoice).where(SalesInvoice.company_id == cid, or_(
                SalesInvoice.number == nf, func.ltrim(SalesInvoice.number, "0") == stripped
            )).order_by(SalesInvoice.id.desc()).limit(1))
        if not inv and order_number:
            inv = _unique_invoice_by_order(db, cid, order_number, "RECEBER")
        return ("faturamento", inv) if inv else (None, None)
    if title.kind == "PAGAR":
        inv = None
        if nf:
            stripped = nf.lstrip("0") or "0"
            inv = db.scalar(select(PurchaseInvoice).where(PurchaseInvoice.company_id == cid, or_(
                PurchaseInvoice.number == nf, func.ltrim(PurchaseInvoice.number, "0") == stripped
            )).order_by(PurchaseInvoice.id.desc()).limit(1))
        if not inv and order_number:
            inv = _unique_invoice_by_order(db, cid, order_number, "PAGAR")
        return ("compras", inv) if inv else (None, None)
    return None, None


@app.get("/financeiro/cobranca/configuracao", response_class=HTMLResponse)
def cobranca_configuracao(request: Request, saved: int = 0, auto_sent: int = 0, auto_failed: int = 0, q: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        cfg = _collection_config(db, cid)
        stmt = select(Partner).where(Partner.company_id == cid, Partner.role_customer == True)
        if q.strip():
            stmt = stmt.where(or_(Partner.name.ilike(f"%{q.strip()}%"), Partner.trade_name.ilike(f"%{q.strip()}%"), Partner.cnpj_cpf.ilike(f"%{q.strip()}%")))
        customers = db.scalars(stmt.order_by(Partner.name).limit(1000)).all()
        rules = _collection_contract_map(db, cid, [p.id for p in customers])
        recent = db.execute(select(CollectionSendLog, FinancialTitle).join(FinancialTitle, CollectionSendLog.financial_title_id == FinancialTitle.id).where(
            CollectionSendLog.company_id == cid
        ).order_by(CollectionSendLog.sent_at.desc()).limit(30)).all()
    return render(request, "cobranca_configuracao.html", company=company, cfg=cfg, customers=customers, rules=rules, recent=recent, q=q, saved=saved, auto_sent=auto_sent, auto_failed=auto_failed)


@app.post("/financeiro/cobranca/configuracao")
def cobranca_configuracao_salvar(request: Request, active: str = Form(""), days_before_due: int = Form(3), include_overdue: str = Form(""), resend_interval_hours: int = Form(24), subject_template: str = Form(""), body_template: str = Form(""), contract_customer_ids: list[int] = Form(default=[])):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        cfg = _collection_config(db, cid)
        cfg.active = active == "1"
        cfg.days_before_due = max(0, min(int(days_before_due or 0), 60))
        cfg.include_overdue = include_overdue == "1"
        cfg.resend_interval_hours = max(1, min(int(resend_interval_hours or 24), 720))
        cfg.subject_template = subject_template.strip() or None
        cfg.body_template = body_template.strip() or None
        current = {r.partner_id:r for r in db.scalars(select(CollectionCustomerRule).where(CollectionCustomerRule.company_id == cid)).all()}
        selected = set(int(x) for x in contract_customer_ids)
        for pid in selected:
            rule = current.get(pid)
            if not rule:
                db.add(CollectionCustomerRule(company_id=cid, partner_id=pid, has_collection_contract=True))
            else:
                rule.has_collection_contract = True
        for pid, rule in current.items():
            if pid not in selected:
                rule.has_collection_contract = False
        db.commit()
    return RedirectResponse("/financeiro/cobranca/configuracao?saved=1", status_code=303)


@app.post("/financeiro/cobranca/executar")
def cobranca_executar_agora(request: Request):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
    result = _run_automatic_collection_company(cid)
    return RedirectResponse(f"/financeiro/cobranca/configuracao?auto_sent={result['sent']}&auto_failed={result['failed']}", status_code=303)


@app.get("/financeiro", response_class=HTMLResponse)
def financeiro(request: Request, kind: str = "RECEBER", q: str = "", period: str = "all", date_from: str = "", date_to: str = "",
               date_field: str = "due", bank_filter: str = "", wallet_filter: str = "", boleto_filter: str = "", payment_filter: str = "", page: int = 1, per_page: int = 0, cols: str = "", save_view: int = 0,
               enriched: int = 0, enrich_failed: int = 0, batch_sent: int = 0, batch_failed: int = 0, batch_skipped: int = 0):
    kind = kind if kind in {"RECEBER", "PAGAR", "DESPESA"} else "RECEBER"
    date_field = date_field if date_field in {"issue", "due"} else "due"
    boleto_filter = boleto_filter if boleto_filter in {"", "DISPONIVEL", "PRONTO", "PAGO", "INDISPONIVEL", "CONTRATO"} else ""
    payment_filter = payment_filter if payment_filter in {"", "ABERTO", "PARCIAL", "VENCIDO", "PAGO", "CANCELADO"} else ""
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        column_defs, selected_cols, pp = _module_view_settings(db, request, cid, "financeiro", cols, per_page, save_view)
        if kind == "DESPESA":
            # Despesas: consulta operacional simples, sem banco/carteira/boleto.
            selected_cols = ["title_number", "partner", "issue_date", "due_date", "value", "status"]
        period, start, end = _period_bounds(period, date_from, date_to)
        stmt = select(FinancialTitle, Partner).outerjoin(Partner, FinancialTitle.partner_id == Partner.id).where(FinancialTitle.company_id == cid, FinancialTitle.kind == kind)
        count_stmt = select(func.count(FinancialTitle.id)).outerjoin(Partner, FinancialTitle.partner_id == Partner.id).where(FinancialTitle.company_id == cid, FinancialTitle.kind == kind)
        sum_stmt = select(func.sum(FinancialTitle.value)).outerjoin(Partner, FinancialTitle.partner_id == Partner.id).where(FinancialTitle.company_id == cid, FinancialTitle.kind == kind)
        balance_stmt = select(func.sum(FinancialTitle.value - FinancialTitle.paid_value)).outerjoin(Partner, FinancialTitle.partner_id == Partner.id).where(FinancialTitle.company_id == cid, FinancialTitle.kind == kind)
        if q.strip():
            like = f"%{q.strip()}%"
            cond = or_(FinancialTitle.document.ilike(like), Partner.name.ilike(like), FinancialTitle.status.ilike(like), FinancialTitle.bank.ilike(like), FinancialTitle.wallet.ilike(like))
            stmt=stmt.where(cond); count_stmt=count_stmt.where(cond); sum_stmt=sum_stmt.where(cond); balance_stmt=balance_stmt.where(cond)
        if bank_filter.strip():
            bank_like = f"%{bank_filter.strip()}%"
            stmt=stmt.where(FinancialTitle.bank.ilike(bank_like)); count_stmt=count_stmt.where(FinancialTitle.bank.ilike(bank_like)); sum_stmt=sum_stmt.where(FinancialTitle.bank.ilike(bank_like)); balance_stmt=balance_stmt.where(FinancialTitle.bank.ilike(bank_like))
        if wallet_filter.strip():
            wallet_like = f"%{wallet_filter.strip()}%"
            stmt=stmt.where(FinancialTitle.wallet.ilike(wallet_like)); count_stmt=count_stmt.where(FinancialTitle.wallet.ilike(wallet_like)); sum_stmt=sum_stmt.where(FinancialTitle.wallet.ilike(wallet_like)); balance_stmt=balance_stmt.where(FinancialTitle.wallet.ilike(wallet_like))
        date_col = FinancialTitle.issue_date if date_field == "issue" else FinancialTitle.due_date
        stmt=_apply_date_range(stmt,date_col,start,end); count_stmt=_apply_date_range(count_stmt,date_col,start,end); sum_stmt=_apply_date_range(sum_stmt,date_col,start,end); balance_stmt=_apply_date_range(balance_stmt,date_col,start,end)
        boleto_statuses = {}
        payment_statuses = {}
        if (kind == "RECEBER" and boleto_filter) or (kind == "PAGAR" and payment_filter):
            candidate_rows = db.execute(stmt.order_by(date_col.desc(), FinancialTitle.id.desc())).all()
            filtered_rows = []
            for t, p in candidate_rows:
                if kind == "RECEBER":
                    code, label = _boleto_status(db, t)
                    boleto_statuses[t.id] = {"code": code, "label": label}
                    keep = code == boleto_filter
                else:
                    code, label = _payment_status(t)
                    payment_statuses[t.id] = {"code": code, "label": label}
                    keep = code == payment_filter
                if keep:
                    filtered_rows.append((t, p))
            total_rows = len(filtered_rows)
            pages=max(1,(total_rows+pp-1)//pp); page=min(max(1,page),pages)
            rows=filtered_rows[(page-1)*pp:page*pp]
            total=sum(float(t.value or 0) for t,_ in filtered_rows)
            total_balance=sum(float(t.value or 0)-float(t.paid_value or 0) for t,_ in filtered_rows)
        else:
            total_rows=db.scalar(count_stmt) or 0; pages=max(1,(total_rows+pp-1)//pp); page=min(max(1,page),pages)
            rows=db.execute(stmt.order_by(date_col.desc(),FinancialTitle.id.desc()).offset((page-1)*pp).limit(pp)).all()
            total=db.scalar(sum_stmt) or 0
            total_balance=db.scalar(balance_stmt) or 0
            if kind == "RECEBER":
                for t,_p in rows:
                    code,label=_boleto_status(db,t); boleto_statuses[t.id]={"code":code,"label":label}
            elif kind == "PAGAR":
                for t,_p in rows:
                    code,label=_payment_status(t); payment_statuses[t.id]={"code":code,"label":label}
        links={}
        display={}
        remittance_statuses = {}
        partner_ids = [t.partner_id for t,_p in rows if t.partner_id]
        collection_contracts = _collection_contract_map(db, cid, partner_ids) if kind == "RECEBER" else {}
        bank_names, wallet_names = _finance_maps(db,cid)
        for t,p in rows:
            display[t.id]=_financial_display(db,t,p,bank_names,wallet_names)
            module, inv = _linked_invoice_for_title(db,cid,t)
            if inv:
                links[t.id]=(module,inv.id)
                if not display[t.id].get("nf_number") and getattr(inv, "number", None):
                    display[t.id]["nf_number"] = inv.number
            if kind == "RECEBER":
                rcode,rlabel = _remittance_status(db,t); remittance_statuses[t.id]={"code":rcode,"label":rlabel}
            if kind == "PAGAR" and t.id not in payment_statuses:
                code,label=_payment_status(t); payment_statuses[t.id]={"code":code,"label":label}
        params={"kind":kind,"q":q,"period":period,"date_from":date_from,"date_to":date_to,"date_field":date_field,"bank_filter":bank_filter,"wallet_filter":wallet_filter,"boleto_filter":boleto_filter,"payment_filter":payment_filter,"per_page":pp,"cols":",".join(selected_cols)}
        prev_url,next_url=_page_urls("/financeiro",page,pages,params)
    return render(request,"financeiro.html",company=company,rows=rows,kind=kind,q=q,total=total,
                  period=period,date_from=date_from,date_to=date_to,date_field=date_field,bank_filter=bank_filter,wallet_filter=wallet_filter,boleto_filter=boleto_filter,payment_filter=payment_filter,period_label=_period_label(period,start,end),
                  column_defs=column_defs,selected_cols=selected_cols,per_page=pp,page=page,pages=pages,total_rows=total_rows,
                  prev_url=prev_url,next_url=next_url,links=links,display=display,total_balance=total_balance,enriched=enriched,enrich_failed=enrich_failed,
                  boleto_statuses=boleto_statuses,payment_statuses=payment_statuses,remittance_statuses=remittance_statuses,collection_contracts=collection_contracts,batch_sent=batch_sent,batch_failed=batch_failed,batch_skipped=batch_skipped)


@app.post("/financeiro/enriquecer")
def financeiro_enriquecer(request: Request, kind: str = Form("RECEBER"), q: str = Form(""), period: str = Form("all"),
                         date_from: str = Form(""), date_to: str = Form(""), date_field: str = Form("due"),
                         per_page: int = Form(25), page: int = Form(1)):
    if kind != "RECEBER":
        return RedirectResponse("/financeiro?kind=" + quote_plus(kind), status_code=303)
    success = failed = 0
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        period, start, end = _period_bounds(period, date_from, date_to)
        stmt = select(FinancialTitle).where(FinancialTitle.company_id == cid, FinancialTitle.kind == "RECEBER")
        if q.strip():
            like = f"%{q.strip()}%"
            stmt = stmt.outerjoin(Partner, FinancialTitle.partner_id == Partner.id).where(or_(FinancialTitle.document.ilike(like), Partner.name.ilike(like)))
        date_col = FinancialTitle.issue_date if date_field == "issue" else FinancialTitle.due_date
        stmt = _apply_date_range(stmt, date_col, start, end)
        page_size = max(int(per_page or 25), 1)
        lim = min(page_size, 100)
        titles = db.scalars(stmt.order_by(date_col.desc(), FinancialTitle.id.desc()).offset(max(0,(int(page or 1)-1)*page_size)).limit(lim)).all()
        settings = get_erpflex_settings(db, cid)
        bank_names, wallet_names = _finance_maps(db, cid)
        with ERPFlexClient(settings) as client:
            # O V7.8 resolve o nome do banco a partir de /api/bancos/. Atualiza o catálogo
            # uma única vez antes de enriquecer os títulos (inclui campos SA6_ID/SA6_Desc).
            _refresh_bank_catalog(db, client, cid)
            _refresh_boleto_options(db, client, cid)
            bank_names, wallet_names = _finance_maps(db, cid)
            for title in titles:
                try:
                    detail = client.receivable_title_detail(title.erpflex_id)
                    if detail:
                        bank_id = text(deep_find(detail, ("id_banco", "banco_id", "idBanco")))
                        if bank_id and bank_id not in bank_names:
                            if _ensure_bank_detail(db, client, cid, bank_id):
                                bank_names, wallet_names = _finance_maps(db, cid)
                        _save_receivable_detail(db, title, detail, bank_names, wallet_names)
                        success += 1
                    else:
                        failed += 1
                except Exception:
                    failed += 1
        db.commit()
    params = {"kind":kind,"q":q,"period":period,"date_from":date_from,"date_to":date_to,"date_field":date_field,"per_page":per_page,"page":page,"enriched":success,"enrich_failed":failed}
    return RedirectResponse("/financeiro?" + urlencode(params), status_code=303)


@app.post("/financeiro/lote")
def financeiro_lote(request: Request, title_ids: list[int] = Form([]), action: str = Form("")):
    if not title_ids:
        return RedirectResponse("/financeiro?kind=RECEBER&batch_failed=1", status_code=303)
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        titles = db.scalars(select(FinancialTitle).where(
            FinancialTitle.company_id == cid, FinancialTitle.kind == "RECEBER", FinancialTitle.id.in_(title_ids)
        ).order_by(FinancialTitle.due_date, FinancialTitle.id)).all()
        if not titles:
            return RedirectResponse("/financeiro?kind=RECEBER&batch_failed=1", status_code=303)
        if action == "imprimir_boletos":
            pages = []; failed = 0
            settings = get_erpflex_settings(db, cid)
            with ERPFlexClient(settings) as client:
                for title in titles:
                    code, _label = _boleto_status(db, title)
                    if code in {"PAGO", "CONTRATO"}:
                        failed += 1; continue
                    try:
                        html, _ids = _fetch_boleto_html(db, cid, title, client, force=False)
                        pages.append({"title": title, "html": html})
                    except Exception:
                        failed += 1
                db.commit()
            return render(request, "boletos_lote.html", company=company, pages=pages, failed=failed)
        if action == "enviar_boletos_xml":
            smtp = get_smtp_settings(db, cid)
            if not smtp.get("active") or not smtp.get("host") or not smtp.get("from_email"):
                return RedirectResponse("/financeiro?kind=RECEBER&batch_failed=" + str(len(titles)), status_code=303)
            sent = failed = skipped = 0
            settings = get_erpflex_settings(db, cid)
            with ERPFlexClient(settings) as client:
                for title in titles:
                    code, _label = _boleto_status(db, title)
                    if code in {"PAGO", "CONTRATO"}:
                        skipped += 1; continue
                    try:
                        recipient = _customer_email_for_title(db, cid, title, client)
                        if not recipient or "@" not in recipient:
                            failed += 1; continue
                        html, _ids = _fetch_boleto_html(db, cid, title, client, force=False)
                        display = _financial_display(db, title, db.get(Partner, title.partner_id) if title.partner_id else None)
                        nf = display.get("nf_number") or ""
                        xml_text, xml_filename = _sales_invoice_xml_for_title(db, cid, title)
                        subject = f"Boleto bancário - NF-e {nf}" if nf else f"Boleto bancário - {title.document or title.erpflex_id}"
                        body = "Prezados,\n\nSegue anexo o boleto bancário"
                        if xml_text: body += " e o XML da NF-e"
                        if nf: body += f" referente à NF-e {nf}"
                        if title.due_date: body += f", com vencimento em {title.due_date}"
                        body += ".\n\nAtenciosamente."
                        filename = f"boleto_NF_{nf}.html" if nf else f"boleto_{title.erpflex_id}.html"
                        _smtp_send_boleto(smtp, recipient, subject, body, html, filename, xml_text, xml_filename)
                        ac = access_context(db, request)
                        add_audit(db, ac, "BOLETO_XML_LOTE_ENVIADO", "financial_title", title.id, f"destinatário {recipient}; xml={'sim' if xml_text else 'não'}")
                        db.commit()
                        sent += 1
                    except Exception:
                        db.rollback()
                        failed += 1
            return RedirectResponse(f"/financeiro?kind=RECEBER&batch_sent={sent}&batch_failed={failed}&batch_skipped={skipped}", status_code=303)
    return RedirectResponse("/financeiro?kind=RECEBER&batch_failed=1", status_code=303)


@app.post("/financeiro/{title_id}/enriquecer")
def financeiro_enriquecer_titulo(request: Request, title_id: int):
    ok = 0
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        title = db.scalar(select(FinancialTitle).where(FinancialTitle.company_id == cid, FinancialTitle.id == title_id).limit(1))
        if not title or title.kind != "RECEBER":
            return RedirectResponse(f"/financeiro/{title_id}", status_code=303)
        settings = get_erpflex_settings(db, cid)
        with ERPFlexClient(settings) as client:
            try:
                detail = client.receivable_title_detail(title.erpflex_id)
                if detail:
                    bank_id = text(deep_find(detail, ("id_banco", "banco_id", "idBanco")))
                    bank_names, wallet_names = _finance_maps(db, cid)
                    if bank_id and bank_id not in bank_names:
                        _refresh_bank_catalog(db, client, cid)
                    _refresh_boleto_options(db, client, cid)
                    bank_names, wallet_names = _finance_maps(db, cid)
                    _save_receivable_detail(db, title, detail, bank_names, wallet_names)
                    db.commit(); ok = 1
            except Exception:
                ok = 0
    return RedirectResponse(f"/financeiro/{title_id}?enriched={ok}", status_code=303)


@app.get("/financeiro/{title_id}/boleto", response_class=HTMLResponse)
def financeiro_boleto(request: Request, title_id: int, refresh: int = 0, sent: int = 0, send_error: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        title = db.scalar(select(FinancialTitle).where(
            FinancialTitle.company_id == cid, FinancialTitle.id == title_id, FinancialTitle.kind == "RECEBER"
        ).limit(1))
        if not title:
            return HTMLResponse("Título a receber não encontrado.", status_code=404)
        settings = get_erpflex_settings(db, cid)
        smtp = get_smtp_settings(db, cid)
        html = ""; error = ""; ids = {}
        recipient = ""
        boleto_code, boleto_label = _boleto_status(db, title)
        if boleto_code == "PAGO":
            error = "Título já pago/liquidado. Não há boleto em aberto para impressão ou reenvio."
        else:
            try:
                with ERPFlexClient(settings) as client:
                    html, ids = _fetch_boleto_html(db, cid, title, client, force=bool(refresh))
                    recipient = _customer_email_for_title(db, cid, title, client)
                    db.commit()
            except Exception as exc:
                error = str(exc)
                db.rollback()
        partner = db.get(Partner, title.partner_id) if title.partner_id else None
        display = _financial_display(db, title, partner)
        nf_number = display.get("nf_number") or ""
    return render(request, "boleto.html", company=company, title=title, display=display, boleto_html=html,
                  boleto_error=error, boleto_ids=ids, recipient=recipient, smtp=smtp, nf_number=nf_number,
                  boleto_status=boleto_label, boleto_status_code=boleto_code, sent=sent, send_error=send_error)


@app.get("/financeiro/{title_id}/boleto/arquivo", response_class=HTMLResponse)
def financeiro_boleto_arquivo(request: Request, title_id: int, imprimir: int = 0):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        title = db.scalar(select(FinancialTitle).where(
            FinancialTitle.company_id == cid, FinancialTitle.id == title_id, FinancialTitle.kind == "RECEBER"
        ).limit(1))
        if not title:
            return HTMLResponse("Título não encontrado.", status_code=404)
        code, _label = _boleto_status(db, title)
        if code == "PAGO":
            return HTMLResponse("Título já pago/liquidado. Não há boleto em aberto para impressão.", status_code=409)
        try:
            with ERPFlexClient(get_erpflex_settings(db, cid)) as client:
                html, _ids = _fetch_boleto_html(db, cid, title, client, force=False)
                db.commit()
        except Exception as exc:
            return HTMLResponse(f"Não foi possível obter o boleto: {str(exc)}", status_code=502)
    if imprimir:
        import re
        script = "<script>window.addEventListener('load',()=>window.print());</script>"
        if re.search(r"</body\s*>", html, flags=re.IGNORECASE):
            html = re.sub(r"</body\s*>", script + "</body>", html, count=1, flags=re.IGNORECASE)
        else:
            html += script
    return HTMLResponse(html)


@app.post("/financeiro/{title_id}/boleto/enviar")
def financeiro_boleto_enviar(request: Request, title_id: int, recipient: str = Form("")):
    recipient = (recipient or "").strip()
    if not recipient or "@" not in recipient:
        return RedirectResponse(f"/financeiro/{title_id}/boleto?send_error=" + quote_plus("Informe um e-mail válido."), status_code=303)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        title = db.scalar(select(FinancialTitle).where(
            FinancialTitle.company_id == cid, FinancialTitle.id == title_id, FinancialTitle.kind == "RECEBER"
        ).limit(1))
        if not title:
            return HTMLResponse("Título não encontrado.", status_code=404)
        smtp = get_smtp_settings(db, cid)
        if not smtp.get("active") or not smtp.get("host") or not smtp.get("from_email"):
            return RedirectResponse(f"/financeiro/{title_id}/boleto?send_error=" + quote_plus("Configure o SMTP em Integrações antes de reenviar o boleto."), status_code=303)
        try:
            with ERPFlexClient(get_erpflex_settings(db, cid)) as client:
                html, _ids = _fetch_boleto_html(db, cid, title, client, force=False)
                display = _financial_display(db, title, db.get(Partner, title.partner_id) if title.partner_id else None)
                nf = display.get("nf_number") or ""
                subject = f"Boleto bancário - NF-e {nf}" if nf else f"Boleto bancário - {title.document or title.erpflex_id}"
                body = "Prezados,\n\nSegue anexo o boleto bancário referente ao título"
                if nf: body += f" da NF-e {nf}"
                if title.due_date: body += f", com vencimento em {title.due_date}"
                body += ".\n\nAtenciosamente."
                filename = f"boleto_NF_{nf}.html" if nf else f"boleto_{title.erpflex_id}.html"
                xml_text, xml_filename = _sales_invoice_xml_for_title(db, cid, title)
                _smtp_send_boleto(smtp, recipient, subject, body, html, filename, xml_text, xml_filename)
                ac = access_context(db, request)
                add_audit(db, ac, "BOLETO_REENVIADO", "financial_title", title.id, f"destinatário {recipient}")
                db.commit()
            return RedirectResponse(f"/financeiro/{title_id}/boleto?sent=1", status_code=303)
        except Exception as exc:
            db.rollback()
            return RedirectResponse(f"/financeiro/{title_id}/boleto?send_error=" + quote_plus(str(exc)[:400]), status_code=303)


@app.get("/financeiro/{title_id}", response_class=HTMLResponse)
def financeiro_detalhe(request: Request, title_id: int, enriched: int = 0):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        row=db.execute(select(FinancialTitle,Partner).outerjoin(Partner,FinancialTitle.partner_id==Partner.id)
                       .where(FinancialTitle.company_id==cid,FinancialTitle.id==title_id)).first()
        if not row:
            return HTMLResponse("Título não encontrado.",status_code=404)
        title,partner=row
        payload=_load_raw_payload(db,title.raw_record_id)
        bank_names, wallet_names = _finance_maps(db,cid)
        display=_financial_display(db,title,partner,bank_names,wallet_names)
        linked_module,linked_inv=_linked_invoice_for_title(db,cid,title)
    return render(request,"financeiro_detalhe.html",company=company,title=title,partner=partner,payload=payload,display=display,
                  linked_module=linked_module,linked_inv=linked_inv,enriched=enriched)


def _delivery_rows(db, cid: int, *, status: str = "", region: str = "", responsible_id: int = 0, q: str = "", limit: int = 500):
    stmt = (select(Delivery, SalesInvoice, Partner, DeliveryAssignment, Responsible)
            .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
            .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
            .outerjoin(DeliveryAssignment, DeliveryAssignment.delivery_id == Delivery.id)
            .outerjoin(Responsible, DeliveryAssignment.responsible_id == Responsible.id)
            .where(Delivery.company_id == cid))
    if status:
        stmt = stmt.where(Delivery.status == status)
    if region:
        stmt = stmt.where(Delivery.region == region)
    if responsible_id:
        stmt = stmt.where(DeliveryAssignment.responsible_id == responsible_id)
    if q.strip():
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(SalesInvoice.number.ilike(like), Partner.name.ilike(like),
                              SalesInvoice.delivery_city.ilike(like), SalesInvoice.delivery_district.ilike(like)))
    return db.execute(stmt.order_by(Delivery.id.desc()).limit(limit)).all()


def _active_logistic_catalogs(db, cid: int):
    responsibles = db.scalars(select(Responsible).where(Responsible.company_id == cid, Responsible.active.is_(True)).order_by(Responsible.name)).all()
    regions = db.scalars(select(Region).where(Region.company_id == cid, Region.active.is_(True)).order_by(Region.sort_order, Region.name)).all()
    boxes = db.scalars(select(BoxType).where(BoxType.company_id == cid, BoxType.active.is_(True)).order_by(BoxType.sort_order, BoxType.name)).all()
    return responsibles, regions, boxes


def _set_delivery_status(db, cid: int, delivery: Delivery, new_status: str, note: str = "", movement_type: str = "STATUS"):
    old = delivery.status
    delivery.status = new_status
    db.add(DeliveryMovement(company_id=cid, delivery_id=delivery.id, movement_type=movement_type,
                            old_status=old, new_status=new_status, note=(note or None)))


@app.get("/logistica/inicio", response_class=HTMLResponse)
def logistica_inicio(request: Request):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        stats = dict(db.execute(select(Delivery.status, func.count(Delivery.id)).where(Delivery.company_id == cid).group_by(Delivery.status)).all())
        total_routes = db.scalar(select(func.count(Route.id)).where(Route.company_id == cid)) or 0
        active_responsibles = db.scalar(select(func.count(Responsible.id)).where(Responsible.company_id == cid, Responsible.active.is_(True))) or 0
    return render(request, "logistica_inicio.html", company=company, stats=stats, total_routes=total_routes, active_responsibles=active_responsibles)


@app.get("/logistica", response_class=HTMLResponse)
def logistica(request: Request, status: str = "", region: str = "", responsible_id: int = 0, q: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        rows = _delivery_rows(db, cid, status=status, region=region, responsible_id=responsible_id, q=q)
        responsibles, regions, boxes = _active_logistic_catalogs(db, cid)
        stats = dict(db.execute(select(Delivery.status, func.count(Delivery.id)).where(Delivery.company_id == cid).group_by(Delivery.status)).all())
    return render(request, "logistica.html", company=company, rows=rows, status=status, region=region,
                  responsible_id=responsible_id, q=q, responsibles=responsibles, regions=regions, boxes=boxes, stats=stats)


@app.post("/logistica/{delivery_id}/editar")
def editar_entrega(request: Request, delivery_id: int, region: str = Form(""), responsible_id: int = Form(0),
                   box_type: str = Form(""), volumes: float = Form(0), notes: str = Form(""),
                   remember_rule: str | None = Form(None)):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        delivery = db.get(Delivery, delivery_id)
        if not delivery or delivery.company_id != cid:
            return JSONResponse({"error": "entrega não encontrada"}, status_code=404)
        invoice = db.get(SalesInvoice, delivery.sales_invoice_id)
        partner = db.get(Partner, invoice.customer_id) if invoice and invoice.customer_id else None
        delivery.region = region.strip() or None
        delivery.box_type = box_type.strip() or None
        delivery.volumes = max(0, volumes or 0)
        delivery.notes = notes.strip() or None
        assignment = db.scalar(select(DeliveryAssignment).where(DeliveryAssignment.delivery_id == delivery.id).limit(1))
        if responsible_id:
            resp = db.get(Responsible, responsible_id)
            if not resp or resp.company_id != cid:
                return JSONResponse({"error": "responsável inválido"}, status_code=400)
            if assignment:
                assignment.responsible_id = resp.id
                assignment.assigned_at = datetime.utcnow()
            else:
                db.add(DeliveryAssignment(company_id=cid, delivery_id=delivery.id, responsible_id=resp.id))
            delivery.responsible = resp.name
            delivery.responsible_type = resp.kind
        elif assignment:
            db.delete(assignment)
            delivery.responsible = None
            delivery.responsible_type = None
        if remember_rule and partner:
            rule = None
            if partner.cnpj_cpf:
                rule = db.scalar(select(DeliveryRule).where(DeliveryRule.company_id == cid, DeliveryRule.customer_cnpj == partner.cnpj_cpf).limit(1))
            if not rule:
                rule = DeliveryRule(company_id=cid, customer_cnpj=partner.cnpj_cpf, customer_name=partner.name)
                db.add(rule)
            rule.customer_name = partner.name
            rule.region = delivery.region
            rule.responsible_id = responsible_id or None
            rule.box_type = delivery.box_type
            rule.active = True
        db.commit()
    return RedirectResponse("/logistica", status_code=303)

def _assign_delivery(db, cid: int, delivery: Delivery, responsible_id: int | None):
    assignment = db.scalar(select(DeliveryAssignment).where(DeliveryAssignment.delivery_id == delivery.id).limit(1))
    if responsible_id:
        resp = db.get(Responsible, responsible_id)
        if not resp or resp.company_id != cid or not resp.active:
            raise ValueError("Responsável inválido")
        if assignment:
            assignment.responsible_id = resp.id
            assignment.assigned_at = datetime.utcnow()
        else:
            db.add(DeliveryAssignment(company_id=cid, delivery_id=delivery.id, responsible_id=resp.id))
        delivery.responsible = resp.name
        delivery.responsible_type = resp.kind
        return resp
    if assignment:
        db.delete(assignment)
    delivery.responsible = None
    delivery.responsible_type = None
    return None


def _remember_delivery_rule(db, cid: int, delivery: Delivery, invoice: SalesInvoice, partner: Partner | None, responsible_id: int | None):
    if not partner:
        return
    rule = None
    if partner.cnpj_cpf:
        rule = db.scalar(select(DeliveryRule).where(DeliveryRule.company_id == cid, DeliveryRule.customer_cnpj == partner.cnpj_cpf).order_by(DeliveryRule.id.desc()).limit(1))
    if not rule and partner.name:
        rule = db.scalar(select(DeliveryRule).where(DeliveryRule.company_id == cid, func.lower(DeliveryRule.customer_name) == partner.name.lower()).order_by(DeliveryRule.id.desc()).limit(1))
    if not rule:
        rule = DeliveryRule(company_id=cid, customer_cnpj=partner.cnpj_cpf, customer_name=partner.name)
        db.add(rule)
    rule.customer_cnpj = partner.cnpj_cpf
    rule.customer_name = partner.name
    rule.region = delivery.region
    rule.responsible_id = responsible_id or None
    rule.box_type = delivery.box_type
    rule.active = True


def _remember_region_rule(db, cid: int, partner: Partner | None, region: str):
    """Memoriza somente a região do cliente, preservando motorista/caixa já cadastrados."""
    if not partner or not region.strip():
        return
    rule = None
    if partner.cnpj_cpf:
        rule = db.scalar(select(DeliveryRule).where(
            DeliveryRule.company_id == cid, DeliveryRule.customer_cnpj == partner.cnpj_cpf
        ).order_by(DeliveryRule.id.desc()).limit(1))
    if not rule and partner.name:
        rule = db.scalar(select(DeliveryRule).where(
            DeliveryRule.company_id == cid, func.lower(DeliveryRule.customer_name) == partner.name.lower()
        ).order_by(DeliveryRule.id.desc()).limit(1))
    if not rule:
        rule = DeliveryRule(company_id=cid, customer_cnpj=partner.cnpj_cpf, customer_name=partner.name)
        db.add(rule)
    rule.customer_cnpj = partner.cnpj_cpf
    rule.customer_name = partner.name
    rule.region = region.strip()
    rule.active = True


def _routing_rows(db, cid: int, *, region: str = "", responsible_id: int = 0, q: str = "", include_routed: bool = False, limit: int = 2000):
    stmt = (select(Delivery, SalesInvoice, Partner, DeliveryAssignment, Responsible)
            .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
            .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
            .outerjoin(DeliveryAssignment, DeliveryAssignment.delivery_id == Delivery.id)
            .outerjoin(Responsible, DeliveryAssignment.responsible_id == Responsible.id)
            .where(Delivery.company_id == cid, Delivery.status.in_(["PENDENTE", "REPROGRAMADA"])))
    if not include_routed:
        active_ids = (select(RouteStop.delivery_id).join(Route, RouteStop.route_id == Route.id)
                      .where(Route.company_id == cid, Route.status.in_(["PLANEJADO", "EM ROTA"])))
        stmt = stmt.where(~Delivery.id.in_(active_ids))
    if region:
        stmt = stmt.where(Delivery.region == region)
    if responsible_id == -1:
        stmt = stmt.where(DeliveryAssignment.id.is_(None))
    elif responsible_id > 0:
        stmt = stmt.where(DeliveryAssignment.responsible_id == responsible_id)
    if q.strip():
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(SalesInvoice.number.ilike(like), Partner.name.ilike(like), Partner.trade_name.ilike(like),
                              SalesInvoice.delivery_city.ilike(like), SalesInvoice.delivery_district.ilike(like),
                              SalesInvoice.delivery_address.ilike(like)))
    return db.execute(stmt.order_by(Delivery.region, Responsible.name, SalesInvoice.number, Delivery.id).limit(limit)).all()


def _active_route_rows(db, cid: int, delivery_id: int):
    return db.execute(
        select(RouteStop, Route)
        .join(Route, RouteStop.route_id == Route.id)
        .where(
            Route.company_id == cid,
            RouteStop.delivery_id == delivery_id,
            Route.status.in_(["PLANEJADO", "EM ROTA"]),
        )
        .order_by(Route.id.desc())
    ).all()


def _renumber_route(db, route_id: int):
    stops = db.scalars(select(RouteStop).where(RouteStop.route_id == route_id).order_by(RouteStop.stop_order, RouteStop.id)).all()
    for order, stop in enumerate(stops, 1):
        stop.stop_order = order


def _remove_from_active_routes(db, cid: int, delivery_id: int):
    """Retira a entrega da roteirização ativa sem apagar o histórico de movimentos."""
    affected = []
    for stop, route in _active_route_rows(db, cid, delivery_id):
        affected.append(route)
        db.delete(stop)
        db.flush()
        _renumber_route(db, route.id)
        remaining = db.scalar(select(func.count(RouteStop.id)).where(RouteStop.route_id == route.id)) or 0
        if remaining == 0:
            if route.status == "PLANEJADO":
                cost = db.scalar(select(RouteCost).where(RouteCost.route_id == route.id).limit(1))
                if cost:
                    db.delete(cost)
                db.delete(route)
            elif route.status == "EM ROTA":
                route.status = "FINALIZADO"
                route.finished_at = datetime.utcnow()
    return affected


def _apply_delivery_closeout(db, cid: int, delivery: Delivery, new_status: str, note: str = "", occurrence_type: str = "Ocorrência"):
    active_routes = _active_route_rows(db, cid, delivery.id)
    if not active_routes:
        return False
    assignment = db.scalar(select(DeliveryAssignment).where(DeliveryAssignment.delivery_id == delivery.id).limit(1))
    _set_delivery_status(db, cid, delivery, new_status, note, "BAIXA")
    if new_status == "OCORRENCIA":
        db.add(DeliveryOccurrence(
            company_id=cid,
            delivery_id=delivery.id,
            responsible_id=assignment.responsible_id if assignment else None,
            occurrence_type=occurrence_type.strip() or "Ocorrência",
            note=note.strip() or None,
        ))
    if new_status == "REPROGRAMADA":
        _remove_from_active_routes(db, cid, delivery.id)
        db.add(DeliveryMovement(
            company_id=cid,
            delivery_id=delivery.id,
            movement_type="RETIRADA_ROTEIRO_REPROGRAMACAO",
            old_status="REPROGRAMADA",
            new_status="REPROGRAMADA",
            note="Entrega retirada do roteiro ativo para nova roteirização",
        ))
    return True


@app.get("/roteirizar", response_class=HTMLResponse)
def roteirizar(request: Request, region: str = "", q: str = "", msg: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        # A operação é orientada por ZONA. Só exibimos cards das regiões que
        # realmente possuem NFs disponíveis para roteirização naquele momento.
        card_rows = _routing_rows(db, cid, region="", responsible_id=0, q=q)
        responsibles, regions, boxes = _active_logistic_catalogs(db, cid)
        zone_counts: dict[str, int] = {}
        for delivery, _inv, _p, _assignment, _rsp in card_rows:
            zone = (delivery.region or "A Classificar").strip() or "A Classificar"
            zone_counts[zone] = zone_counts.get(zone, 0) + 1
        order_map = {r.name: (r.sort_order, r.name.lower()) for r in regions}
        region_cards = [
            {"name": name, "count": count}
            for name, count in sorted(
                zone_counts.items(),
                key=lambda item: order_map.get(item[0], (9999, item[0].lower())),
            )
            if count > 0
        ]
        rows = _routing_rows(db, cid, region=region, responsible_id=0, q=q)
        total_volumes = sum((d.volumes or inv.volumes or 0) for d, inv, _p, _a, _r in rows)
        total_value = sum((inv.total or 0) for _d, inv, _p, _a, _r in rows)
        unassigned = sum(1 for _d, _inv, _p, a, _r in rows if not a)
        alerts = sum(1 for d, inv, _p, _a, _r in rows if alerta_endereco(inv.carrier, inv.carrier2) or "VERIFICAR ENDEREÇO" in (d.notes or "").upper())
    return render(
        request, "roteirizar.html", company=company, rows=rows, responsibles=responsibles, regions=regions, boxes=boxes,
        region=region, q=q, msg=msg, total_volumes=total_volumes, total_value=total_value,
        unassigned=unassigned, alerts=alerts, today=date.today().isoformat(),
        zone_counts=zone_counts, region_cards=region_cards, card_total=len(card_rows)
    )


def _roteirizar_redirect(msg: str = "", region: str = "", q: str = "") -> RedirectResponse:
    params = []
    if region.strip():
        params.append(f"region={quote_plus(region.strip())}")
    if q.strip():
        params.append(f"q={quote_plus(q.strip())}")
    if msg:
        params.append(f"msg={quote_plus(msg)}")
    target = "/roteirizar" + ("?" + "&".join(params) if params else "")
    return RedirectResponse(target, status_code=303)


@app.post("/roteirizar/lote")
def roteirizar_lote(request: Request, delivery_ids: list[int] = Form([]), region: str = Form(""),
                    box_type: str = Form(""), remember_rule: str | None = Form(None),
                    return_region: str = Form(""), return_q: str = Form("")):
    if not delivery_ids:
        return _roteirizar_redirect("selecione", return_region, return_q)
    changed = 0
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        deliveries = db.scalars(select(Delivery).where(Delivery.company_id == cid, Delivery.id.in_(delivery_ids))).all()
        for delivery in deliveries:
            invoice = db.get(SalesInvoice, delivery.sales_invoice_id)
            partner = db.get(Partner, invoice.customer_id) if invoice and invoice.customer_id else None
            if region.strip():
                delivery.region = region.strip()
            if box_type.strip():
                delivery.box_type = box_type.strip()
            assignment = db.scalar(select(DeliveryAssignment).where(DeliveryAssignment.delivery_id == delivery.id).limit(1))
            resp_id_for_rule = assignment.responsible_id if assignment else None
            if remember_rule and invoice:
                _remember_delivery_rule(db, cid, delivery, invoice, partner, resp_id_for_rule)
            db.add(DeliveryMovement(company_id=cid, delivery_id=delivery.id, movement_type="AJUSTE_LOGISTICA",
                                    old_status=delivery.status, new_status=delivery.status,
                                    note="Ajuste em lote no Roteirizador"))
            changed += 1
        db.commit()
    return _roteirizar_redirect(f"alteradas-{changed}", return_region, return_q)


@app.post("/roteirizar/classificar-regiao")
def classificar_regiao_selecionadas(
    request: Request, delivery_ids: list[int] = Form([]), region: str = Form(""),
    remember_region: str | None = Form(None), return_region: str = Form(""), return_q: str = Form("")
):
    if not delivery_ids:
        return _roteirizar_redirect("selecione", return_region, return_q)
    if not region.strip():
        return _roteirizar_redirect("selecione-regiao", return_region, return_q)
    changed = 0
    memorized = 0
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        valid_region = db.scalar(select(Region).where(
            Region.company_id == cid, Region.active.is_(True), Region.name == region.strip()
        ).limit(1))
        if not valid_region:
            return _roteirizar_redirect("selecione-regiao", return_region, return_q)
        rows = db.execute(
            select(Delivery, SalesInvoice, Partner)
            .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
            .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
            .where(Delivery.company_id == cid, Delivery.id.in_(delivery_ids), Delivery.status.in_(["PENDENTE", "REPROGRAMADA"]))
        ).all()
        for delivery, invoice, partner in rows:
            old_region = delivery.region
            delivery.region = valid_region.name
            db.add(DeliveryMovement(
                company_id=cid, delivery_id=delivery.id, movement_type="CLASSIFICACAO_REGIAO",
                old_status=delivery.status, new_status=delivery.status,
                note=f"Região: {old_region or 'A Classificar'} → {valid_region.name}",
            ))
            if remember_region and partner:
                _remember_region_rule(db, cid, partner, valid_region.name)
                memorized += 1
            changed += 1
        db.commit()
    suffix = f"regiao-{changed}-{memorized}"
    return _roteirizar_redirect(suffix, return_region, return_q)


@app.post("/roteirizar/classificar-regiao-uma")
def classificar_regiao_uma(
    request: Request, delivery_id: int = Form(...), region: str = Form(""),
    remember_region: str | None = Form(None), return_region: str = Form(""), return_q: str = Form("")
):
    """Informa/corrige a região de uma NF sem interferir nas demais seleções da tela."""
    if not region.strip():
        return _roteirizar_redirect("selecione-regiao", return_region, return_q)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        valid_region = db.scalar(select(Region).where(
            Region.company_id == cid, Region.active.is_(True), Region.name == region.strip()
        ).limit(1))
        if not valid_region:
            return _roteirizar_redirect("selecione-regiao", return_region, return_q)
        row = db.execute(
            select(Delivery, SalesInvoice, Partner)
            .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
            .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
            .where(Delivery.company_id == cid, Delivery.id == delivery_id, Delivery.status.in_(["PENDENTE", "REPROGRAMADA"]))
        ).first()
        if not row:
            return _roteirizar_redirect("entrega-indisponivel", return_region, return_q)
        delivery, invoice, partner = row
        if _active_route_rows(db, cid, delivery.id):
            return _roteirizar_redirect("entrega-ja-roteirizada", return_region, return_q)
        old_region = delivery.region
        delivery.region = valid_region.name
        db.add(DeliveryMovement(
            company_id=cid, delivery_id=delivery.id, movement_type="CLASSIFICACAO_REGIAO",
            old_status=delivery.status, new_status=delivery.status,
            note=f"Região: {old_region or 'A Classificar'} → {valid_region.name}",
        ))
        memorized = 0
        if remember_region and partner:
            _remember_region_rule(db, cid, partner, valid_region.name)
            memorized = 1
        db.commit()
    return _roteirizar_redirect(f"regiao-1-{memorized}", return_region, return_q)


@app.post("/roteirizar/reclassificar")
def reclassificar_pendentes(request: Request, return_region: str = Form(""), return_q: str = Form("")):
    changed = 0
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        rows = db.execute(select(Delivery, SalesInvoice, RawRecord)
                          .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
                          .outerjoin(RawRecord, SalesInvoice.raw_record_id == RawRecord.id)
                          .where(Delivery.company_id == cid, Delivery.status.in_(["PENDENTE", "REPROGRAMADA"]))).all()
        for delivery, inv, raw in rows:
            if delivery.region and delivery.region != "A Classificar":
                continue
            special = regiao_especial(inv.carrier, inv.carrier2)
            channel = raw_channel(raw.payload_json if raw else None)
            new_region = special or classificar_regiao(inv.delivery_city, inv.delivery_district, inv.delivery_zip, channel)
            if new_region and new_region != delivery.region:
                delivery.region = new_region
                changed += 1
            alert = alerta_endereco(inv.carrier, inv.carrier2)
            if alert and alert not in (delivery.notes or "").upper():
                delivery.notes = ((delivery.notes or "") + (" | " if delivery.notes else "") + alert).strip()
            if new_region == "COLETA":
                assignment = db.scalar(select(DeliveryAssignment).where(DeliveryAssignment.delivery_id == delivery.id).limit(1))
                if not assignment:
                    coleta = db.scalar(select(Responsible).where(Responsible.company_id == cid, Responsible.name == "COLETA", Responsible.active.is_(True)).limit(1))
                    if coleta:
                        _assign_delivery(db, cid, delivery, coleta.id)
        db.commit()
    return _roteirizar_redirect(f"reclassificadas-{changed}", return_region, return_q)


@app.post("/roteirizar/criar-roteiro")
def criar_roteiro_selecionado(
    request: Request, delivery_ids: list[int] = Form([]), responsible_id: int = Form(0),
    route_date: str = Form(...), remember_rule: str | None = Form(None),
    return_region: str = Form(""), return_q: str = Form("")
):
    if not delivery_ids:
        return _roteirizar_redirect("selecione", return_region, return_q)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        resp = db.get(Responsible, responsible_id)
        if not resp or resp.company_id != cid or not resp.active:
            return _roteirizar_redirect("selecione-motorista", return_region, return_q)
        active_ids = set(db.scalars(select(RouteStop.delivery_id).join(Route, RouteStop.route_id == Route.id)
                                    .where(Route.company_id == cid, Route.status.in_(["PLANEJADO", "EM ROTA"]))).all())
        candidates = db.execute(select(Delivery, SalesInvoice, Partner)
                                .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
                                .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
                                .where(Delivery.company_id == cid, Delivery.id.in_(delivery_ids),
                                       Delivery.status.in_(["PENDENTE", "REPROGRAMADA"]))
                                .order_by(Delivery.region, SalesInvoice.number, Delivery.id)).all()
        candidates = [(d, inv, partner) for d, inv, partner in candidates if d.id not in active_ids]
        if not candidates:
            return _roteirizar_redirect("sem-entregas", return_region, return_q)
        # Um motorista/data deve ficar concentrado em um roteiro PLANEJADO.
        # Isso evita fragmentar a operação quando o operador roteiriza zona por zona.
        route = db.scalar(select(Route).where(
            Route.company_id == cid, Route.responsible_id == resp.id,
            Route.route_date == route_date, Route.status == "PLANEJADO"
        ).order_by(Route.id.desc()).limit(1))
        if not route:
            route = Route(company_id=cid, responsible_id=resp.id, route_date=route_date, status="PLANEJADO")
            db.add(route)
            db.flush()
            db.add(RouteCost(company_id=cid, route_id=route.id, charge_type=resp.default_freight_type,
                             freight_value=resp.default_freight_value or 0, freight_percent=resp.default_freight_percent))
        next_order = (db.scalar(select(func.max(RouteStop.stop_order)).where(RouteStop.route_id == route.id)) or 0) + 1
        for offset, (delivery, inv, partner) in enumerate(candidates):
            order = next_order + offset
            _assign_delivery(db, cid, delivery, resp.id)
            if remember_rule:
                _remember_delivery_rule(db, cid, delivery, inv, partner, resp.id)
            db.add(RouteStop(route_id=route.id, delivery_id=delivery.id, stop_order=order))
            db.add(DeliveryMovement(company_id=cid, delivery_id=delivery.id, movement_type="ROTEIRIZADA",
                                    old_status=delivery.status, new_status=delivery.status,
                                    note=f"Roteiro #{route.id} - {resp.name} - ordem {order}"))
        routed_count = len(candidates)
        db.commit()
    # O operador volta à mesma zona para continuar roteirizando as notas restantes.
    return _roteirizar_redirect(f"roteirizadas-{routed_count}", return_region, return_q)


@app.get("/importar-logistica", response_class=HTMLResponse)
def importar_logistica_page(request: Request, msg: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        history = db.scalars(select(LogisticsImport).where(LogisticsImport.company_id == cid).order_by(LogisticsImport.id.desc()).limit(30)).all()
    return render(request, "importar_logistica.html", company=company, history=history, msg=msg)


@app.post("/importar-logistica")
async def importar_logistica(request: Request, arquivo: UploadFile = File(...)):
    filename = (arquivo.filename or "arquivo").strip()
    data = await arquivo.read()
    max_bytes = int(os.getenv("MAX_UPLOAD_MB", "25")) * 1024 * 1024
    if len(data) > max_bytes:
        return RedirectResponse("/importar-logistica?msg=arquivo-grande", status_code=303)
    try:
        with SessionLocal() as db:
            cid, _company = company_context(db, request)
            stats = import_logistics_file(db, cid, filename, data)
            db.commit()
        return RedirectResponse(f"/importar-logistica?msg=ok-{stats['invoices']}-{len(stats['errors'])}", status_code=303)
    except Exception as exc:
        return render(request, "importar_logistica.html", company=None, history=[], msg="", error=str(exc))


@app.get("/entregas/{delivery_id}", response_class=HTMLResponse)
def entrega_detalhe(request: Request, delivery_id: int):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        row = db.execute(select(Delivery, SalesInvoice, Partner, DeliveryAssignment, Responsible)
                         .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
                         .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
                         .outerjoin(DeliveryAssignment, DeliveryAssignment.delivery_id == Delivery.id)
                         .outerjoin(Responsible, DeliveryAssignment.responsible_id == Responsible.id)
                         .where(Delivery.company_id == cid, Delivery.id == delivery_id)).first()
        if not row:
            return HTMLResponse("Entrega não encontrada", status_code=404)
        delivery, invoice, partner, assignment, responsible = row
        movements = db.scalars(select(DeliveryMovement).where(DeliveryMovement.delivery_id == delivery.id).order_by(DeliveryMovement.id.desc()).limit(100)).all()
        occurrences = db.scalars(select(DeliveryOccurrence).where(DeliveryOccurrence.delivery_id == delivery.id).order_by(DeliveryOccurrence.id.desc()).limit(100)).all()
        routes = db.execute(select(Route, RouteStop, Responsible)
                            .join(RouteStop, RouteStop.route_id == Route.id)
                            .join(Responsible, Route.responsible_id == Responsible.id)
                            .where(RouteStop.delivery_id == delivery.id).order_by(Route.id.desc())).all()
    return render(request, "entrega_detalhe.html", company=company, delivery=delivery, invoice=invoice, partner=partner,
                  assignment=assignment, responsible=responsible, movements=movements, occurrences=occurrences, routes=routes)


def _dashboard_aggregates(rows):
    total = {"nfs": len(rows), "volumes": 0.0, "value": 0.0, "cost": 0.0, "km": 0.0}
    by_region = {}
    by_resp = {}
    status = {}
    matrix = {}
    for d, inv, p, a, rsp in rows:
        vol = float(d.volumes or inv.volumes or 0)
        value = float(inv.total or 0)
        cost = float(d.freight_cost or 0)
        km = float(d.km or 0)
        region = (d.region or "A Classificar").strip() or "A Classificar"
        resp_name = (rsp.name if rsp else d.responsible) or "Não vinculado"
        resp_kind = (rsp.kind if rsp else d.responsible_type) or ""
        total["volumes"] += vol; total["value"] += value; total["cost"] += cost; total["km"] += km
        status[d.status] = status.get(d.status, 0) + 1
        rg = by_region.setdefault(region, {"name": region, "nfs": 0, "volumes": 0.0, "value": 0.0, "delivered": 0, "occurrences": 0, "responsibles": set()})
        rg["nfs"] += 1; rg["volumes"] += vol; rg["value"] += value; rg["responsibles"].add(resp_name)
        rg["delivered"] += int(d.status == "ENTREGUE"); rg["occurrences"] += int(d.status == "OCORRENCIA")
        rr = by_resp.setdefault(resp_name, {"name": resp_name, "kind": resp_kind, "nfs": 0, "volumes": 0.0, "value": 0.0, "cost": 0.0, "km": 0.0, "delivered": 0, "occurrences": 0, "regions": set()})
        rr["nfs"] += 1; rr["volumes"] += vol; rr["value"] += value; rr["cost"] += cost; rr["km"] += km
        rr["delivered"] += int(d.status == "ENTREGUE"); rr["occurrences"] += int(d.status == "OCORRENCIA"); rr["regions"].add(region)
        matrix.setdefault(region, {})[resp_name] = matrix.setdefault(region, {}).get(resp_name, 0) + 1
    for g in by_region.values():
        g["responsible_count"] = len(g.pop("responsibles"))
    for g in by_resp.values():
        g["region_count"] = len(g.pop("regions"))
        g["cost_per_delivery"] = g["cost"] / g["nfs"] if g["nfs"] else 0
        g["freight_pct"] = (g["cost"] / g["value"] * 100) if g["value"] and g["cost"] else 0
    cost_revenue_base = sum(g["value"] for g in by_resp.values() if g["cost"] > 0)
    total["freight_revenue_base"] = cost_revenue_base
    total["freight_pct"] = (total["cost"] / cost_revenue_base * 100) if cost_revenue_base else 0
    total["delivered_pct"] = (status.get("ENTREGUE", 0) / total["nfs"] * 100) if total["nfs"] else 0
    return total, status, by_region, by_resp, matrix


@app.get("/painel-entregas", response_class=HTMLResponse)
def painel_entregas(request: Request, status_filter: str = "", responsible_id: int = 0, q: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        rows = _delivery_rows(db, cid, status=status_filter, responsible_id=responsible_id, q=q, limit=5000)
        all_rows = _delivery_rows(db, cid, limit=5000)
        responsibles, _regions, _boxes = _active_logistic_catalogs(db, cid)
        total, statuses, _by_region, by_resp, _matrix = _dashboard_aggregates(all_rows)
    return render(request, "painel_entregas.html", company=company, rows=rows, responsibles=responsibles,
                  status_filter=status_filter, responsible_id=responsible_id, q=q, total=total, statuses=statuses, by_resp=by_resp)


@app.get("/painel-gerencial", response_class=HTMLResponse)
def painel_gerencial(request: Request, detail_region: str = "", detail_responsible_id: int = 0):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        rows = _delivery_rows(db, cid, limit=10000)
        total, statuses, by_region, by_resp, matrix = _dashboard_aggregates(rows)
        detail_rows = []
        detail_title = ""
        if detail_region:
            detail_rows = [r for r in rows if (r[0].region or "A Classificar") == detail_region]
            detail_title = f"Região: {detail_region}"
        elif detail_responsible_id:
            detail_rows = [r for r in rows if r[4] and r[4].id == detail_responsible_id]
            resp = next((r[4] for r in rows if r[4] and r[4].id == detail_responsible_id), None)
            detail_title = f"Responsável: {resp.name if resp else detail_responsible_id}"
        resp_ids = {r[4].name: r[4].id for r in rows if r[4]}
        detail_total = _dashboard_aggregates(detail_rows)[0] if detail_rows else None
    return render(request, "painel_gerencial.html", company=company, total=total, statuses=statuses,
                  by_region=by_region, by_resp=by_resp, matrix=matrix, detail_rows=detail_rows,
                  detail_title=detail_title, detail_total=detail_total, resp_ids=resp_ids)


def _report_rows(db, cid: int, *, date_from: str = "", date_to: str = "", region: str = "", responsible_id: int = 0,
                 status_filter: str = "", responsible_kind: str = ""):
    stmt = (select(Delivery, SalesInvoice, Partner, DeliveryAssignment, Responsible)
            .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
            .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
            .outerjoin(DeliveryAssignment, DeliveryAssignment.delivery_id == Delivery.id)
            .outerjoin(Responsible, DeliveryAssignment.responsible_id == Responsible.id)
            .where(Delivery.company_id == cid))
    if date_from:
        stmt = stmt.where(SalesInvoice.forecast_date >= date_from)
    if date_to:
        stmt = stmt.where(SalesInvoice.forecast_date <= date_to)
    if region:
        stmt = stmt.where(Delivery.region == region)
    if responsible_id:
        stmt = stmt.where(DeliveryAssignment.responsible_id == responsible_id)
    if status_filter:
        stmt = stmt.where(Delivery.status == status_filter)
    if responsible_kind:
        stmt = stmt.where(Responsible.kind == responsible_kind)
    return db.execute(stmt.order_by(SalesInvoice.forecast_date, Delivery.region, Responsible.name, SalesInvoice.number).limit(50000)).all()


def _report_context(db, cid: int, date_from: str = "", date_to: str = "", region: str = "", responsible_id: int = 0,
                    status_filter: str = "", responsible_kind: str = ""):
    rows = _report_rows(db, cid, date_from=date_from, date_to=date_to, region=region, responsible_id=responsible_id,
                        status_filter=status_filter, responsible_kind=responsible_kind)
    total, statuses, by_region, by_resp, matrix = _dashboard_aggregates(rows)
    third_party = {k: v for k, v in by_resp.items() if v["kind"] in {"Motorista Externo", "Transportadora"}}
    return rows, total, statuses, by_region, by_resp, matrix, third_party


@app.get("/relatorios-gerenciais", response_class=HTMLResponse)
def relatorios_gerenciais(request: Request, date_from: str = "", date_to: str = "", region: str = "", responsible_id: int = 0,
                          status_filter: str = "", responsible_kind: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        rows, total, statuses, by_region, by_resp, matrix, third_party = _report_context(
            db, cid, date_from, date_to, region, responsible_id, status_filter, responsible_kind)
        responsibles, regions, _boxes = _active_logistic_catalogs(db, cid)
    return render(request, "relatorios_gerenciais.html", company=company, rows=rows, total=total, statuses=statuses,
                  by_region=by_region, by_resp=by_resp, matrix=matrix, third_party=third_party,
                  responsibles=responsibles, regions=regions, date_from=date_from, date_to=date_to, region=region,
                  responsible_id=responsible_id, status_filter=status_filter, responsible_kind=responsible_kind)


@app.get("/relatorios-gerenciais/imprimir", response_class=HTMLResponse)
def relatorios_gerenciais_imprimir(request: Request, date_from: str = "", date_to: str = "", region: str = "", responsible_id: int = 0,
                                   status_filter: str = "", responsible_kind: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        rows, total, statuses, by_region, by_resp, matrix, third_party = _report_context(
            db, cid, date_from, date_to, region, responsible_id, status_filter, responsible_kind)
    return render(request, "relatorios_gerenciais_impressao.html", company=company, rows=rows, total=total, statuses=statuses,
                  by_region=by_region, by_resp=by_resp, third_party=third_party, date_from=date_from, date_to=date_to,
                  region=region, status_filter=status_filter, responsible_kind=responsible_kind)


def _xlsx_sheet(wb, title: str, headers: list[str], rows: list[list]):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for col in range(1, len(headers) + 1):
        max_len = max([len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 500) + 1)] + [8])
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


@app.get("/relatorios-gerenciais/excel")
def relatorios_gerenciais_excel(request: Request, date_from: str = "", date_to: str = "", region: str = "", responsible_id: int = 0,
                                status_filter: str = "", responsible_kind: str = ""):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        rows, total, statuses, by_region, by_resp, matrix, third_party = _report_context(
            db, cid, date_from, date_to, region, responsible_id, status_filter, responsible_kind)
    wb = Workbook(); wb.remove(wb.active)
    _xlsx_sheet(wb, "Resumo", ["Indicador", "Valor"], [
        ["NFs / entregas", total["nfs"]], ["Volumes", total["volumes"]], ["Valor transportado", total["value"]],
        ["Custo de frete", total["cost"]], ["Faturamento com custo informado", total["freight_revenue_base"]],
        ["Frete / faturamento com custo (%)", total["freight_pct"]], ["KM", total["km"]], ["% entregues", total["delivered_pct"]],
    ])
    _xlsx_sheet(wb, "Por Motorista", ["Responsável", "Tipo", "NFs", "Volumes", "Valor", "Custo", "Frete/Faturamento %", "KM", "Entregues", "Ocorrências"], [
        [g["name"], g["kind"], g["nfs"], g["volumes"], g["value"], g["cost"], g["freight_pct"], g["km"], g["delivered"], g["occurrences"]]
        for g in sorted(by_resp.values(), key=lambda x: x["name"])
    ])
    _xlsx_sheet(wb, "Por Região", ["Região", "NFs", "Volumes", "Valor", "Responsáveis", "Entregues", "Ocorrências"], [
        [g["name"], g["nfs"], g["volumes"], g["value"], g["responsible_count"], g["delivered"], g["occurrences"]]
        for g in sorted(by_region.values(), key=lambda x: x["name"])
    ])
    resp_names = sorted(by_resp)
    _xlsx_sheet(wb, "Motorista x Região", ["Região", *resp_names, "Total"], [
        [rg, *[matrix.get(rg, {}).get(rn, 0) for rn in resp_names], sum(matrix.get(rg, {}).values())]
        for rg in sorted(matrix)
    ])
    _xlsx_sheet(wb, "Terceirizados", ["Responsável", "Tipo", "NFs", "Volumes", "Valor", "Custo", "Frete/Faturamento %", "KM"], [
        [g["name"], g["kind"], g["nfs"], g["volumes"], g["value"], g["cost"], g["freight_pct"], g["km"]]
        for g in sorted(third_party.values(), key=lambda x: x["name"])
    ])
    _xlsx_sheet(wb, "NFs Detalhadas", ["Previsão/Entrega", "NF", "Cliente", "Região", "Responsável", "Tipo", "Volumes", "Caixa", "Valor", "Status", "Transportadora 1", "Transportadora 2", "KM", "Custo"], [
        [inv.forecast_date or "", inv.number or "", cliente_exibicao(p.name if p else "", p.trade_name if p else ""), d.region or "A Classificar",
         rsp.name if rsp else (d.responsible or "Não vinculado"), rsp.kind if rsp else (d.responsible_type or ""), d.volumes or inv.volumes or 0,
         d.box_type or "", inv.total or 0, d.status, inv.carrier or "", inv.carrier2 or "", d.km or 0, d.freight_cost or 0]
        for d, inv, p, a, rsp in rows
    ])
    out = io.BytesIO(); wb.save(out)
    filename = f"relatorios_gerenciais_{date.today().isoformat()}.xlsx"
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/roteiros/{route_id}/imprimir", response_class=HTMLResponse)
def imprimir_roteiro(request: Request, route_id: int):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        route = db.get(Route, route_id)
        if not route or route.company_id != cid:
            return HTMLResponse("Roteiro não encontrado", status_code=404)
        resp = db.get(Responsible, route.responsible_id)
        rows = db.execute(select(RouteStop, Delivery, SalesInvoice, Partner)
                          .join(Delivery, RouteStop.delivery_id == Delivery.id)
                          .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
                          .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
                          .where(RouteStop.route_id == route.id).order_by(RouteStop.stop_order)).all()
        total_value = sum(inv.total or 0 for _stop, _d, inv, _p in rows)
        total_volumes = sum(d.volumes or inv.volumes or 0 for _stop, d, inv, _p in rows)
    return render(request, "roteiro_impressao.html", company=company, route=route, responsible=resp, rows=rows,
                  total_value=total_value, total_volumes=total_volumes)


@app.post("/roteiros/{route_id}/ordem")
def salvar_ordem_roteiro(request: Request, route_id: int, stop_ids: list[int] = Form([]), stop_orders: list[int] = Form([])):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        route = db.get(Route, route_id)
        if not route or route.company_id != cid or route.status != "PLANEJADO":
            return JSONResponse({"error": "roteiro inválido para edição"}, status_code=400)
        for sid, order in zip(stop_ids, stop_orders):
            stop = db.get(RouteStop, sid)
            if stop and stop.route_id == route.id:
                stop.stop_order = max(1, int(order or 1))
        db.commit()
    return RedirectResponse(f"/roteiros/{route_id}", status_code=303)


@app.get("/responsaveis", response_class=HTMLResponse)
def responsaveis(request: Request):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        rows = db.scalars(select(Responsible).where(Responsible.company_id == cid).order_by(Responsible.active.desc(), Responsible.name)).all()
    return render(request, "responsaveis.html", company=company, rows=rows)


@app.post("/responsaveis")
def criar_responsavel(request: Request, name: str = Form(...), kind: str = Form("Motorista Interno"),
                      default_freight_type: str = Form("Valor fechado"), default_freight_value: float = Form(0),
                      default_freight_percent: str = Form(""), notes: str = Form("")):
    allowed = {"Motorista Interno", "Motorista Externo", "Transportadora", "COLETA"}
    try:
        default_percent_value = float(str(default_freight_percent).replace(",", ".")) if str(default_freight_percent).strip() else None
    except Exception:
        default_percent_value = None
    if kind not in allowed or not name.strip():
        return JSONResponse({"error": "dados inválidos"}, status_code=400)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        exists = db.scalar(select(Responsible).where(Responsible.company_id == cid, func.lower(Responsible.name) == name.strip().lower(), Responsible.kind == kind).limit(1))
        if not exists:
            db.add(Responsible(company_id=cid, name=name.strip(), kind=kind,
                               default_freight_type=default_freight_type, default_freight_value=max(0, default_freight_value or 0),
                               default_freight_percent=default_percent_value, notes=notes.strip() or None))
            db.commit()
    return RedirectResponse("/responsaveis", status_code=303)


@app.post("/responsaveis/{responsible_id}/toggle")
def toggle_responsavel(request: Request, responsible_id: int):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        row = db.get(Responsible, responsible_id)
        if not row or row.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        row.active = not row.active
        db.commit()
    return RedirectResponse("/responsaveis", status_code=303)


@app.get("/parametros-logistica", response_class=HTMLResponse)
def parametros_logistica(request: Request):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        regions = db.scalars(select(Region).where(Region.company_id == cid).order_by(Region.sort_order, Region.name)).all()
        boxes = db.scalars(select(BoxType).where(BoxType.company_id == cid).order_by(BoxType.sort_order, BoxType.name)).all()
    return render(request, "parametros_logistica.html", company=company, regions=regions, boxes=boxes)


@app.post("/parametros-logistica/regiao")
def criar_regiao(request: Request, name: str = Form(...)):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        clean = name.strip()
        if clean and not db.scalar(select(Region).where(Region.company_id == cid, func.lower(Region.name) == clean.lower()).limit(1)):
            db.add(Region(company_id=cid, name=clean, sort_order=100))
            db.commit()
    return RedirectResponse("/parametros-logistica", status_code=303)


@app.post("/parametros-logistica/caixa")
def criar_caixa(request: Request, name: str = Form(...)):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        clean = name.strip()
        if clean and not db.scalar(select(BoxType).where(BoxType.company_id == cid, func.lower(BoxType.name) == clean.lower()).limit(1)):
            db.add(BoxType(company_id=cid, name=clean, sort_order=100))
            db.commit()
    return RedirectResponse("/parametros-logistica", status_code=303)


@app.post("/parametros-logistica/{kind}/{item_id}/toggle")
def toggle_parametro_logistica(request: Request, kind: str, item_id: int):
    cls = Region if kind == "regiao" else BoxType if kind == "caixa" else None
    if not cls:
        return JSONResponse({"error": "tipo inválido"}, status_code=400)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        row = db.get(cls, item_id)
        if not row or row.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        row.active = not row.active
        db.commit()
    return RedirectResponse("/parametros-logistica", status_code=303)


@app.get("/regras-entrega", response_class=HTMLResponse)
def regras_entrega(request: Request):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        rows = db.execute(select(DeliveryRule, Responsible).outerjoin(Responsible, DeliveryRule.responsible_id == Responsible.id)
                          .where(DeliveryRule.company_id == cid).order_by(DeliveryRule.id.desc())).all()
        responsibles, regions, boxes = _active_logistic_catalogs(db, cid)
    return render(request, "regras_entrega.html", company=company, rows=rows, responsibles=responsibles, regions=regions, boxes=boxes)


@app.post("/regras-entrega")
def criar_regra_entrega(request: Request, customer_cnpj: str = Form(""), customer_name: str = Form(""), region: str = Form(""),
                        responsible_id: int = Form(0), box_type: str = Form(""), priority: str = Form("Normal"),
                        window_start: str = Form(""), window_end: str = Form(""), stop_minutes: int = Form(10), notes: str = Form("")):
    cnpj = "".join(ch for ch in customer_cnpj if ch.isdigit()) or None
    name = customer_name.strip() or None
    if not cnpj and not name:
        return JSONResponse({"error": "informe CNPJ ou cliente"}, status_code=400)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        rule = None
        if cnpj:
            rule = db.scalar(select(DeliveryRule).where(DeliveryRule.company_id == cid, DeliveryRule.customer_cnpj == cnpj).limit(1))
        if not rule and name:
            rule = db.scalar(select(DeliveryRule).where(DeliveryRule.company_id == cid, DeliveryRule.customer_name == name).limit(1))
        if not rule:
            rule = DeliveryRule(company_id=cid, customer_cnpj=cnpj, customer_name=name)
            db.add(rule)
        rule.customer_cnpj, rule.customer_name = cnpj, name
        rule.region = region.strip() or None
        rule.responsible_id = responsible_id or None
        rule.box_type = box_type.strip() or None
        rule.priority = priority.strip() or "Normal"
        rule.window_start = window_start.strip() or None
        rule.window_end = window_end.strip() or None
        rule.stop_minutes = max(0, stop_minutes or 0)
        rule.notes = notes.strip() or None
        rule.active = True
        db.commit()
    return RedirectResponse("/regras-entrega", status_code=303)


@app.post("/regras-entrega/{rule_id}/toggle")
def toggle_regra_entrega(request: Request, rule_id: int):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        row = db.get(DeliveryRule, rule_id)
        if not row or row.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        row.active = not row.active
        db.commit()
    return RedirectResponse("/regras-entrega", status_code=303)


@app.get("/roteiros", response_class=HTMLResponse)
def roteiros(request: Request, responsible_id: int = 0, msg: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        responsibles = db.scalars(select(Responsible).where(
            Responsible.company_id == cid, Responsible.active.is_(True)
        ).order_by(Responsible.name)).all()
        pending_counts = dict(db.execute(select(Responsible.id, func.count(Delivery.id))
            .join(DeliveryAssignment, DeliveryAssignment.responsible_id == Responsible.id)
            .join(Delivery, Delivery.id == DeliveryAssignment.delivery_id)
            .where(Responsible.company_id == cid, Delivery.status.in_(["PENDENTE", "REPROGRAMADA"]))
            .group_by(Responsible.id)).all())

        stmt = (select(Route, Responsible, RouteStop, Delivery, SalesInvoice, Partner)
            .join(Responsible, Route.responsible_id == Responsible.id)
            .join(RouteStop, RouteStop.route_id == Route.id)
            .join(Delivery, RouteStop.delivery_id == Delivery.id)
            .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
            .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
            .where(Route.company_id == cid, Route.status.in_(["PLANEJADO", "EM ROTA"])))
        if responsible_id:
            stmt = stmt.where(Route.responsible_id == responsible_id)
        active_rows = db.execute(stmt.order_by(Responsible.name, Route.route_date.desc(), Route.id.desc(), RouteStop.stop_order)).all()

        # Cards consideram todos os roteiros ativos, mesmo quando a tela está filtrada.
        card_counts = dict(db.execute(
            select(Route.responsible_id, func.count(RouteStop.id))
            .join(RouteStop, RouteStop.route_id == Route.id)
            .where(Route.company_id == cid, Route.status.in_(["PLANEJADO", "EM ROTA"]))
            .group_by(Route.responsible_id)
        ).all())
        driver_cards = [r for r in responsibles if card_counts.get(r.id, 0) > 0]

        groups_map: dict[int, dict] = {}
        for route, resp, stop, delivery, invoice, partner in active_rows:
            group = groups_map.setdefault(resp.id, {
                "responsible": resp, "routes": {}, "delivery_count": 0,
                "total_volumes": 0.0, "total_value": 0.0,
            })
            route_data = group["routes"].setdefault(route.id, {"route": route, "rows": []})
            route_data["rows"].append((stop, delivery, invoice, partner))
            group["delivery_count"] += 1
            group["total_volumes"] += (delivery.volumes or invoice.volumes or 0)
            group["total_value"] += (invoice.total or 0)
        route_groups = []
        for group in groups_map.values():
            group["routes"] = list(group["routes"].values())
            route_groups.append(group)

        finalized_route_rows = db.execute(
            select(Route, Responsible, func.count(RouteStop.id))
            .join(Responsible, Route.responsible_id == Responsible.id)
            .outerjoin(RouteStop, RouteStop.route_id == Route.id)
            .where(Route.company_id == cid, Route.status == "FINALIZADO")
            .group_by(Route.id, Responsible.id)
            .order_by(Route.id.desc()).limit(40)
        ).all()
    return render(
        request, "roteiros.html", company=company, responsibles=responsibles, pending_counts=pending_counts,
        driver_cards=driver_cards, card_counts=card_counts, route_groups=route_groups,
        finalized_route_rows=finalized_route_rows, responsible_id=responsible_id, msg=msg,
        today=date.today().isoformat()
    )


@app.post("/roteiros/trocar-motorista-lote")
def trocar_motorista_lote(
    request: Request, delivery_ids: list[int] = Form([]), responsible_id: int = Form(0),
    return_responsible_id: int = Form(0)
):
    return_base = f"/roteiros?responsible_id={return_responsible_id}" if return_responsible_id else "/roteiros"
    if not delivery_ids:
        sep = "&" if "?" in return_base else "?"
        return RedirectResponse(f"{return_base}{sep}msg=selecione", status_code=303)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        target_resp = db.get(Responsible, responsible_id) if responsible_id else None
        if not target_resp or target_resp.company_id != cid or not target_resp.active:
            sep = "&" if "?" in return_base else "?"
            return RedirectResponse(f"{return_base}{sep}msg=selecione-motorista", status_code=303)

        rows = db.execute(
            select(RouteStop, Route, Delivery)
            .join(Route, RouteStop.route_id == Route.id)
            .join(Delivery, RouteStop.delivery_id == Delivery.id)
            .where(
                Route.company_id == cid, Route.status.in_(["PLANEJADO", "EM ROTA"]),
                Delivery.id.in_(delivery_ids),
            )
            .order_by(Route.id, RouteStop.stop_order)
        ).all()
        if not rows:
            sep = "&" if "?" in return_base else "?"
            return RedirectResponse(f"{return_base}{sep}msg=selecione", status_code=303)

        target_cache: dict[tuple[str, str], Route] = {}
        next_order_cache: dict[int, int] = {}
        touched_sources: dict[int, Route] = {}
        moved = 0
        for stop, source, delivery in rows:
            if source.responsible_id == target_resp.id:
                _assign_delivery(db, cid, delivery, target_resp.id)
                continue
            key = (source.route_date, source.status)
            target = target_cache.get(key)
            if not target:
                target = db.scalar(select(Route).where(
                    Route.company_id == cid, Route.responsible_id == target_resp.id,
                    Route.route_date == source.route_date, Route.status == source.status
                ).order_by(Route.id.desc()).limit(1))
                if not target:
                    target = Route(
                        company_id=cid, responsible_id=target_resp.id, route_date=source.route_date,
                        status=source.status,
                        started_at=(source.started_at or datetime.utcnow()) if source.status == "EM ROTA" else None,
                    )
                    db.add(target); db.flush()
                    db.add(RouteCost(
                        company_id=cid, route_id=target.id, charge_type=target_resp.default_freight_type,
                        freight_value=target_resp.default_freight_value or 0,
                        freight_percent=target_resp.default_freight_percent,
                    ))
                target_cache[key] = target
                next_order_cache[target.id] = (db.scalar(select(func.max(RouteStop.stop_order)).where(RouteStop.route_id == target.id)) or 0) + 1
            next_order = next_order_cache[target.id]
            touched_sources[source.id] = source
            db.delete(stop)
            db.flush()
            db.add(RouteStop(route_id=target.id, delivery_id=delivery.id, stop_order=next_order))
            next_order_cache[target.id] = next_order + 1
            _assign_delivery(db, cid, delivery, target_resp.id)
            db.add(DeliveryMovement(
                company_id=cid, delivery_id=delivery.id, movement_type="TROCA_MOTORISTA",
                old_status=delivery.status, new_status=delivery.status,
                note=f"Roteiro #{source.id} → #{target.id}; novo responsável: {target_resp.name}",
            ))
            moved += 1

        for source in touched_sources.values():
            _renumber_route(db, source.id)
            remaining = db.scalar(select(func.count(RouteStop.id)).where(RouteStop.route_id == source.id)) or 0
            if remaining == 0:
                if source.status == "PLANEJADO":
                    source_cost = db.scalar(select(RouteCost).where(RouteCost.route_id == source.id).limit(1))
                    if source_cost:
                        db.delete(source_cost)
                    db.delete(source)
                elif source.status == "EM ROTA":
                    source.status = "FINALIZADO"
                    source.finished_at = datetime.utcnow()
        db.commit()
    sep = "&" if "?" in return_base else "?"
    return RedirectResponse(f"{return_base}{sep}msg=motorista-{moved}", status_code=303)


@app.post("/roteiros/criar")
def criar_roteiro(request: Request, responsible_id: int = Form(...), route_date: str = Form(...)):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        resp = db.get(Responsible, responsible_id)
        if not resp or resp.company_id != cid or not resp.active:
            return JSONResponse({"error": "responsável inválido"}, status_code=400)
        active_route_deliveries = (select(RouteStop.delivery_id)
            .join(Route, RouteStop.route_id == Route.id)
            .where(Route.company_id == cid, Route.status.in_(["PLANEJADO", "EM ROTA"])))
        deliveries = db.scalars(select(Delivery)
            .join(DeliveryAssignment, DeliveryAssignment.delivery_id == Delivery.id)
            .where(Delivery.company_id == cid, DeliveryAssignment.responsible_id == responsible_id,
                   Delivery.status.in_(["PENDENTE", "REPROGRAMADA"]),
                   ~Delivery.id.in_(active_route_deliveries))
            .order_by(Delivery.region, Delivery.id)).all()
        if not deliveries:
            return RedirectResponse("/roteiros?msg=sem_entregas", status_code=303)
        route = db.scalar(select(Route).where(
            Route.company_id == cid, Route.responsible_id == responsible_id,
            Route.route_date == route_date, Route.status == "PLANEJADO"
        ).order_by(Route.id.desc()).limit(1))
        if not route:
            route = Route(company_id=cid, responsible_id=responsible_id, route_date=route_date, status="PLANEJADO")
            db.add(route); db.flush()
            db.add(RouteCost(company_id=cid, route_id=route.id, charge_type=resp.default_freight_type,
                             freight_value=resp.default_freight_value or 0, freight_percent=resp.default_freight_percent))
        next_order = (db.scalar(select(func.max(RouteStop.stop_order)).where(RouteStop.route_id == route.id)) or 0) + 1
        for offset, delivery in enumerate(deliveries):
            order = next_order + offset
            db.add(RouteStop(route_id=route.id, delivery_id=delivery.id, stop_order=order))
            db.add(DeliveryMovement(company_id=cid, delivery_id=delivery.id, movement_type="ROTEIRIZADA",
                                    old_status=delivery.status, new_status=delivery.status,
                                    note=f"Roteiro #{route.id} - ordem {order}"))
        db.commit()
    return RedirectResponse(f"/roteiros/{route.id}", status_code=303)


@app.get("/roteiros/{route_id}", response_class=HTMLResponse)
def roteiro_detalhe(request: Request, route_id: int):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        route = db.get(Route, route_id)
        if not route or route.company_id != cid:
            return HTMLResponse("Roteiro não encontrado", status_code=404)
        resp = db.get(Responsible, route.responsible_id)
        rows = db.execute(select(RouteStop, Delivery, SalesInvoice, Partner)
            .join(Delivery, RouteStop.delivery_id == Delivery.id)
            .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
            .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
            .where(RouteStop.route_id == route.id).order_by(RouteStop.stop_order)).all()
        cost = db.scalar(select(RouteCost).where(RouteCost.route_id == route.id).limit(1))
        total_value = sum((inv.total or 0) for _, _, inv, _ in rows)
        total_volumes = sum((d.volumes or inv.volumes or 0) for _, d, inv, _ in rows)
        responsibles = db.scalars(select(Responsible).where(Responsible.company_id == cid, Responsible.active.is_(True)).order_by(Responsible.name)).all()
    return render(request, "roteiro_detalhe.html", company=company, route=route, responsible=resp, rows=rows,
                  cost=cost, total_value=total_value, total_volumes=total_volumes, responsibles=responsibles)


@app.post("/roteiros/{route_id}/trocar-motorista")
def trocar_motorista_roteiro(request: Request, route_id: int, delivery_ids: list[int] = Form([]), responsible_id: int = Form(0)):
    if not delivery_ids:
        return RedirectResponse(f"/roteiros/{route_id}?msg=selecione", status_code=303)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        source = db.get(Route, route_id)
        if not source or source.company_id != cid:
            return JSONResponse({"error": "roteiro não encontrado"}, status_code=404)
        if source.status == "FINALIZADO":
            return RedirectResponse(f"/roteiros/{route_id}?msg=finalizado", status_code=303)
        target_resp = db.get(Responsible, responsible_id) if responsible_id else None
        if not target_resp or target_resp.company_id != cid or not target_resp.active:
            return RedirectResponse(f"/roteiros/{route_id}?msg=selecione-motorista", status_code=303)

        valid_stops = db.scalars(
            select(RouteStop).where(RouteStop.route_id == route_id, RouteStop.delivery_id.in_(delivery_ids)).order_by(RouteStop.stop_order)
        ).all()
        if not valid_stops:
            return RedirectResponse(f"/roteiros/{route_id}?msg=selecione", status_code=303)

        # Se o motorista escolhido é o atual, apenas garante o vínculo individual.
        if target_resp.id == source.responsible_id:
            for stop in valid_stops:
                delivery = db.get(Delivery, stop.delivery_id)
                if delivery:
                    _assign_delivery(db, cid, delivery, target_resp.id)
            db.commit()
            return RedirectResponse(f"/roteiros/{route_id}?msg=motorista-{len(valid_stops)}", status_code=303)

        target = db.scalar(
            select(Route).where(
                Route.company_id == cid,
                Route.responsible_id == target_resp.id,
                Route.route_date == source.route_date,
                Route.status == source.status,
            ).order_by(Route.id.desc()).limit(1)
        )
        if not target:
            target = Route(
                company_id=cid,
                responsible_id=target_resp.id,
                route_date=source.route_date,
                status=source.status,
                started_at=(source.started_at or datetime.utcnow()) if source.status == "EM ROTA" else None,
            )
            db.add(target); db.flush()
            db.add(RouteCost(
                company_id=cid, route_id=target.id,
                charge_type=target_resp.default_freight_type,
                freight_value=target_resp.default_freight_value or 0,
                freight_percent=target_resp.default_freight_percent,
            ))

        next_order = (db.scalar(select(func.max(RouteStop.stop_order)).where(RouteStop.route_id == target.id)) or 0) + 1
        moved = 0
        for stop in valid_stops:
            delivery = db.get(Delivery, stop.delivery_id)
            if not delivery:
                continue
            db.delete(stop)
            db.flush()
            db.add(RouteStop(route_id=target.id, delivery_id=delivery.id, stop_order=next_order))
            next_order += 1
            _assign_delivery(db, cid, delivery, target_resp.id)
            db.add(DeliveryMovement(
                company_id=cid, delivery_id=delivery.id, movement_type="TROCA_MOTORISTA",
                old_status=delivery.status, new_status=delivery.status,
                note=f"Roteiro #{source.id} → #{target.id}; novo responsável: {target_resp.name}",
            ))
            moved += 1

        _renumber_route(db, source.id)
        remaining = db.scalar(select(func.count(RouteStop.id)).where(RouteStop.route_id == source.id)) or 0
        source_removed = False
        if remaining == 0:
            if source.status == "PLANEJADO":
                source_cost = db.scalar(select(RouteCost).where(RouteCost.route_id == source.id).limit(1))
                if source_cost:
                    db.delete(source_cost)
                db.delete(source)
                source_removed = True
            elif source.status == "EM ROTA":
                source.status = "FINALIZADO"
                source.finished_at = datetime.utcnow()
        db.commit()
        redirect_id = target.id if source_removed else route_id
    return RedirectResponse(f"/roteiros/{redirect_id}?msg=motorista-{moved}", status_code=303)


@app.post("/roteiros/{route_id}/iniciar")
def iniciar_roteiro(request: Request, route_id: int):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        route = db.get(Route, route_id)
        if not route or route.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        route.status = "EM ROTA"; route.started_at = datetime.utcnow()
        deliveries = db.scalars(select(Delivery).join(RouteStop, RouteStop.delivery_id == Delivery.id).where(RouteStop.route_id == route.id)).all()
        for delivery in deliveries:
            if delivery.status not in {"ENTREGUE", "NAO ENTREGUE"}:
                _set_delivery_status(db, cid, delivery, "EM ROTA", f"Roteiro #{route.id} iniciado", "INICIO_ROTA")
        db.commit()
    return RedirectResponse(f"/roteiros/{route_id}", status_code=303)


@app.post("/roteiros/{route_id}/finalizar")
def finalizar_roteiro(request: Request, route_id: int):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        route = db.get(Route, route_id)
        if not route or route.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        route.status = "FINALIZADO"; route.finished_at = datetime.utcnow()
        db.commit()
    return RedirectResponse(f"/roteiros/{route_id}", status_code=303)


@app.get("/baixa-entregas", response_class=HTMLResponse)
def baixa_entregas(request: Request, status: str = "", msg: str = ""):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        active_route_rows = db.execute(
            select(RouteStop.delivery_id, Route)
            .join(Route, RouteStop.route_id == Route.id)
            .where(Route.company_id == cid, Route.status.in_(["PLANEJADO", "EM ROTA"]))
            .order_by(Route.id.desc())
        ).all()
        route_by_delivery = {}
        for delivery_id, route in active_route_rows:
            route_by_delivery.setdefault(delivery_id, route)
        base_rows = _delivery_rows(db, cid, status=status, limit=2000)
        rows = [row for row in base_rows if row[0].id in route_by_delivery]
    return render(request, "baixa_entregas.html", company=company, rows=rows, status=status, msg=msg,
                  route_by_delivery=route_by_delivery)


@app.post("/baixa-entregas/lote")
def baixa_entregas_lote(request: Request, delivery_ids: list[int] = Form([]), action: str = Form(""), note: str = Form("")):
    if not delivery_ids:
        return RedirectResponse("/baixa-entregas?msg=selecione", status_code=303)
    status_map = {"ENTREGUE": "ENTREGUE", "REPROGRAMAR": "REPROGRAMADA"}
    new_status = status_map.get(action)
    if not new_status:
        return RedirectResponse("/baixa-entregas?msg=acao-invalida", status_code=303)
    changed = 0
    skipped = 0
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        deliveries = db.scalars(select(Delivery).where(Delivery.company_id == cid, Delivery.id.in_(delivery_ids))).all()
        for delivery in deliveries:
            if _apply_delivery_closeout(db, cid, delivery, new_status, note, "Ocorrência"):
                changed += 1
            else:
                skipped += 1
        db.commit()
    return RedirectResponse(f"/baixa-entregas?msg=lote-{changed}-{skipped}", status_code=303)


@app.post("/entregas/{delivery_id}/status")
def alterar_status_entrega(request: Request, delivery_id: int, new_status: str = Form(...), note: str = Form(""), occurrence_type: str = Form("Ocorrência")):
    allowed = {"ENTREGUE", "NAO ENTREGUE", "OCORRENCIA", "REPROGRAMADA"}
    if new_status not in allowed:
        return JSONResponse({"error": "status inválido"}, status_code=400)
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        delivery = db.get(Delivery, delivery_id)
        if not delivery or delivery.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        if not _apply_delivery_closeout(db, cid, delivery, new_status, note, occurrence_type):
            return RedirectResponse("/baixa-entregas?msg=nao-roteirizada", status_code=303)
        db.commit()
    return RedirectResponse("/baixa-entregas?msg=ok", status_code=303)

@app.get("/ocorrencias", response_class=HTMLResponse)
def ocorrencias(request: Request, resolved: str = "0"):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        stmt = (select(DeliveryOccurrence, Delivery, SalesInvoice, Partner, Responsible)
                .join(Delivery, DeliveryOccurrence.delivery_id == Delivery.id)
                .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
                .outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
                .outerjoin(Responsible, DeliveryOccurrence.responsible_id == Responsible.id)
                .where(DeliveryOccurrence.company_id == cid))
        if resolved != "all":
            stmt = stmt.where(DeliveryOccurrence.resolved.is_(resolved == "1"))
        rows = db.execute(stmt.order_by(DeliveryOccurrence.id.desc()).limit(500)).all()
    return render(request, "ocorrencias.html", company=company, rows=rows, resolved=resolved)


@app.post("/ocorrencias/{occurrence_id}/resolver")
def resolver_ocorrencia(request: Request, occurrence_id: int):
    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        occ = db.get(DeliveryOccurrence, occurrence_id)
        if not occ or occ.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        occ.resolved = True; occ.resolved_at = datetime.utcnow()
        db.commit()
    return RedirectResponse("/ocorrencias", status_code=303)


@app.get("/custos-roteiro", response_class=HTMLResponse)
def custos_roteiro(request: Request):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        rows = db.execute(select(Route, Responsible, RouteCost)
            .join(Responsible, Route.responsible_id == Responsible.id)
            .outerjoin(RouteCost, RouteCost.route_id == Route.id)
            .where(Route.company_id == cid).order_by(Route.id.desc()).limit(200)).all()
    return render(request, "custos_roteiro.html", company=company, rows=rows)


@app.post("/custos-roteiro/{route_id}")
def salvar_custo_roteiro(request: Request, route_id: int, charge_type: str = Form("Valor fechado"), freight_value: float = Form(0),
                         freight_percent: str = Form(""), km_initial: str = Form(""),
                         km_final: str = Form(""), notes: str = Form("")):
    def _opt_float(value):
        txt = str(value or "").strip().replace(",", ".")
        if not txt:
            return None
        try:
            return float(txt)
        except Exception:
            return None
    freight_percent_value = _opt_float(freight_percent)
    km_initial_value = _opt_float(km_initial)
    km_final_value = _opt_float(km_final)

    with SessionLocal() as db:
        cid, _company = company_context(db, request)
        route = db.get(Route, route_id)
        if not route or route.company_id != cid:
            return JSONResponse({"error": "não encontrado"}, status_code=404)
        cost = db.scalar(select(RouteCost).where(RouteCost.route_id == route.id).limit(1))
        if not cost:
            cost = RouteCost(company_id=cid, route_id=route.id)
            db.add(cost)
        cost.charge_type = charge_type
        cost.freight_percent = freight_percent_value
        cost.km_initial = km_initial_value
        cost.km_final = km_final_value
        cost.km_realized = max(0, (km_final_value - km_initial_value)) if km_initial_value is not None and km_final_value is not None else (cost.km_realized or 0)
        cost.notes = notes.strip() or None
        # Percentual é calculado somente sobre o faturamento transportado por este roteiro.
        route_value = db.scalar(select(func.sum(SalesInvoice.total))
                                .select_from(RouteStop)
                                .join(Delivery, RouteStop.delivery_id == Delivery.id)
                                .join(SalesInvoice, Delivery.sales_invoice_id == SalesInvoice.id)
                                .where(RouteStop.route_id == route.id)) or 0
        if freight_percent_value is not None and ("percent" in (charge_type or "").lower() or "%" in (charge_type or "")):
            cost.freight_value = max(0, float(route_value) * float(freight_percent_value or 0) / 100.0)
        else:
            cost.freight_value = max(0, freight_value or 0)
        # Replica somente os valores logísticos derivados no registro de entrega; a NF continua única.
        deliveries = db.scalars(select(Delivery).join(RouteStop, RouteStop.delivery_id == Delivery.id).where(RouteStop.route_id == route.id)).all()
        per_delivery = (cost.freight_value / len(deliveries)) if deliveries and cost.freight_value else 0
        per_km = (cost.km_realized / len(deliveries)) if deliveries and cost.km_realized else 0
        for delivery in deliveries:
            delivery.freight_cost = per_delivery
            delivery.km = per_km
        db.commit()
    return RedirectResponse("/custos-roteiro", status_code=303)


CADASTRO_COLUMN_DEFS = {
    "produtos": [
        ("code", "Código"), ("description", "Produto"), ("ean", "EAN"), ("ncm", "NCM"),
        ("unit", "Unidade"), ("category", "Categoria"), ("erpflex_id", "ID ERPFlex"), ("updated_at", "Atualizado"),
    ],
    "clientes": [
        ("cnpj_cpf", "CNPJ/CPF"), ("name", "Razão social"), ("trade_name", "Fantasia"),
        ("address", "Endereço"), ("district", "Bairro"), ("city", "Cidade"), ("state", "UF"),
        ("zip_code", "CEP"), ("erpflex_id", "ID ERPFlex"), ("updated_at", "Atualizado"),
    ],
    "fornecedores": [
        ("cnpj_cpf", "CNPJ/CPF"), ("name", "Fornecedor"), ("trade_name", "Fantasia"),
        ("address", "Endereço"), ("district", "Bairro"), ("city", "Cidade"), ("state", "UF"),
        ("zip_code", "CEP"), ("erpflex_id", "ID ERPFlex"), ("updated_at", "Atualizado"),
    ],
    "parceiros": [
        ("cnpj_cpf", "CNPJ/CPF"), ("name", "Razão social"), ("trade_name", "Fantasia"),
        ("roles", "Tipo"), ("city", "Cidade"), ("state", "UF"), ("erpflex_id", "ID ERPFlex"), ("updated_at", "Atualizado"),
    ],
    "bancos": [("external_id", "ID"), ("name", "Banco"), ("code", "Código"), ("updated_at", "Atualizado")],
}

CADASTRO_DEFAULT_COLS = {
    "produtos": ["code", "description", "ean", "ncm", "unit", "category"],
    "clientes": ["cnpj_cpf", "name", "trade_name", "city", "state", "erpflex_id"],
    "fornecedores": ["cnpj_cpf", "name", "trade_name", "city", "state", "erpflex_id"],
    "parceiros": ["cnpj_cpf", "name", "trade_name", "roles", "city", "state"],
    "bancos": ["external_id", "name", "code", "updated_at"],
}


def _cadastro_pref(db, request: Request, company_id: int, tipo: str):
    ac = access_context(db, request)
    if not ac.user:
        return None
    return db.scalar(select(UserViewPreference).where(
        UserViewPreference.user_id == ac.user.id,
        UserViewPreference.company_id == company_id,
        UserViewPreference.view_key == f"cadastros:{tipo}",
    ).limit(1))


def _format_dt(value):
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M")
    return value or "—"


@app.get("/cadastros", response_class=HTMLResponse)
def cadastros(
    request: Request, tipo: str = "produtos", q: str = "", page: int = 1,
    per_page: int = 0, cols: str = "", save_view: int = 0,
):
    tipo = tipo if tipo in CADASTRO_COLUMN_DEFS else "produtos"
    allowed_pp = {25, 50, 100, 250}
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        pref = _cadastro_pref(db, request, cid, tipo)
        default_cols = CADASTRO_DEFAULT_COLS[tipo]
        pref_cols = []
        if pref and pref.columns_json:
            try:
                pref_cols = [x for x in json.loads(pref.columns_json) if x in dict(CADASTRO_COLUMN_DEFS[tipo])]
            except Exception:
                pref_cols = []
        requested_cols = [x for x in cols.split(",") if x in dict(CADASTRO_COLUMN_DEFS[tipo])] if cols else []
        selected_cols = requested_cols or pref_cols or default_cols
        pp = per_page if per_page in allowed_pp else (pref.per_page if pref and pref.per_page in allowed_pp else 50)
        if save_view and access_context(db, request).user:
            ac = access_context(db, request)
            if not pref:
                pref = UserViewPreference(user_id=ac.user.id, company_id=cid, view_key=f"cadastros:{tipo}")
                db.add(pref)
            pref.columns_json = json.dumps(selected_cols, ensure_ascii=False)
            pref.per_page = pp
            db.commit()
        page = max(1, int(page or 1))
        rows_view = []

        if tipo == "produtos":
            stmt = select(Product).where(Product.company_id == cid)
            count_stmt = select(func.count(Product.id)).where(Product.company_id == cid)
            if q.strip():
                like = f"%{q.strip()}%"
                cond = or_(Product.description.ilike(like), Product.code.ilike(like), Product.ean.ilike(like), Product.ncm.ilike(like), Product.category.ilike(like))
                stmt = stmt.where(cond); count_stmt = count_stmt.where(cond)
            total = db.scalar(count_stmt) or 0
            pages = max(1, (total + pp - 1) // pp)
            page = min(page, pages)
            rows = db.scalars(stmt.order_by(Product.description, Product.code).offset((page-1)*pp).limit(pp)).all()
            for r in rows:
                rows_view.append({"_id": r.id, "_name": r.description or r.code or f"Produto {r.id}",
                    "code": r.code or "—", "description": r.description or "—", "ean": r.ean or "—", "ncm": r.ncm or "—",
                    "unit": r.unit or "—", "category": r.category or "—", "erpflex_id": r.erpflex_id or "—", "updated_at": _format_dt(r.updated_at)})
        elif tipo in {"clientes", "fornecedores", "parceiros"}:
            stmt = select(Partner).where(Partner.company_id == cid)
            count_stmt = select(func.count(Partner.id)).where(Partner.company_id == cid)
            if tipo == "clientes":
                stmt = stmt.where(Partner.role_customer.is_(True)); count_stmt = count_stmt.where(Partner.role_customer.is_(True))
            elif tipo == "fornecedores":
                stmt = stmt.where(Partner.role_supplier.is_(True)); count_stmt = count_stmt.where(Partner.role_supplier.is_(True))
            if q.strip():
                like = f"%{q.strip()}%"
                cond = or_(Partner.name.ilike(like), Partner.trade_name.ilike(like), Partner.cnpj_cpf.ilike(like), Partner.city.ilike(like), Partner.erpflex_id.ilike(like))
                stmt = stmt.where(cond); count_stmt = count_stmt.where(cond)
            total = db.scalar(count_stmt) or 0
            pages = max(1, (total + pp - 1) // pp)
            page = min(page, pages)
            rows = db.scalars(stmt.order_by(Partner.name).offset((page-1)*pp).limit(pp)).all()
            for r in rows:
                roles = " / ".join(x for x, ok in (("Cliente", r.role_customer), ("Fornecedor", r.role_supplier)) if ok) or "—"
                rows_view.append({"_id": r.id, "_name": r.name or r.trade_name or f"Parceiro {r.id}",
                    "cnpj_cpf": r.cnpj_cpf or "—", "name": r.name or "—", "trade_name": r.trade_name or "—", "roles": roles,
                    "address": r.address or "—", "district": r.district or "—", "city": r.city or "—", "state": r.state or "—",
                    "zip_code": r.zip_code or "—", "erpflex_id": r.erpflex_id or "—", "updated_at": _format_dt(r.updated_at)})
        else:  # bancos preservados integralmente em raw_records
            stmt = select(RawRecord).where(RawRecord.company_id == cid, RawRecord.source == "ERPFLEX", RawRecord.module == "banks")
            count_stmt = select(func.count(RawRecord.id)).where(RawRecord.company_id == cid, RawRecord.source == "ERPFLEX", RawRecord.module == "banks")
            if q.strip():
                like = f"%{q.strip()}%"
                stmt = stmt.where(RawRecord.payload_json.ilike(like)); count_stmt = count_stmt.where(RawRecord.payload_json.ilike(like))
            total = db.scalar(count_stmt) or 0
            pages = max(1, (total + pp - 1) // pp)
            page = min(page, pages)
            rows = db.scalars(stmt.order_by(RawRecord.id).offset((page-1)*pp).limit(pp)).all()
            for r in rows:
                try: payload = json.loads(r.payload_json or "{}")
                except Exception: payload = {}
                name = str(payload.get("nome") or payload.get("descricao") or payload.get("banco") or payload.get("razao_social") or r.external_id)
                code = str(payload.get("codigo") or payload.get("banco_id") or payload.get("id") or "—")
                rows_view.append({"_id": r.id, "_name": name, "external_id": r.external_id, "name": name, "code": code, "updated_at": _format_dt(r.updated_at)})

        col_map = dict(CADASTRO_COLUMN_DEFS[tipo])
        selected_columns = [(key, col_map[key]) for key in selected_cols if key in col_map]
        all_columns = CADASTRO_COLUMN_DEFS[tipo]
        base = f"/cadastros?tipo={quote_plus(tipo)}&q={quote_plus(q)}&per_page={pp}"
        prev_url = f"{base}&page={page-1}" if page > 1 else None
        next_url = f"{base}&page={page+1}" if page < pages else None
    return render(request, "cadastros.html", company=company, rows=rows_view, tipo=tipo, q=q, page=page, pages=pages,
                  total=total, per_page=pp, selected_columns=selected_columns, selected_keys=selected_cols,
                  all_columns=all_columns, prev_url=prev_url, next_url=next_url)


@app.get("/cadastros/{tipo}/{record_id}", response_class=HTMLResponse)
def cadastro_detalhe(request: Request, tipo: str, record_id: int):
    tipo = tipo if tipo in CADASTRO_COLUMN_DEFS else "produtos"
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        summary = []
        raw_rows = []
        title = "Detalhe"
        if tipo == "produtos":
            row = db.get(Product, record_id)
            if not row or row.company_id != cid:
                return HTMLResponse("Registro não encontrado.", status_code=404)
            title = row.description or row.code or f"Produto {row.id}"
            summary = [("Código", row.code), ("Descrição", row.description), ("EAN", row.ean), ("NCM", row.ncm), ("Unidade", row.unit), ("Categoria", row.category), ("ID ERPFlex", row.erpflex_id)]
            if row.erpflex_id:
                raw_rows = db.scalars(select(RawRecord).where(RawRecord.company_id == cid, RawRecord.source == "ERPFLEX", RawRecord.module == "products", RawRecord.external_id == row.erpflex_id).order_by(RawRecord.id.desc()).limit(5)).all()
        elif tipo in {"clientes", "fornecedores", "parceiros"}:
            row = db.get(Partner, record_id)
            if not row or row.company_id != cid:
                return HTMLResponse("Registro não encontrado.", status_code=404)
            title = row.name or row.trade_name or f"Parceiro {row.id}"
            summary = [("CNPJ/CPF", row.cnpj_cpf), ("Razão social", row.name), ("Fantasia", row.trade_name), ("Endereço", row.address), ("Bairro", row.district), ("Cidade", row.city), ("UF", row.state), ("CEP", row.zip_code), ("ID ERPFlex", row.erpflex_id)]
            conditions = []
            if row.erpflex_id:
                conditions.append(and_(RawRecord.module == "clientes", RawRecord.external_id == row.erpflex_id))
            if row.cnpj_cpf:
                conditions.append(RawRecord.payload_json.ilike(f"%{row.cnpj_cpf}%"))
            if conditions:
                raw_rows = db.scalars(select(RawRecord).where(RawRecord.company_id == cid, RawRecord.source == "ERPFLEX", or_(*conditions)).order_by(RawRecord.id.desc()).limit(12)).all()
        else:
            raw = db.get(RawRecord, record_id)
            if not raw or raw.company_id != cid or raw.module != "banks":
                return HTMLResponse("Registro não encontrado.", status_code=404)
            raw_rows = [raw]
            try: data = json.loads(raw.payload_json or "{}")
            except Exception: data = {}
            title = str(data.get("nome") or data.get("descricao") or data.get("banco") or raw.external_id)
            summary = [("ID", raw.external_id), ("Atualizado", _format_dt(raw.updated_at))]

        payloads = []
        for raw in raw_rows:
            try: data = json.loads(raw.payload_json or "{}")
            except Exception: data = {"raw": raw.payload_json}
            payloads.append({"module": raw.module, "external_id": raw.external_id, "updated_at": _format_dt(raw.updated_at), "data": data, "pretty": json.dumps(data, ensure_ascii=False, indent=2, default=str)})
    return render(request, "cadastro_detalhe.html", company=company, tipo=tipo, title=title, summary=summary, payloads=payloads)


@app.get("/diagnostico", response_class=HTMLResponse)
def diagnostico(request: Request):
    with SessionLocal() as db:
        cid, company = company_context(db, request)
        runs = db.scalars(select(SyncRun).where(SyncRun.company_id == cid).order_by(SyncRun.id.desc()).limit(100)).all()
        raw_counts = db.execute(select(RawRecord.source, RawRecord.module, func.count(RawRecord.id)).where(RawRecord.company_id == cid).group_by(RawRecord.source, RawRecord.module).order_by(RawRecord.source, RawRecord.module)).all()
        integrations = integration_public_view(db, cid)
    return render(request, "diagnostico.html", company=company, runs=runs, raw_counts=raw_counts,
                  erp_configured=integrations["erpflex"]["configured"], nf_configured=integrations["nfstock"]["configured"],
                  integrations=integrations, db_status=database_status(),
                  erpflex_engine=(os.getenv("ERPFLEX_ENGINE") or "go_v78"),
                  erpflex_bridge_exists=(Path(__file__).resolve().parent.parent / "bin" / ("erpflex_v78_bridge_windows.exe" if os.name == "nt" else "erpflex_v78_bridge")).exists())


PERMISSION_CATALOG = [
    ("dashboard.view", "Início / indicadores"), ("sync.manage", "Sincronização ERPFlex / NF Stock"),
    ("cadastros.view", "Cadastros"), ("faturamento.view", "Faturamento"),
    ("financeiro.view", "Financeiro"), ("compras.view", "Compras / NF-e"),
    ("logistica.view", "Consultar logística"), ("logistica.edit", "Editar logística"),
    ("roteirizar.view", "Consultar Roteirizador"), ("roteirizar.manage", "Classificar, vincular e montar roteiros"),
    ("importar.manage", "Importação Excel de contingência"), ("painel_entregas.view", "Painel de Entregas"),
    ("painel_gerencial.view", "Painel Gerencial de Logística"), ("relatorios.view", "Relatórios Gerenciais de Logística"),
    ("roteiros.view", "Consultar roteiros e custos"), ("roteiros.manage", "Criar / iniciar / finalizar roteiros e custos"),
    ("baixa.manage", "Baixa de entregas"), ("ocorrencias.view", "Consultar ocorrências"),
    ("ocorrencias.manage", "Resolver ocorrências"), ("responsaveis.view", "Consultar responsáveis"),
    ("responsaveis.manage", "Gerenciar responsáveis"), ("regras.manage", "Gerenciar regras de entrega"),
    ("parametros.manage", "Gerenciar parâmetros logísticos"), ("diagnostico.view", "Diagnóstico"),
    ("api.read", "Consumir API interna somente leitura"), ("admin.manage", "Administração de usuários, perfis, empresas e filiais"),
]


@app.post("/contexto")
def trocar_contexto(request: Request, company_id: int = Form(...), branch_id: int = Form(0), next_url: str = Form("/")):
    with SessionLocal() as db:
        ac = access_context(db, request)
        allowed_companies = {c.id for c in ac.companies}
        if company_id not in allowed_companies:
            return JSONResponse({"error": "empresa não autorizada"}, status_code=403)
        request.session["company_id"] = company_id
        # Recalcula filiais após trocar empresa.
        ac2 = access_context(db, request)
        allowed_branches = {b.id for b in ac2.branches}
        if branch_id and branch_id in allowed_branches:
            request.session["branch_id"] = branch_id
        elif ac2.branches:
            request.session["branch_id"] = ac2.branches[0].id
        add_audit(db, ac2, "CONTEXTO_ALTERADO", "company", company_id, f"Filial {request.session.get('branch_id')}")
        db.commit()
    return RedirectResponse(next_url if next_url.startswith("/") else "/", status_code=303)


@app.get("/administracao", response_class=HTMLResponse)
def administracao(request: Request):
    with SessionLocal() as db:
        ac = access_context(db, request)
        allowed_company_ids = {c.id for c in ac.companies}
        companies = db.scalars(select(Company).where(Company.id.in_(allowed_company_ids)).order_by(Company.name)).all()
        branches = db.execute(select(Branch, Company).join(Company, Branch.company_id == Company.id)
                              .where(Company.id.in_(allowed_company_ids)).order_by(Company.name, Branch.name)).all()
        profiles = db.scalars(select(AccessProfile).where(AccessProfile.active.is_(True)).order_by(AccessProfile.name)).all()
        profile_permissions = {p.id: set(db.scalars(select(ProfilePermission.permission).where(ProfilePermission.profile_id == p.id)).all()) for p in profiles}
        if ac.user and ac.user.is_superuser:
            users = db.scalars(select(AppUser).order_by(AppUser.username)).all()
        else:
            users = db.scalars(select(AppUser).join(UserCompanyAccess, UserCompanyAccess.user_id == AppUser.id)
                               .where(UserCompanyAccess.company_id.in_(allowed_company_ids)).distinct().order_by(AppUser.username)).all()
        user_access = {}
        user_branches = {}
        for u in users:
            user_access[u.id] = db.execute(
                select(UserCompanyAccess, Company, AccessProfile)
                .join(Company, UserCompanyAccess.company_id == Company.id)
                .join(AccessProfile, UserCompanyAccess.profile_id == AccessProfile.id)
                .where(UserCompanyAccess.user_id == u.id, UserCompanyAccess.company_id.in_(allowed_company_ids))
            ).all()
            user_branches[u.id] = db.execute(
                select(UserBranchAccess, Branch).join(Branch, UserBranchAccess.branch_id == Branch.id)
                .where(UserBranchAccess.user_id == u.id, UserBranchAccess.active.is_(True), Branch.company_id.in_(allowed_company_ids))
            ).all()
        audit = db.execute(select(AuditLog, AppUser).outerjoin(AppUser, AuditLog.user_id == AppUser.id)
                           .where(AuditLog.company_id.in_(allowed_company_ids)).order_by(AuditLog.id.desc()).limit(80)).all()
    return render(request, "administracao.html", companies=companies, branches=branches, profiles=profiles,
                  profile_permissions=profile_permissions, users=users, user_access=user_access, user_branches=user_branches,
                  permission_catalog=PERMISSION_CATALOG, audit=audit)


@app.post("/administracao/empresa")
def admin_empresa(request: Request, name: str = Form(...), trade_name: str = Form(""), cnpj: str = Form("")):
    clean = name.strip()
    if not clean:
        return JSONResponse({"error": "nome obrigatório"}, status_code=400)
    with SessionLocal() as db:
        ac = access_context(db, request)
        if not ac.user or not ac.user.is_superuser:
            return JSONResponse({"error": "somente superusuário pode criar empresas"}, status_code=403)
        company = Company(name=clean, trade_name=trade_name.strip() or None, cnpj="".join(ch for ch in cnpj if ch.isdigit()) or None)
        db.add(company); db.flush()
        branch = Branch(company_id=company.id, name="Matriz", code="MATRIZ", cnpj=company.cnpj)
        db.add(branch); db.flush()
        default_regions = [
            ("Outras Regiões", 10), ("COLETA", 20), ("A Classificar", 30), ("Zona Oeste", 40),
            ("Centro", 50), ("Zona Sul", 60), ("Osasco / Barueri", 70), ("Zona Norte", 80),
            ("Zona Leste", 90), ("ABC Paulista", 100), ("Guarulhos", 110), ("Alto Tietê", 120),
            ("Interior SP", 130), ("Litoral SP", 140),
        ]
        for region_name, order in default_regions:
            db.add(Region(company_id=company.id, name=region_name, sort_order=order))
        db.add(BoxType(company_id=company.id, name="Papelão", sort_order=10))
        db.add(BoxType(company_id=company.id, name="Caixa Vermelha", sort_order=20))
        db.add(Responsible(company_id=company.id, name="COLETA", kind="COLETA", active=True))
        if ac.user:
            admin_profile = db.scalar(select(AccessProfile).where(AccessProfile.name == "Administrador").limit(1))
            db.add(UserCompanyAccess(user_id=ac.user.id, company_id=company.id, profile_id=admin_profile.id))
            db.add(UserBranchAccess(user_id=ac.user.id, branch_id=branch.id))
        add_audit(db, ac, "EMPRESA_CRIADA", "company", company.id, clean)
        db.commit()
    return RedirectResponse("/administracao", status_code=303)


@app.post("/administracao/filial")
def admin_filial(request: Request, company_id: int = Form(...), name: str = Form(...), code: str = Form(""), cnpj: str = Form("")):
    with SessionLocal() as db:
        ac = access_context(db, request)
        company = db.get(Company, company_id)
        if company_id not in {c.id for c in ac.companies}:
            return JSONResponse({"error": "empresa não autorizada"}, status_code=403)
        if not company or not name.strip():
            return JSONResponse({"error": "dados inválidos"}, status_code=400)
        branch = Branch(company_id=company_id, name=name.strip(), code=code.strip() or None, cnpj="".join(ch for ch in cnpj if ch.isdigit()) or None)
        db.add(branch); db.flush()
        if ac.user and ac.user.is_superuser:
            db.add(UserBranchAccess(user_id=ac.user.id, branch_id=branch.id))
        add_audit(db, ac, "FILIAL_CRIADA", "branch", branch.id, f"{company.name} / {branch.name}")
        db.commit()
    return RedirectResponse("/administracao", status_code=303)


@app.post("/administracao/perfil")
def admin_perfil(request: Request, name: str = Form(...), description: str = Form(""), permissions: list[str] = Form([])):
    valid = {p for p, _ in PERMISSION_CATALOG}
    chosen = [p for p in permissions if p in valid]
    with SessionLocal() as db:
        ac = access_context(db, request)
        profile = db.scalar(select(AccessProfile).where(func.lower(AccessProfile.name) == name.strip().lower()).limit(1))
        if not profile:
            profile = AccessProfile(name=name.strip(), description=description.strip() or None, system=False)
            db.add(profile); db.flush()
        elif profile.system:
            return JSONResponse({"error": "perfil de sistema não pode ser sobrescrito por este formulário"}, status_code=400)
        else:
            profile.description = description.strip() or None
            db.query(ProfilePermission).filter(ProfilePermission.profile_id == profile.id).delete(synchronize_session=False)
        for perm in chosen:
            db.add(ProfilePermission(profile_id=profile.id, permission=perm))
        add_audit(db, ac, "PERFIL_SALVO", "profile", profile.id, profile.name)
        db.commit()
    return RedirectResponse("/administracao", status_code=303)


@app.post("/administracao/usuario")
def admin_usuario(request: Request, username: str = Form(...), name: str = Form(""), email: str = Form(""),
                  password: str = Form(...), company_id: int = Form(...), profile_id: int = Form(...), branch_ids: list[int] = Form([])):
    if not username.strip() or len(password) < 6:
        return JSONResponse({"error": "usuário obrigatório e senha com pelo menos 6 caracteres"}, status_code=400)
    with SessionLocal() as db:
        ac = access_context(db, request)
        if db.scalar(select(AppUser).where(func.lower(AppUser.username) == username.strip().lower()).limit(1)):
            return JSONResponse({"error": "usuário já existe"}, status_code=400)
        profile = db.get(AccessProfile, profile_id); company = db.get(Company, company_id)
        if company_id not in {c.id for c in ac.companies}:
            return JSONResponse({"error": "empresa não autorizada"}, status_code=403)
        if not profile or not company:
            return JSONResponse({"error": "empresa/perfil inválido"}, status_code=400)
        user = AppUser(username=username.strip(), name=name.strip() or username.strip(), email=email.strip() or None, password_hash=hash_password(password))
        db.add(user); db.flush()
        db.add(UserCompanyAccess(user_id=user.id, company_id=company_id, profile_id=profile_id))
        allowed_branches = db.scalars(select(Branch).where(Branch.company_id == company_id, Branch.id.in_(branch_ids))).all() if branch_ids else []
        if not allowed_branches:
            first = db.scalar(select(Branch).where(Branch.company_id == company_id, Branch.active.is_(True)).order_by(Branch.id).limit(1))
            allowed_branches = [first] if first else []
        for branch in allowed_branches:
            db.add(UserBranchAccess(user_id=user.id, branch_id=branch.id))
        add_audit(db, ac, "USUARIO_CRIADO", "user", user.id, user.username)
        db.commit()
    return RedirectResponse("/administracao", status_code=303)


@app.post("/administracao/usuario/{user_id}/acesso")
def admin_usuario_acesso(request: Request, user_id: int, company_id: int = Form(...), profile_id: int = Form(...), branch_ids: list[int] = Form([])):
    with SessionLocal() as db:
        ac = access_context(db, request)
        user = db.get(AppUser, user_id); profile = db.get(AccessProfile, profile_id); company = db.get(Company, company_id)
        if company_id not in {c.id for c in ac.companies}:
            return JSONResponse({"error": "empresa não autorizada"}, status_code=403)
        if not user or not profile or not company:
            return JSONResponse({"error": "dados inválidos"}, status_code=400)
        access = db.scalar(select(UserCompanyAccess).where(UserCompanyAccess.user_id == user_id, UserCompanyAccess.company_id == company_id).limit(1))
        if not access:
            db.add(UserCompanyAccess(user_id=user_id, company_id=company_id, profile_id=profile_id))
        else:
            access.profile_id = profile_id; access.active = True
        valid = set(db.scalars(select(Branch.id).where(Branch.company_id == company_id, Branch.id.in_(branch_ids))).all()) if branch_ids else set()
        existing = db.scalars(select(UserBranchAccess).join(Branch, UserBranchAccess.branch_id == Branch.id).where(UserBranchAccess.user_id == user_id, Branch.company_id == company_id)).all()
        for ba in existing:
            ba.active = ba.branch_id in valid
        existing_ids = {ba.branch_id for ba in existing}
        for bid in valid - existing_ids:
            db.add(UserBranchAccess(user_id=user_id, branch_id=bid, active=True))
        add_audit(db, ac, "ACESSO_USUARIO_ATUALIZADO", "user", user_id, f"empresa={company_id}; perfil={profile.name}")
        db.commit()
    return RedirectResponse("/administracao", status_code=303)


@app.post("/administracao/usuario/{user_id}/toggle")
def admin_usuario_toggle(request: Request, user_id: int):
    with SessionLocal() as db:
        ac = access_context(db, request)
        user = db.get(AppUser, user_id)
        if not user:
            return JSONResponse({"error": "usuário não encontrado"}, status_code=404)
        if not (ac.user and ac.user.is_superuser):
            allowed_ids = {c.id for c in ac.companies}
            visible = db.scalar(select(UserCompanyAccess.id).where(UserCompanyAccess.user_id == user_id, UserCompanyAccess.company_id.in_(allowed_ids)).limit(1))
            if not visible:
                return JSONResponse({"error": "usuário não autorizado"}, status_code=403)
        if ac.user and ac.user.id == user.id:
            return JSONResponse({"error": "não é permitido desativar o próprio usuário"}, status_code=400)
        user.active = not user.active
        add_audit(db, ac, "USUARIO_STATUS", "user", user.id, f"ativo={user.active}")
        db.commit()
    return RedirectResponse("/administracao", status_code=303)


@app.post("/administracao/usuario/{user_id}/senha")
def admin_usuario_senha(request: Request, user_id: int, new_password: str = Form(...)):
    if len(new_password) < 6:
        return JSONResponse({"error": "a senha deve ter pelo menos 6 caracteres"}, status_code=400)
    with SessionLocal() as db:
        ac = access_context(db, request)
        user = db.get(AppUser, user_id)
        if not user:
            return JSONResponse({"error": "usuário não encontrado"}, status_code=404)
        if not (ac.user and ac.user.is_superuser):
            allowed_ids = {c.id for c in ac.companies}
            visible = db.scalar(select(UserCompanyAccess.id).where(UserCompanyAccess.user_id == user_id, UserCompanyAccess.company_id.in_(allowed_ids)).limit(1))
            if not visible:
                return JSONResponse({"error": "usuário não autorizado"}, status_code=403)
        user.password_hash = hash_password(new_password)
        add_audit(db, ac, "SENHA_REDEFINIDA", "user", user.id, user.username)
        db.commit()
    return RedirectResponse("/administracao", status_code=303)


@app.get("/minha-conta", response_class=HTMLResponse)
def minha_conta(request: Request):
    with SessionLocal() as db:
        ac = access_context(db, request)
    return render(request, "minha_conta.html", account=ac.user)


@app.post("/minha-conta/senha")
def minha_conta_senha(request: Request, current_password: str = Form(...), new_password: str = Form(...), confirm_password: str = Form(...)):
    if new_password != confirm_password or len(new_password) < 6:
        return RedirectResponse("/minha-conta?erro=senha", status_code=303)
    with SessionLocal() as db:
        ac = access_context(db, request)
        user = db.get(AppUser, ac.user.id) if ac.user else None
        if not user or not verify_password(current_password, user.password_hash):
            return RedirectResponse("/minha-conta?erro=atual", status_code=303)
        user.password_hash = hash_password(new_password)
        add_audit(db, ac, "SENHA_ALTERADA", "user", user.id, user.username)
        db.commit()
    return RedirectResponse("/minha-conta?ok=1", status_code=303)


@app.get("/api/v1/context")
def api_context(request: Request):
    with SessionLocal() as db:
        ac = access_context(db, request)
        return {
            "user": {"id": ac.user.id, "username": ac.user.username, "name": ac.user.name} if ac.user else None,
            "company": {"id": ac.company.id, "name": ac.company.name, "trade_name": ac.company.trade_name, "cnpj": ac.company.cnpj},
            "branch": {"id": ac.branch.id, "name": ac.branch.name, "code": ac.branch.code, "cnpj": ac.branch.cnpj} if ac.branch else None,
            "profile": ac.profile.name if ac.profile else ("Administrador" if ac.user and ac.user.is_superuser else None),
            "permissions": sorted(ac.permissions),
        }


@app.get("/api/v1/faturamento")
def api_faturamento(request: Request, limit: int = 100):
    limit = max(1, min(500, limit))
    with SessionLocal() as db:
        cid, _ = company_context(db, request)
        rows = db.execute(select(SalesInvoice, Partner).outerjoin(Partner, SalesInvoice.customer_id == Partner.id)
                          .where(SalesInvoice.company_id == cid).order_by(SalesInvoice.id.desc()).limit(limit)).all()
        return {"items": [{"id": n.id, "nf": n.number, "chave": n.access_key, "cliente": p.name if p else None,
                            "emissao": n.issue_date, "valor": n.total, "transportadora": n.carrier, "volumes": n.volumes} for n, p in rows]}


@app.get("/api/v1/compras")
def api_compras(request: Request, limit: int = 100):
    limit = max(1, min(500, limit))
    with SessionLocal() as db:
        cid, _ = company_context(db, request)
        rows = db.execute(select(PurchaseInvoice, Partner).outerjoin(Partner, PurchaseInvoice.supplier_id == Partner.id)
                          .where(PurchaseInvoice.company_id == cid).order_by(PurchaseInvoice.id.desc()).limit(limit)).all()
        return {"items": [{"id": n.id, "nf": n.number, "chave": n.access_key, "fornecedor": p.name if p else None,
                            "emissao": n.issue_date, "valor_erpflex": n.total_erpflex, "valor_nfstock": n.total_nfstock,
                            "conciliacao": n.reconciliation_status} for n, p in rows]}


@app.get("/api/v1/entregas")
def api_entregas(request: Request, status: str = "", limit: int = 200):
    limit = max(1, min(500, limit))
    with SessionLocal() as db:
        cid, _ = company_context(db, request)
        rows = _delivery_rows(db, cid, status=status, limit=limit)
        return {"items": [{"id": d.id, "nf": n.number, "cliente": p.name if p else None, "regiao": d.region,
                            "responsavel": r.name if r else d.responsible, "status": d.status, "volumes": d.volumes or n.volumes,
                            "valor_nf": n.total} for d, n, p, _a, r in rows]}
