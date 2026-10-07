from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import func, select

from .db import (
    Delivery, DeliveryAssignment, DeliveryRule, LogisticsImport, Partner, RawRecord,
    Responsible, SalesInvoice,
)
from .logistics import alerta_endereco, classificar_regiao, digits, regiao_especial


def _norm_header(value) -> str:
    txt = "" if value is None else str(value).strip().lower()
    txt = unicodedata.normalize("NFKD", txt)
    txt = "".join(c for c in txt if not unicodedata.combining(c))
    txt = re.sub(r"[^a-z0-9]+", " ", txt).strip()
    return txt


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _fnum(value) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    txt = str(value).strip().replace("R$", "").replace(" ", "")
    if not txt:
        return 0.0
    if "," in txt:
        txt = txt.replace(".", "").replace(",", ".")
    try:
        return float(txt)
    except Exception:
        return 0.0


def _date(value) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    txt = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(txt[:10], fmt).strftime("%Y-%m-%d")
        except Exception:
            pass
    return txt[:40] or None


ALIASES = {
    "nf": ["nfe", "nf", "numero nf", "numero_nf"],
    "fantasia": ["fantasia", "nome fantasia", "nome_fantasia"],
    "razao": ["razao social", "razao_social", "cliente", "cliente razao social"],
    "cnpj": ["cnpj", "cpf cnpj", "cpf_cnpj"],
    "endereco": ["endereco de entrega", "endereco entrega", "endereco", "logradouro entrega", "logradouro"],
    "numero": ["numero", "numero endereco", "numero_endereco"],
    "complemento": ["complemento"],
    "bairro": ["bairro", "bairro entrega"],
    "cidade": ["municipio", "cidade", "municipio entrega", "cidade entrega"],
    "cep": ["cep", "cep entrega"],
    "uf": ["uf", "estado", "uf entrega"],
    "transportadora": ["transportadora", "transportadora 1", "transportadora1"],
    "transportadora2": ["transportadora 2", "transportadora2", "modalidade", "tipo entrega"],
    "canal": ["canal", "canal venda"],
    "volumes": ["quant volume", "qtd volume", "volumes", "volume"],
    "valor": ["valor total produto nf", "valor nf", "valor total", "total nf", "valor"],
    "previsao": ["previsao faturamento", "previsao entrega", "data entrega", "data_entrega"],
    "emissao": ["emissao", "data emissao", "data_emissao"],
}


def _map_row(row: dict[str, object]) -> dict[str, object]:
    normalized = {_norm_header(k): v for k, v in row.items() if k is not None}
    out = {}
    for key, aliases in ALIASES.items():
        val = None
        for alias in aliases:
            nk = _norm_header(alias)
            if nk in normalized and normalized[nk] not in (None, ""):
                val = normalized[nk]
                break
        out[key] = val
    return out


def _iter_rows(filename: str, data: bytes):
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        iterator = ws.iter_rows(values_only=True)
        try:
            headers = [_text(v) for v in next(iterator)]
        except StopIteration:
            return
        for row in iterator:
            yield {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
    elif lower.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=";,\t,")
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ";"
        yield from csv.DictReader(io.StringIO(text), delimiter=delimiter)
    else:
        raise ValueError("Formato não aceito. Envie XLSX ou CSV.")


def _partner(db, company_id: int, *, cnpj: str, razao: str, fantasia: str, address: str, district: str, city: str, state: str, zip_code: str):
    cnpj = digits(cnpj) or None
    name = razao.strip() or fantasia.strip() or "Não identificado"
    row = None
    if cnpj:
        row = db.scalar(select(Partner).where(Partner.company_id == company_id, Partner.cnpj_cpf == cnpj).limit(1))
    if not row and name:
        row = db.scalar(select(Partner).where(Partner.company_id == company_id, func.lower(Partner.name) == name.lower()).limit(1))
    if not row:
        row = Partner(company_id=company_id, cnpj_cpf=cnpj, name=name, trade_name=fantasia or None, role_customer=True)
        db.add(row); db.flush()
    row.role_customer = True
    row.cnpj_cpf = row.cnpj_cpf or cnpj
    if razao:
        row.name = razao
    if fantasia:
        row.trade_name = fantasia
    row.address = address or row.address
    row.district = district or row.district
    row.city = city or row.city
    row.state = state or row.state
    row.zip_code = zip_code or row.zip_code
    return row


def _apply_rule(db, company_id: int, delivery: Delivery, customer: Partner | None):
    if not customer:
        return
    rule = None
    if customer.cnpj_cpf:
        rule = db.scalar(select(DeliveryRule).where(DeliveryRule.company_id == company_id, DeliveryRule.active.is_(True), DeliveryRule.customer_cnpj == customer.cnpj_cpf).order_by(DeliveryRule.id.desc()).limit(1))
    if not rule and customer.name:
        rule = db.scalar(select(DeliveryRule).where(DeliveryRule.company_id == company_id, DeliveryRule.active.is_(True), func.lower(DeliveryRule.customer_name) == customer.name.lower()).order_by(DeliveryRule.id.desc()).limit(1))
    if not rule:
        return
    if rule.region and (not delivery.region or delivery.region == "A Classificar"):
        delivery.region = rule.region
    if rule.box_type and not delivery.box_type:
        delivery.box_type = rule.box_type
    if rule.responsible_id:
        current = db.scalar(select(DeliveryAssignment).where(DeliveryAssignment.delivery_id == delivery.id).limit(1))
        if not current:
            resp = db.get(Responsible, rule.responsible_id)
            if resp and resp.active:
                db.add(DeliveryAssignment(company_id=company_id, delivery_id=delivery.id, responsible_id=resp.id))
                delivery.responsible = resp.name
                delivery.responsible_type = resp.kind


def import_logistics_file(db, company_id: int, filename: str, data: bytes) -> dict:
    if not data:
        raise ValueError("Arquivo vazio.")
    file_hash = hashlib.sha256(data).hexdigest()
    rows = []
    for raw in _iter_rows(filename, data):
        rows.append(_map_row(raw))
    if not rows:
        raise ValueError("O arquivo não possui linhas de dados.")

    grouped: dict[str, dict] = {}
    ignored = 0
    for row in rows:
        nf = _text(row.get("nf"))
        if not nf or nf == "0":
            ignored += 1
            continue
        g = grouped.setdefault(nf, {
            "nf": nf, "fantasia": "", "razao": "", "cnpj": "", "endereco": "", "bairro": "",
            "cidade": "", "cep": "", "uf": "", "transportadora": "", "transportadora2": "", "canal": "",
            "volumes": 0.0, "valor": 0.0, "previsao": None, "emissao": None,
        })
        endereco = _text(row.get("endereco"))
        numero = _text(row.get("numero"))
        complemento = _text(row.get("complemento"))
        if numero and numero not in endereco:
            endereco = f"{endereco}, {numero}" if endereco else numero
        if complemento:
            endereco = f"{endereco}, {complemento}" if endereco else complemento
        updates = {
            "fantasia": _text(row.get("fantasia")), "razao": _text(row.get("razao")), "cnpj": _text(row.get("cnpj")),
            "endereco": endereco, "bairro": _text(row.get("bairro")), "cidade": _text(row.get("cidade")),
            "cep": _text(row.get("cep")), "uf": _text(row.get("uf")).upper(), "transportadora": _text(row.get("transportadora")),
            "transportadora2": _text(row.get("transportadora2")), "canal": _text(row.get("canal")),
        }
        for k, v in updates.items():
            if v and not g.get(k):
                g[k] = v
        vol = _fnum(row.get("volumes"))
        if vol:
            g["volumes"] = max(g["volumes"], vol)
        g["valor"] += _fnum(row.get("valor"))
        g["previsao"] = _date(row.get("previsao")) or g["previsao"]
        g["emissao"] = _date(row.get("emissao")) or g["emissao"]

    if not grouped:
        raise ValueError("Nenhuma NF válida foi encontrada. Verifique a coluna NFe/NF.")

    batch = LogisticsImport(company_id=company_id, filename=filename[:255], file_hash=file_hash, status="EM_EXECUCAO", total_rows=len(rows), ignored=ignored)
    db.add(batch); db.flush()
    inserted = updated = 0
    errors = []

    for nf, g in grouped.items():
        try:
            customer = _partner(db, company_id, cnpj=g["cnpj"], razao=g["razao"], fantasia=g["fantasia"], address=g["endereco"], district=g["bairro"], city=g["cidade"], state=g["uf"], zip_code=g["cep"])
            raw_payload = {"arquivo": filename, **g}
            ext = f"{file_hash[:16]}:{nf}"
            raw = db.scalar(select(RawRecord).where(RawRecord.company_id == company_id, RawRecord.source == "EXCEL", RawRecord.module == "faturamento", RawRecord.external_id == ext).limit(1))
            payload_json = json.dumps(raw_payload, ensure_ascii=False, default=str)
            rec_hash = hashlib.sha256(payload_json.encode("utf-8")).hexdigest()
            if not raw:
                raw = RawRecord(company_id=company_id, source="EXCEL", module="faturamento", external_id=ext, record_hash=rec_hash, payload_json=payload_json)
                db.add(raw); db.flush()
            else:
                raw.record_hash = rec_hash; raw.payload_json = payload_json; raw.updated_at = datetime.utcnow()

            invoice = db.scalar(select(SalesInvoice).where(SalesInvoice.company_id == company_id, SalesInvoice.number == nf).order_by(SalesInvoice.id).limit(1))
            if not invoice:
                invoice = SalesInvoice(company_id=company_id, erpflex_id=f"EXCEL:{file_hash[:12]}:{nf}", number=nf)
                db.add(invoice); db.flush(); inserted += 1
            else:
                updated += 1
            invoice.customer_id = customer.id
            invoice.issue_date = g["emissao"] or invoice.issue_date
            invoice.forecast_date = g["previsao"] or invoice.forecast_date
            invoice.total = g["valor"] or invoice.total
            invoice.carrier = g["transportadora"] or invoice.carrier
            invoice.carrier2 = g["transportadora2"] or invoice.carrier2
            invoice.delivery_address = g["endereco"] or invoice.delivery_address
            invoice.delivery_district = g["bairro"] or invoice.delivery_district
            invoice.delivery_city = g["cidade"] or invoice.delivery_city
            invoice.delivery_state = g["uf"] or invoice.delivery_state
            invoice.delivery_zip = g["cep"] or invoice.delivery_zip
            invoice.volumes = g["volumes"] or invoice.volumes
            invoice.raw_record_id = raw.id

            delivery = db.scalar(select(Delivery).where(Delivery.sales_invoice_id == invoice.id).limit(1))
            if not delivery:
                delivery = Delivery(company_id=company_id, sales_invoice_id=invoice.id, volumes=invoice.volumes or 0, box_type="Papelão")
                db.add(delivery); db.flush()
            elif invoice.volumes and not delivery.volumes:
                delivery.volumes = invoice.volumes
            _apply_rule(db, company_id, delivery, customer)

            special = regiao_especial(invoice.carrier, invoice.carrier2)
            region = special or classificar_regiao(invoice.delivery_city, invoice.delivery_district, invoice.delivery_zip, g["canal"])
            if not delivery.region or delivery.region == "A Classificar":
                delivery.region = region
            delivery.box_type = delivery.box_type or "Papelão"
            alert = alerta_endereco(invoice.carrier, invoice.carrier2)
            if alert and alert not in (delivery.notes or "").upper():
                delivery.notes = ((delivery.notes or "") + (" | " if delivery.notes else "") + alert).strip()
            if delivery.region == "COLETA":
                current = db.scalar(select(DeliveryAssignment).where(DeliveryAssignment.delivery_id == delivery.id).limit(1))
                if not current:
                    coleta = db.scalar(select(Responsible).where(Responsible.company_id == company_id, Responsible.name == "COLETA", Responsible.active.is_(True)).limit(1))
                    if coleta:
                        db.add(DeliveryAssignment(company_id=company_id, delivery_id=delivery.id, responsible_id=coleta.id))
                        delivery.responsible = coleta.name; delivery.responsible_type = coleta.kind
        except Exception as exc:
            errors.append(f"NF {nf}: {exc}")

    batch.total_invoices = len(grouped)
    batch.inserted = inserted
    batch.updated = updated
    batch.status = "CONCLUIDO_COM_ERROS" if errors else "CONCLUIDO"
    batch.message = "\n".join(errors[:20]) if errors else None
    db.flush()
    return {
        "batch_id": batch.id, "filename": filename, "rows": len(rows), "invoices": len(grouped),
        "inserted": inserted, "updated": updated, "ignored": ignored, "errors": errors,
    }
